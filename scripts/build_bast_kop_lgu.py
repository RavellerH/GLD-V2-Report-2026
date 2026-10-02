"""Draf BAST Termin 1 Field Testing dengan kop format laporan LGU (logo LGU | judul | logo Pertamina),
mengikuti header laporan URS LGU. Isi BAST tidak diubah.

Output: Paket LGU/02_Penagihan_Termin_1/02_Draft_BAST_Termin_1_Field_Testing_20Persen.docx
(PDF dirender lewat Word COM).
"""
import os

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
ASSET = os.path.join(HERE, "assets", "lgu_cover")
SRC = os.path.join(ROOT, "Paket Pertamina", "04_Laporan_Termin_1", "Draft_BAST_Termin_1_FieldTesting_20Persen.docx")
OUT = os.path.join(ROOT, "Paket LGU", "02_Penagihan_Termin_1", "02_Draft_BAST_Termin_1_Field_Testing_20Persen.docx")

JUDUL = "Berita Acara Serah Terima Pekerjaan Termin 1"
SUB = "Pengembangan dan Field Testing Sistem Gas Leak Detection Tahap 2"


def borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)


def kop(header):
    for p in list(header.paragraphs):
        p._p.getparent().remove(p._p)
    t = header.add_table(rows=1, cols=3, width=Inches(6.5))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    widths = [Inches(1.5), Inches(3.5), Inches(1.5)]
    for c, w in zip(t.rows[0].cells, widths):
        c.width = w
    c0, c1, c2 = t.rows[0].cells
    p0 = c0.paragraphs[0]; p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.add_run().add_picture(os.path.join(ASSET, "logo_lgu.png"), height=Inches(0.42))
    p1 = c1.paragraphs[0]; p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p1.add_run(JUDUL); r.font.size = Pt(8.5); r.italic = True
    p1b = c1.add_paragraph(); p1b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p1b.add_run(SUB); r.font.size = Pt(8.5); r.italic = True
    p2 = c2.paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run().add_picture(os.path.join(ASSET, "logo_pertamina.png"), height=Inches(0.32))
    for c in (c0, c1, c2):
        tcPr = c._tc.get_or_add_tcPr()
        va = OxmlElement("w:vAlign"); va.set(qn("w:val"), "center"); tcPr.append(va)
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
    header.add_paragraph().paragraph_format.space_after = Pt(0)


d = docx.Document(SRC)
for s in d.sections:
    s.different_first_page_header_footer = False
    s.header.is_linked_to_previous = False
    s.header_distance = Inches(0.35)
    if s.top_margin < Inches(1.25):
        s.top_margin = Inches(1.25)
    kop(s.header)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
d.save(OUT)
print("written", OUT)
