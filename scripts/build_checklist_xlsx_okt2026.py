# -*- coding: utf-8 -*-
"""Dua workbook Excel (28 Okt 2026):

1. Deliverables/Checklist_Instalasi_Block_Diagram_RU-IV.xlsx
   Checklist dari block diagram lingkup final RU IV (Field/Building, LGU vs PTMN).
2. Deliverables/Checklist_Rencana_Sertifikasi_GLD_EMC_IP_RF_ATEX.xlsx
   Checklist EMC · Rencana uji IP · Rencana uji RF · Monitor ATEX + log.

Status dipilih dari dropdown; ringkasan dihitung dengan rumus COUNTIF.
Run: python3 scripts/build_checklist_xlsx_okt2026.py
"""
import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DELIV = os.path.join(REPO, "Deliverables")
DIAGRAM = os.path.join(REPO, "Sumber Dokumen", "Pertamina RU IV 8Okt2026",
                       "block_diagram_scope_final_router_LGU_8Okt2026.png")
DATE = "28 Oktober 2026"

NAVY = "1B2A4A"
BLUE = "2B5FCB"
THIN = Side(style="thin", color="C9CED6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {
    "Selesai": ("EAF4EC", "1E7B3A"),
    "Proses": ("FFF3DC", "9A6200"),
    "Sebagian": ("FFF3DC", "9A6200"),
    "Menunggu": ("E8EEFA", "2B5FCB"),
    "Belum": ("FBE6E4", "B3261E"),
    "N.A.": ("EEEEEE", "666666"),
}


def title_block(ws, title, subtitle, ncols):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, size=9.5, color="555B66")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 42


def write_table(ws, start_row, head, rows, widths, status_col=None, statuses=None):
    for i, h in enumerate(head, 1):
        c = ws.cell(row=start_row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BORDER
    for r, row in enumerate(rows, start_row + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.alignment = WRAP
            c.border = BORDER
            c.font = Font(size=9.5)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    end = start_row + len(rows)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=3)
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(head))}{end}"
    if status_col:
        col = get_column_letter(status_col)
        rng = f"{col}{start_row + 1}:{col}{end + 50}"
        dv = DataValidation(type="list", formula1='"' + ",".join(statuses) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(rng)
        for s in statuses:
            bg, fg = STATUS_FILL[s]
            ws.conditional_formatting.add(rng, CellIsRule(
                operator="equal", formula=[f'"{s}"'],
                fill=PatternFill("solid", fgColor=bg), font=Font(bold=True, color=fg)))
    return end


def summary(ws, row, label, sheet, col, first, last, statuses):
    ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=11, color=NAVY)
    row += 1
    for i, h in enumerate(["Status", "Jumlah", "%"], 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=BLUE); c.border = BORDER
    rng = f"'{sheet}'!${col}${first}:${col}${last}"
    total_row = row + len(statuses) + 1
    for s in statuses:
        row += 1
        ws.cell(row=row, column=1, value=s).border = BORDER
        c = ws.cell(row=row, column=2, value=f'=COUNTIF({rng},"{s}")'); c.border = BORDER
        c = ws.cell(row=row, column=3, value=f"=IF($B${total_row}=0,0,B{row}/$B${total_row})")
        c.number_format = "0%"; c.border = BORDER
        bg, fg = STATUS_FILL[s]
        ws.cell(row=row, column=1).fill = PatternFill("solid", fgColor=bg)
        ws.cell(row=row, column=1).font = Font(bold=True, color=fg)
    row += 1
    ws.cell(row=row, column=1, value="Total").font = Font(bold=True)
    c = ws.cell(row=row, column=2, value=f"=SUM(B{row - len(statuses)}:B{row - 1})")
    c.font = Font(bold=True)
    return row + 2


# ======================================================================= 1. block diagram
BD_STATUS = ["Belum", "Proses", "Menunggu", "Selesai", "N.A."]
BD_HEAD = ["No", "Area", "Komponen / jalur", "Spesifikasi (dari diagram)", "Item yang dicek",
           "Terkait pihak lain", "Status", "PIC", "Tanggal cek", "Bukti (foto / dokumen)", "Catatan"]
FAT = "Laporan FAT GLD Tahap 2 (LGU/GLD/FAT/2026-001)"
SPEK = "Spesifikasi Material Instalasi rev 1.5"
PERANGKAT = "Daftar Perangkat & Kebutuhan Pemasangan rev 1.2"
WAIT = "Dilakukan di lokasi setelah pekerjaan RU IV selesai"
# (area, komponen, spesifikasi, item cek, terkait, status, bukti, catatan)
LGU_ROWS = [
    ("Field", "GL Detector", "Node GLD V2, catu 24 VDC (disuplai LGU)",
     "Unit tersedia; nomor seri tercatat; uji fungsi (sensor, AI, alarm, LoRa)", "—", "Selesai",
     f"{FAT} FAT-01, FAT-02, FAT-06; {PERANGKAT}", "FAT 7/7 lulus (lab)"),
    ("Field", "GL Detector — cable entry", "Gland M20×1,5 (butir 5)",
     "Gland sesuai kabel 2C 16AWG tersedia di unit", "RU IV memasukkan kabel (R-07)", "Selesai",
     f"{SPEK} (terminasi L+/L−); {PERANGKAT}", ""),
    ("Field", "GL Detector — antena & konfigurasi", "LoRa 920 MHz ke CH",
     "Antena terpasang; ID node & frekuensi terkonfigurasi", "—", "Selesai", f"{FAT} FAT-03", ""),
    ("Field", "Cluster Head (CH)", "CH, LoRa dari GLD dan ke Gateway",
     "Unit tersedia & teruji; konfigurasi star/mesh", "RU IV memasang (R-08)", "Selesai",
     f"{FAT} FAT-04; {PERANGKAT}", "CH tersedia > 5 unit"),
    ("Field", "Solar cell & bracket CH", "Solar cell → CH",
     "Panel, bracket 2 panel & klem CH ke tiang 2\" disiapkan LGU", "Tiang CH dari RU IV (R-09)", "Selesai",
     f"{SPEK} BoQ no. 8 & Gambar 3.1 (foto uji)", ""),
    ("Building", "Gateway", "Gateway (Provided by LGU), 5 VDC (butir 6)",
     "Unit tersedia; frekuensi 921 MHz; menerima data CH", "RU IV memasang (R-10)", "Selesai",
     f"{FAT} FAT-05; {PERANGKAT}", "Antena fiber omni 8 dBi"),
    ("Building", "Adaptor Gateway", "Adaptor 5 VDC, catu dari UPS",
     "Adaptor tersedia; tegangan keluaran sesuai", "Outlet UPS dari RU IV (R-12)", "Selesai", SPEK, ""),
    ("Building", "WiFi Router 2,4 GHz", "WiFi Router 2.4 GHz (Provided by LGU) (butir 7)",
     "Router tersedia & dikonfigurasi (SSID GLD, WPA2, tanpa internet)", "—", "Selesai", f"{SPEK} §5.2", ""),
    ("Building", "Adaptor Router", "Adaptor router, catu dari UPS",
     "Adaptor tersedia & sesuai router", "Outlet UPS dari RU IV (R-12)", "Selesai", f"{SPEK} §5", ""),
    ("Building", "Spesifikasi jaringan lokal", "LAN Cat6 lokal tanpa internet (butir 8)",
     "Skema GW → router → PC server & kebutuhan kabel Cat6 diserahkan ke RU IV", "Kabel Cat6 dari RU IV (R-16)",
     "Selesai", f"{SPEK} §5.2 (dikirim ke grup Pertamina 1 Okt)", ""),
    ("Building", "Paket aplikasi server", "Server (disediakan RU IV)",
     "Broker MQTT, backend & dashboard GLD siap dipasang", "Server dari RU IV (R-13)", "Selesai",
     f"{FAT} FAT-05, FAT-07", ""),
    ("Field", "Dokumen supervisi/QA pemasangan", "Pekerjaan instalasi oleh RU IV",
     "Pembagian tugas, daftar persiapan pelaksana & spesifikasi material diserahkan", "R-01…R-10", "Selesai",
     "Pembagian Persiapan Instalasi rev 1.5; Daftar Persiapan Pelaksana Instalasi rev 1.2", ""),
    ("Field", "Prosedur terminasi & energize GLD", "24 VDC dari PSU di junction box",
     "Prosedur terminasi L+/L−, cek polaritas & tegangan tertulis", "R-04…R-06", "Selesai",
     f"{SPEK}; Instruction Manual GLD", ""),
    ("Sistem", "Uji link LoRa GLD → CH (lab)", "Sinyal LoRa (wireless)", "PDR & RSSI terukur", "—", "Selesai",
     f"{FAT} FAT-03; Laporan Uji Lab 02", "PDR 100% pada 177 m & 243 m"),
    ("Sistem", "Uji link LoRa CH → Gateway & mesh (lab)", "Sinyal LoRa (wireless)",
     "Mesh, failover & downlink teruji", "—", "Selesai", f"{FAT} FAT-04; Laporan Uji Lab 03", ""),
    ("Sistem", "Uji link WiFi Gateway → Router (lab)", "WiFi 2,4 GHz", "Gateway terhubung ke router", "—",
     "Selesai", f"{FAT} FAT-05", ""),
    ("Sistem", "Uji end-to-end & alarm (lab)", "GLD → CH → Gateway → Router → Server → Dashboard",
     "Data normal tampil; uji gas → alarm otomatis", "—", "Selesai", f"{FAT} FAT-05…FAT-07", ""),
    ("Sistem", "Kesiapan persiapan LGU", "Persiapan instalasi",
     "Seluruh item persiapan sisi LGU & ITB", "—", "Selesai",
     "Kurva-S Persiapan Instalasi RU IV: LGU & ITB 100%", "Tidak termasuk izin masuk kilang (sisi RU IV)"),
    ("Field", "Mounting CH", "Tidak digambarkan di diagram",
     "Pembagian mounting CH ditetapkan", "R-09", "Selesai",
     f"{SPEK} BoQ no. 2 & 8", "Tiang 2\" Sch40 4,0 m = RU IV; bracket panel & klem CH = LGU"),
    # --- pelaksanaan di lokasi
    ("Lokasi", "Supervisi/QA saat pemasangan", "Pekerjaan RU IV",
     "QA posisi, orientasi & kekencangan GLD/CH/antena; foto tiap titik", "R-01…R-10", "Menunggu", "", WAIT),
    ("Lokasi", "Terminasi & energize GLD", "24 VDC dari PSU",
     "Polaritas & tegangan terukur di terminal GLD; GLD menyala normal", "R-04…R-07", "Menunggu", "", WAIT),
    ("Lokasi", "Pasang aplikasi di PC server RU IV", "Server", "Aplikasi berjalan di PC server", "R-13", "Menunggu",
     "", WAIT),
    ("Lokasi", "Commissioning (link & alarm di lokasi)", "End-to-end",
     "RSSI/SNR tiap link tercatat; uji gas → alarm; waktu respons dicatat", "Disaksikan RU IV (R-18)", "Menunggu",
     "", WAIT),
    ("Lokasi", "Berita acara instalasi", "Serah terima", "Foto as-built, nomor seri, hasil uji",
     "Ditandatangani RU IV (R-19)", "Menunggu", "", WAIT),
]
RU_ROWS = [
    ("Field", "Stanchion GLD", "Stanchion support 2\" HDG (butir 1)",
     "Stanchion berdiri tegak & kokoh di titik sesuai layout", "—"),
    ("Field", "Mounting support & pemasangan GLD", "Mounting support by PTMN",
     "GLD terpasang dengan plat U-bolt; tinggi sensor sesuai rencana", "QA oleh LGU (L-20)"),
    ("Field", "Junction box", "Junction box SS316 (butir 2)",
     "JB terpasang; entry kabel rapat; grounding JB tersambung", "—"),
    ("Field", "PSU", "PSU 220VAC in, 24VDC out (butir 3), di dalam JB",
     "Merek/model tercatat; rating ≥ 1 A per GLD; keluaran 24 VDC terukur", "Energize oleh LGU (L-21)"),
    ("Field", "Local receptacle", "220 VAC ke PSU",
     "Sumber 220 VAC tersedia; dilindungi MCB; L/N/PE benar", "—"),
    ("Field", "Kabel daya 24 VDC", "Power cable 24VDC, 2C, 16AWG (butir 4)",
     "Kabel tergelar & terlindung; terminasi di JB; panjang tercatat", "—"),
    ("Field", "Kabel masuk GLD", "Melalui gland M20×1,5",
     "Kabel dimasukkan lewat gland & dikencangkan", "Gland dari LGU (L-02)"),
    ("Field", "Pemasangan CH & solar cell", "CH + solar cell (disuplai LGU)",
     "CH & panel terpasang di atas tiang; line-of-sight ke GLD & Gateway", "QA oleh LGU (L-20)"),
    ("Field", "Tiang CH", "Tidak digambarkan di diagram",
     "Tiang galvanis 2\" Sch40 4,0 m terpasang & dicor (Spesifikasi BoQ no. 2)", "Bracket & klem dari LGU (L-05)"),
    ("Building", "Mounting support Gateway", "Mounting support by PTMN",
     "Lokasi Gateway disetujui; dudukan terpasang; antena dapat menjangkau CH", "Gateway dari LGU (L-06)"),
    ("Building", "UPS", "UPS mencatu Gateway, Router & Server",
     "Kapasitas cukup untuk 3 beban; waktu cadangan tercatat", "—"),
    ("Building", "Outlet UPS", "Untuk adaptor Gateway & Router + Server 220 VAC",
     "Outlet UPS tersedia di dekat Gateway, router & server", "Adaptor dari LGU (L-07, L-09)"),
    ("Building", "Server", "Server (catu 220 VAC dari UPS)",
     "PC server tersedia (usulan: 4 core, 8 GB, SSD 256 GB, 2 port LAN)", "Aplikasi dipasang LGU (L-22)"),
    ("Building", "Ruang server", "Building",
     "Ruang/meja untuk server, router & Gateway; akses bagi tim LGU", "—"),
    ("Building", "Server → IT Network", "IT Network (intranet kilang)",
     "Port LAN ke intranet tersedia; persetujuan tim IT; dashboard dapat diakses dari kantor", "—"),
    ("Building", "Kabel LAN Cat6 router → server", "LAN Cable Cat6 lokal, tanpa internet (butir 8)",
     "Kabel Cat6 tergelar antar-ruang (≤ 90 m per segmen) & berlabel", "Spesifikasi dari LGU (L-10)"),
    ("Umum", "Izin kerja & HSE", "Seluruh pekerjaan instalasi",
     "Izin masuk kilang, izin kerja, TRA/JSA, klasifikasi area titik pemasangan disetujui", "—"),
    ("Sistem", "Saksi commissioning", "Commissioning",
     "Perwakilan RU IV hadir saat uji alarm", "Uji oleh LGU (L-23)"),
    ("Sistem", "Berita acara", "Serah terima",
     "Berita acara instalasi & commissioning ditandatangani", "Draf dari LGU (L-24)"),
]
AREA_FILL = {"Field": "EAF2FD", "Building": "FFF8E1", "Sistem": "F2F2F2", "Umum": "F2F2F2", "Lokasi": "EDE7F6"}


def checklist_sheet(wb, name, title, subtitle, prefix, rows, tab_color):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab_color
    title_block(ws, title, subtitle, len(BD_HEAD))
    data = []
    for i, row in enumerate(rows):
        a, k, s, cek, rel = row[:5]
        st, bukti, cat = (row[5:] + ("Belum", "", ""))[:3] if len(row) > 5 else ("Belum", "", "")
        data.append([f"{prefix}-{i + 1:02d}", a, k, s, cek, rel, st, "LGU" if prefix == "L" and st == "Selesai"
                     else "", DATE if st == "Selesai" else "", bukti, cat])
    end = write_table(ws, 4, BD_HEAD, data, [8, 10, 28, 30, 44, 24, 11, 10, 16, 40, 30], 7, BD_STATUS)
    for r in range(5, end + 1):
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=AREA_FILL[ws.cell(row=r, column=2).value])
    return end


def build_block_diagram():
    wb = Workbook()
    sm = wb.active
    sm.title = "Ringkasan"
    sub = (f"Disusun dari block diagram lingkup final. Status per {DATE}. Pilih status dari dropdown kolom G. Kolom 'Terkait pihak lain' merujuk ID di checklist pihak lain.")
    l_end = checklist_sheet(wb, "Checklist LGU", "Checklist LGU — Penyedia (Supplier) Sistem GLD, RU IV Cilacap",
                            "LGU: menyediakan GLD, CH, solar cell, Gateway, router WiFi & adaptor; konfigurasi "
                            "perangkat & aplikasi; supervisi/QA pemasangan; energize & commissioning. " + sub,
                            "L", LGU_ROWS, "2B5FCB")
    r_end = checklist_sheet(wb, "Checklist RU IV Pertamina",
                            "Checklist RU IV Pertamina — Instalasi, Mounting Support & Infrastruktur",
                            "RU IV (PTMN): stanchion, mounting support, junction box, PSU, kabel 24 VDC, pemasangan "
                            "perangkat, UPS, server, jaringan IT, izin kerja & HSE. " + sub,
                            "R", RU_ROWS, "E0B000")

    sm["A1"] = "Ringkasan status checklist instalasi GLD — RU IV Cilacap"
    sm["A1"].font = Font(bold=True, size=14, color=NAVY)
    r = summary(sm, 3, "Checklist LGU", "Checklist LGU", "G", 5, l_end, BD_STATUS)
    r = summary(sm, r, "Checklist RU IV Pertamina", "Checklist RU IV Pertamina", "G", 5, r_end, BD_STATUS)
    notes = ["Persiapan LGU selesai (L-01…L-19, bukti di kolom J). L-20…L-24 = pekerjaan LGU di lokasi yang baru "
             "bisa dilakukan setelah pekerjaan RU IV selesai (status 'Menunggu').",
             "Warna lingkup (sesuai diagram): biru = disuplai LGU; kuning = disuplai/dipasang PTMN (RU IV).",
             "Jalur: merah = daya (220 VAC / 24 VDC); putus-putus biru = LoRa; titik biru = WiFi 2,4 GHz; "
             "biru tegas = LAN Cat6 lokal tanpa internet.",
             "Dari Spesifikasi Material rev 1.5: tiang CH 2\" Sch40 4,0 m = RU IV, bracket panel & klem CH = LGU; "
             "kabel Cat6 router ↔ PC server = RU IV."]
    for n in notes:
        sm.cell(row=r, column=1, value=n).alignment = Alignment(wrap_text=True, vertical="top")
        sm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        sm.row_dimensions[r].height = 30
        r += 1
    sm.column_dimensions["A"].width = 26
    for c in "BCDEF":
        sm.column_dimensions[c].width = 14

    dg = wb.create_sheet("Diagram")
    dg["A1"] = "Block diagram lingkup (versi final, sumber checklist)"
    dg["A1"].font = Font(bold=True, size=12, color=NAVY)
    img = XLImage(DIAGRAM)
    scale = 1100 / img.width
    img.width, img.height = int(img.width * scale), int(img.height * scale)
    dg.add_image(img, "A3")

    out = os.path.join(DELIV, "Checklist_Instalasi_Block_Diagram_RU-IV.xlsx")
    wb.save(out)
    print("written", out)

# ======================================================================= 2. sertifikasi
ST = ["Belum", "Proses", "Sebagian", "Menunggu", "Selesai", "N.A."]

EMC = [
    ("EMC-01", "Pre-scan", "Emisi radiasi GLD 30 MHz–1 GHz (QP)", "EN 55032 Kelas A", "Selesai",
     "Margin terkecil 0,51 dB @ 117,9 MHz (H, 300 cm) — kritis", "Lab IoT ITB", "Laporan EMC bab 2.3"),
    ("EMC-02", "Pre-scan", "Emisi radiasi GLD 1–6 GHz (Peak/AV)", "EN 55032 Kelas A", "Selesai",
     "Margin ≥ 33,97 dB", "Lab IoT ITB", "Laporan EMC bab 2.4"),
    ("EMC-03", "Pre-scan", "Emisi konduksi GLD sisi AC PSU (L1/L2)", "EN 55032 Kelas A", "Selesai",
     "Margin ≥ 6,03 dB @ 323,9 kHz", "Lab IoT ITB", "Laporan EMC bab 2.1–2.2"),
    ("EMC-04", "Pre-scan", "Emisi radiasi CH besar 30 MHz–6 GHz", "EN 55032 Kelas A", "Selesai",
     "Margin 3,28 dB @ 978,9 MHz; 5,56 dB @ 1,8425 GHz (harmonik)", "Lab IoT ITB", "Laporan EMC bab 3.2–3.3"),
    ("EMC-05", "Pre-scan", "Emisi konduksi port DC CH besar", "EN 55032 Kelas A", "Selesai",
     "Margin ≥ 42,41 dB", "Lab IoT ITB", "Laporan EMC bab 3.1"),
    ("EMC-06", "Pre-scan", "Emisi konduksi port 24 VDC GLD (AAN/LISN DC)", "EN 55032 Kelas A", "Belum",
     "Batas produk = masukan 24 VDC saja (temuan E-4)", "Lab IoT ITB", ""),
    ("EMC-07", "Mitigasi", "Mitigasi & re-scan GLD 117–127 MHz, target margin ≥ 6 dB", "EN 55032 Kelas A", "Belum",
     "Ferit kabel daya/fan, filter input, grounding (E-1)", "Lab IoT ITB", ""),
    ("EMC-08", "Konfirmasi", "Exclusion band LoRa / ulang ukur mode idle", "EN 301 489-1/-3", "Belum",
     "Konfirmasi ke Telkom Test House (E-2)", "LGU", ""),
    ("EMC-09", "Pre-scan", "Pre-scan Gateway (adaptor 5 VDC + WiFi)", "EN 55032", "Belum",
     "Perangkat di gedung", "Lab IoT ITB", ""),
    ("EMC-10", "Pre-scan", "Pre-scan CH kecil", "EN 55032", "Belum", "", "Lab IoT ITB", ""),
    ("EMC-11", "Imunitas", "ESD", "IEC 61000-4-2", "Belum", "Kontak/udara pada enclosure & konektor", "LGU", ""),
    ("EMC-12", "Imunitas", "Medan RF radiasi", "IEC 61000-4-3", "Sebagian",
     "Hanya uji interferensi co-channel LoRa 17 Sep (informal)", "LGU", "Laporan perkembangan 17 Sep"),
    ("EMC-13", "Imunitas", "EFT/burst & surge port 24 VDC", "IEC 61000-4-4 / -4-5", "Belum", "", "LGU", ""),
    ("EMC-14", "Imunitas", "RF konduksi & dip/interupsi DC", "IEC 61000-4-6 / -4-29", "Belum", "", "LGU", ""),
    ("EMC-15", "Persiapan", "Kriteria kinerja saat uji imunitas (tanpa alarm palsu / alarm hilang)",
     "EN 61000-6-2 / EN 50270", "Belum", "Ditetapkan sebelum uji", "LGU · Lab IoT ITB", ""),
    ("EMC-16", "Persiapan", "Identitas EUT tercatat (S/N, firmware, mode, kabel, PSU)", "—", "Belum",
     "Kekurangan draf (E-6)", "Lab IoT ITB", ""),
    ("EMC-17", "Persiapan", "Pilih lab terakreditasi & standar uji resmi", "—", "Belum", "", "LGU", ""),
    ("EMC-18", "Resmi", "Laporan uji EMC resmi", "—", "Belum", "", "Lab", ""),
]

IP = [
    ("IP-01", "Target", "Tetapkan target IP", "IEC 60529", "IP66 (target di dossier & nameplate usulan); IP67 opsional "
     "(proposal)", "Target tertulis, belum diuji", "Selesai", "LGU · PT Galaksi"),
    ("IP-02", "Desain", "Keputusan inlet gas (mesh depan + fan) terhadap IP66",
     "IEC 60529 §14 / IEC 60079-0", "Inlet terbuka berpotensi kemasukan air saat uji semprot. Opsi: membran "
     "hidrofobik (PTFE), pelindung semprot (weather guard), atau ruang sensor terpisah",
     "Keputusan desain tertulis + gambar", "Belum", "PT Galaksi · LGU"),
    ("IP-03", "Desain", "Rating komponen penutup", "IEC 60529", "Cable gland M20×1,5 Ex d berating ≥ IP66; blanking "
     "plug; seal bulkhead SMA antena; O-ring tutup (tali 2,70 mm, OD ≈ 91 mm)", "Datasheet/sertifikat tiap komponen",
     "Belum", "PT Galaksi"),
    ("IP-04", "Sampel", "Siapkan sampel konfigurasi produksi", "—", "Gland Ex d (bukan gland polimer prototipe), "
     "torsi tutup & gland tercatat", "Sampel sesuai gambar", "Belum", "PT Galaksi · LGU"),
    ("IP-05", "Pra-uji", "Pra-uji internal semprot air (setara IPX6)", "IEC 60529 §14.2.6", "Nozzle Ø12,5 mm, "
     "100 L/min ± 5 %, jarak 2,5–3 m, 1 menit/m² (min. 3 menit) dari semua arah", "Tidak ada air di bagian "
     "elektronik; unit berfungsi normal", "Belum", "Lab IoT ITB"),
    ("IP-06", "Pra-uji", "Inspeksi seal pasca pra-uji", "—", "Buka unit, kertas indikator air, foto bagian dalam",
     "Bukti foto & catatan", "Belum", "Lab IoT ITB"),
    ("IP-07", "Uji resmi", "Uji pendahuluan Ex: ketahanan termal bagian non-logam & uji impak",
     "IEC 60079-0 §26", "Panas/dingin pada gasket & bagian non-logam; impak sesuai Grup II; dilakukan SEBELUM uji IP",
     "Tidak ada kerusakan yang menurunkan proteksi", "Belum", "Lab ATEX (via GTS)"),
    ("IP-08", "Uji resmi", "Uji debu IP6X", "IEC 60529 §13.4", "Kategori 1 (tekanan negatif 2 kPa), bubuk talk, "
     "maks. 8 jam / 80 volume udara", "Tidak ada debu masuk", "Belum", "Lab ATEX (via GTS)"),
    ("IP-09", "Uji resmi", "Uji semprot air IPX6", "IEC 60529 §14.2.6", "Seperti IP-05", "Air yang masuk tidak "
     "mengganggu operasi & keselamatan", "Belum", "Lab ATEX (via GTS)"),
    ("IP-10", "Uji resmi", "Uji fungsi setelah uji IP", "—", "Pembacaan sensor, kirim LoRa, alarm", "Fungsi normal",
     "Belum", "Lab · LGU"),
    ("IP-11", "Koordinasi", "Konfirmasi ke GTS: uji IP66 masuk paket uji ATEX atau terpisah", "—",
     "Karena IP66 dicantumkan pada marking Ex, uji IP biasanya dilakukan lab Ex dalam urutan uji tipe",
     "Jawaban tertulis GTS", "Belum", "LGU"),
    ("IP-12", "Dokumen", "Laporan uji IP & pembaruan dossier/nameplate", "—", "", "Laporan resmi", "Belum", "LGU"),
]

RF = [
    ("RF-01", "Regulasi", "Tetapkan sertifikasi target", "SDPPI (Indonesia)", "Pita 920–923 MHz adalah pita "
     "Indonesia; pita SRD Eropa adalah 863–870 MHz sehingga sertifikasi radio CE untuk 920 MHz tidak relevan "
     "kecuali dibuat varian 868 MHz", "Keputusan tertulis", "Belum", "LGU"),
    ("RF-02", "Regulasi", "Verifikasi batas teknis 920–923 MHz (EIRP, bandwidth, duty cycle)",
     "Peraturan SDPPI/Kominfo terbaru", "Acuan umum LoRa untuk Indonesia: EIRP maks. 16 dBm — wajib "
     "dicek ke peraturan yang berlaku", "Tabel batas resmi", "Belum", "LGU"),
    ("RF-03", "Desain", "Hitung EIRP konfigurasi saat ini", "—", "GLD: 22 dBm + 3 dBi ≈ 25 dBm; CH mesh: 22 dBm + "
     "8 dBi ≈ 30 dBm. Pre-scan CH 112,47 dBµV/m @3 m ≈ 17 dBm EIRP (indikatif)", "EIRP ≤ batas resmi",
     "Belum", "Lab IoT ITB"),
    ("RF-04", "Desain", "Setel daya TX per jenis antena & uji ulang jangkauan", "—", "Contoh bila batas 16 dBm: "
     "GLD ≤ 13 dBm, CH mesh ≤ 8 dBm; ulang uji jangkauan/mesh", "Link budget tetap memadai", "Belum",
     "Lab IoT ITB"),
    ("RF-05", "Firmware", "Mode uji radio (TX kontinu / CW / termodulasi pada kanal terpilih)", "—",
     "Diperlukan lab untuk mengukur pemancar", "Firmware mode uji tersedia", "Belum", "Lab IoT ITB"),
    ("RF-06", "Pra-uji", "Daya keluaran konduksi di port antena", "SDPPI / EN 300 220-1 (metode)",
     "Spectrum analyzer + atenuator, 920 & 921 MHz", "Sesuai setelan & batas", "Belum", "Lab IoT ITB"),
    ("RF-07", "Pra-uji", "Kesalahan frekuensi & lebar pita terpakai (99 %)", "SDPPI", "LoRa BW 125 kHz",
     "Dalam pita & toleransi", "Belum", "Lab IoT ITB"),
    ("RF-08", "Pra-uji", "Emisi spurious & harmonik (konduksi + radiasi s.d. 6 GHz)", "SDPPI / EN 300 220-1",
     "Fokus harmonik ke-2 1,84 GHz & ke-3 2,76 GHz (CH)", "Di bawah batas spurious", "Belum", "Lab IoT ITB"),
    ("RF-09", "Pra-uji", "Duty cycle / laju kirim", "Peraturan terkait", "Hitung dari interval kirim normal (±90 dtk) "
     "& alarm (±10 dtk)", "Sesuai batas bila ada", "Belum", "Lab IoT ITB"),
    ("RF-10", "Dokumen", "Data antena (gain, pola) 3 dBi & 8 dBi", "—", "Datasheet pemasok", "Datasheet tersedia",
     "Belum", "LGU"),
    ("RF-11", "Dokumen", "Status sertifikasi modul E22-900MM22S", "—", "Sertifikat modul tidak menggantikan "
     "sertifikasi perangkat di Indonesia", "Data pemasok", "Belum", "LGU"),
    ("RF-12", "Regulasi", "Cek persyaratan TKDN perangkat LPWA", "Peraturan Kominfo/Kemenperin", "Perlu dicek",
     "Status jelas", "Belum", "LGU"),
    ("RF-13", "Koordinasi", "Tentukan pemohon sertifikat & daftar perangkat", "SDPPI", "Badan hukum Indonesia "
     "(LGU atau PT Galaksi); perangkat terpisah: GLD, CH besar, CH kecil, Gateway", "Keputusan tertulis", "Belum",
     "LGU"),
    ("RF-14", "Koordinasi", "Pilih lab uji perangkat telekomunikasi", "—", "Telkom Test House (sudah dipakai pre-scan "
     "EMC) atau lab lain yang diakui SDPPI; gabungkan dengan uji EMC radio", "Penawaran & jadwal", "Belum", "LGU"),
    ("RF-15", "Dokumen", "Berkas permohonan (manual, skematik, foto, block diagram, label)", "SDPPI", "",
     "Berkas lengkap", "Belum", "LGU"),
    ("RF-16", "Pendukung", "Uji ketahanan interferensi co-channel 10–25 dBm", "—", "RSSI stabil −17…−18 dBm "
     "(hasil awal informal, 17 Sep)", "Bukti pendukung", "Sebagian", "Lab Fisika ITB"),
    ("RF-17", "Pendukung", "Router WiFi 2,4 GHz komersial bersertifikat", "SDPPI", "Pastikan berlabel sertifikat",
     "Label terverifikasi", "Belum", "LGU"),
]

ATEX = [
    ("AX-01", "Pengajuan", "Form aplikasi & dossier dikirim ke GTS (agen sertifikasi)", "Selesai", "±29 Sep 2026",
     "Applicant = PT Galaksi; marking usulan II 2G Ex db IIC T4 Gb", "—", "PT Galaksi · LGU"),
    ("AX-02", "Pengajuan", "Umpan balik / iterasi dari GTS atas dossier", "Menunggu", "28 Okt 2026",
     "Belum tercatat ada umpan balik di repo", "Tanyakan status ke GTS", "LGU"),
    ("AX-03", "Pengajuan", "Perbaikan 8 kontradiksi dossier terkirim (marking debu, Tamb, PG13.5, Ex d, partisi, "
     "baterai, material, relay)", "Selesai", "8 Okt 2026", "Sudah dibetulkan di Dokumen Teknis Sertifikasi rev 2.8+",
     "Kirim ulang (AX-04)", "LGU"),
    ("AX-04", "Pengajuan", "Kirim ulang dossier terkoreksi ke GTS", "Belum", "28 Okt 2026", "",
     "Kirim bersama jawaban iterasi GTS", "LGU · PT Galaksi"),
    ("AX-05", "Sampel", "Sampel GLD2-0x1001 + surat pengantar GTS & surat Bea Cukai", "Proses", "12 Okt 2026",
     "Surat siap; status pengiriman/serah terima sampel belum dikonfirmasi", "Konfirmasi tanggal serah terima",
     "LGU"),
    ("AX-06", "Sampel", "Baterai 7×18650 di sampel — kapasitas per sel & keputusan baterai di luar produk",
     "Belum", "12 Okt 2026", ">3860 mAh/sel ⇒ >100 Wh (izin maskapai)", "Konfirmasi kapasitas sel; putuskan "
     "terminal BAT", "LGU"),
    ("AX-07", "Teknis", "Metode proteksi Ex d (flameproof enclosure)", "Proses", "28 Sep 2026",
     "Rekomendasi tim & dicentang di form; perlu diterima ExCB", "Tunggu tanggapan GTS/ExCB", "LGU"),
    ("AX-08", "Teknis", "Gambar enclosure & flame path terkendali (G-03)", "Belum", "7 Okt 2026",
     "Sketsa ukur MS-01…04 hanya ilustrasi", "Minta gambar dari mitra casing", "PT Galaksi"),
    ("AX-09", "Teknis", "Cable gland Ex d IIC (barrier gland) — ganti gland polimer prototipe", "Belum", "7 Okt 2026",
     "", "Pilih gland bersertifikat M20×1,5", "PT Galaksi"),
    ("AX-10", "Teknis", "Bukti inlet gas / mesh sensor (flame arrestor) untuk Ex d", "Belum", "7 Okt 2026", "",
     "Data pabrikan mesh / desain arrestor", "PT Galaksi · LGU"),
    ("AX-11", "Teknis", "Kelas suhu T4 — uji suhu permukaan", "Sebagian", "30 Sep 2026",
     "Ekstrapolasi +60 °C ≈ 114 °C vs batas efektif 130 °C; enclosure ≈ 79 °C; TC-3/4/5, mode alarm & TX "
     "belum", "Lengkapi pengukuran & metadata uji", "Lab IoT ITB"),
    ("AX-12", "Dokumen", "BOM Ex-critical & material non-logam (6.b/6.c)", "Sebagian", "7 Okt 2026",
     "BOM elektronik ada; data material & datasheet pemasok belum", "Minta data mitra casing", "PT Galaksi · LGU"),
    ("AX-13", "Dokumen", "Proses manufaktur (6.d)", "Sebagian", "7 Okt 2026", "Struktur d.1–d.13 ada; data pabrik "
     "belum", "Lengkapi dari PT Galaksi", "PT Galaksi"),
    ("AX-14", "Dokumen", "Nameplate & marking (6.h)", "Sebagian", "7 Okt 2026", "Artwork usulan MS-05 80×55 mm",
     "Final setelah sertifikat", "LGU · PT Galaksi"),
    ("AX-15", "Dokumen", "Sertifikat komponen Ex (6.i)", "Belum", "7 Okt 2026", "Gland, blanking plug, dll.",
     "Kumpulkan sertifikat", "PT Galaksi"),
    ("AX-16", "Organisasi", "ISO 9001 PT Galaksi", "Proses", "7 Okt 2026", "Audit dijadwalkan Okt 2026",
     "Konfirmasi hasil audit", "PT Galaksi"),
    ("AX-17", "Kontrak", "Termin 1 sertifikasi: uji internal disaksikan Pertamina + berita acara", "Belum",
     "30 Sep 2026", "Harus sebelum sampel diuji lab; penilaian oleh Pertamina", "Jadwalkan sesi witness", "LGU"),
    ("AX-18", "Terkait", "Uji IP66 (sheet Rencana IP)", "Belum", DATE, "", "Lihat IP-02 & IP-11", "LGU"),
    ("AX-19", "Terkait", "EMC (sheet Checklist EMC)", "Sebagian", DATE, "Pre-scan emisi selesai; imunitas belum",
     "Lihat EMC-06…EMC-18", "LGU"),
    ("AX-20", "Terkait", "RF / SDPPI (sheet Rencana RF)", "Belum", DATE, "", "Lihat RF-01…RF-03", "LGU"),
]


def build_cert():
    wb = Workbook()
    rs = wb.active
    rs.title = "Ringkasan"

    ws = wb.create_sheet("Checklist EMC")
    head = ["ID", "Kelompok", "Item", "Standar", "Status", "Hasil / catatan", "PIC", "Bukti", "Target", "Update"]
    title_block(ws, "Checklist EMC — Node GLD, Cluster Head, Gateway",
                f"Status per {DATE}, dari pre-scan EN 55032 Kelas A di Telkom Test House (15 Okt 2026). Detail "
                "temuan E-1…E-7 ada di Laporan Teknis Pengujian EMC (LGU/GLD/EMC/2026-001).", len(head))
    rows = [[a, b, c, d, e, f, g, h, "", ""] for a, b, c, d, e, f, g, h in EMC]
    e_end = write_table(ws, 4, head, rows, [9, 11, 40, 20, 11, 42, 16, 22, 12, 12], 5, ST)

    ws = wb.create_sheet("Rencana IP")
    head = ["ID", "Tahap", "Kegiatan", "Standar / klausul", "Metode / parameter", "Kriteria lulus", "Status", "PIC",
            "Target", "Catatan"]
    title_block(ws, "Rencana Uji IP (Ingress Protection) — Node GLD",
                "Target IP66 menurut IEC 60529, sejalan dengan marking Ex usulan. Urutan: keputusan desain → pra-uji "
                "internal → uji resmi di lab Ex (uji termal & impak lebih dulu, lalu IP). Kolom Target diisi tim.",
                len(head))
    rows = [[a, b, c, d, e, f, g, h, "", ""] for a, b, c, d, e, f, g, h in IP]
    i_end = write_table(ws, 4, head, rows, [8, 11, 36, 20, 46, 28, 11, 18, 12, 24], 7, ST)

    ws = wb.create_sheet("Rencana RF")
    title_block(ws, "Rencana Uji & Sertifikasi RF (LoRa 920–923 MHz)",
                "Sertifikasi radio utama: SDPPI (Indonesia). Angka batas regulasi wajib diverifikasi ke peraturan "
                "terbaru sebelum dipakai. RF-03/RF-04 berpotensi mengubah setelan daya pancar & jangkauan.",
                len(head))
    rows = [[a, b, c, d, e, f, g, h, "", ""] for a, b, c, d, e, f, g, h in RF]
    r_end = write_table(ws, 4, head, rows, [8, 11, 38, 22, 48, 24, 11, 16, 12, 24], 7, ST)

    ws = wb.create_sheet("Monitor ATEX")
    head = ["ID", "Kategori", "Item", "Status", "Update terakhir", "Keterangan", "Langkah berikut", "PIC",
            "Target", "Bukti / dokumen"]
    title_block(ws, "Monitor Sertifikasi ATEX/IECEx — Node GLD",
                f"Status per {DATE}. GTS = agen sertifikasi (bukan badan sertifikasi). Perbarui kolom Status & "
                "Update setiap ada perkembangan, lalu catat di sheet Log ATEX.", len(head))
    rows = [[a, b, c, d, e, f, g, h, "", ""] for a, b, c, d, e, f, g, h in ATEX]
    a_end = write_table(ws, 4, head, rows, [8, 12, 46, 11, 14, 44, 30, 18, 12, 22], 4, ST)

    lg = wb.create_sheet("Log ATEX")
    head = ["Tanggal", "ID terkait", "Perkembangan", "Sumber", "Dicatat oleh"]
    title_block(lg, "Log perkembangan ATEX", "Tambahkan satu baris setiap ada perkembangan (umpan balik GTS, uji, "
                "dokumen baru).", len(head))
    log = [
        ["29 Sep 2026", "AX-01", "Form aplikasi & dossier dikirim ke GTS", "Berkas terkirim", "LGU"],
        ["30 Sep 2026", "AX-11", "Data suhu pertama: ekstrapolasi ≈ 114 °C vs batas efektif T4 130 °C", "Uji internal",
         "Lab IoT ITB"],
        ["7–8 Okt 2026", "AX-03", "Dokumen Teknis Sertifikasi diaudit vs Panduan Revisi; kontradiksi dibetulkan",
         "Dokumen rev 2.8–2.9", "LGU"],
        ["12 Okt 2026", "AX-05", "Surat pengantar sampel ke GTS & surat keterangan Bea Cukai disiapkan", "Surat", "LGU"],
        ["15 Okt 2026", "AX-19", "Pre-scan EMC GLD & CH di Telkom Test House", "Laporan EMC", "Lab IoT ITB"],
        ["28 Okt 2026", "—", "Checklist EMC, rencana IP & RF, dan monitor ATEX dibuat", "Workbook ini", "LGU"],
    ]
    write_table(lg, 4, head, log, [14, 12, 70, 22, 16])

    # ---- ringkasan
    rs["A1"] = "Ringkasan Sertifikasi GLD — EMC · IP · RF · ATEX"
    rs["A1"].font = Font(bold=True, size=14, color=NAVY)
    rs["A2"] = f"Status per {DATE}. Angka dihitung otomatis dari kolom Status di setiap sheet."
    rs["A2"].font = Font(italic=True, color="555B66")
    r = 4
    r = summary(rs, r, "Checklist EMC", "Checklist EMC", "E", 5, e_end, ST)
    r = summary(rs, r, "Rencana IP", "Rencana IP", "G", 5, i_end, ST)
    r = summary(rs, r, "Rencana RF", "Rencana RF", "G", 5, r_end, ST)
    r = summary(rs, r, "Monitor ATEX", "Monitor ATEX", "D", 5, a_end, ST)
    rs.column_dimensions["A"].width = 22
    rs.column_dimensions["B"].width = 10
    rs.column_dimensions["C"].width = 10
    rs["E4"] = "Catatan prioritas"
    rs["E4"].font = Font(bold=True, size=11, color=NAVY)
    pr = ["1. EMC: margin GLD 0,51 dB @ 117,9 MHz — mitigasi & re-scan sebelum uji resmi (EMC-07).",
          "2. RF: EIRP konfigurasi saat ini (GLD ≈ 25 dBm, CH mesh ≈ 30 dBm) perlu dicek ke batas SDPPI (RF-02/03).",
          "3. IP: inlet gas + fan terbuka perlu keputusan desain agar IP66 tercapai (IP-02).",
          "4. ATEX: tunggu umpan balik GTS, kirim ulang dossier terkoreksi, gambar flame path & gland Ex d (AX-02…09).",
          "5. Termin 1 sertifikasi: sesi uji internal disaksikan Pertamina + berita acara (AX-17)."]
    for i, t in enumerate(pr):
        c = rs.cell(row=5 + i, column=5, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        rs.merge_cells(start_row=5 + i, start_column=5, end_row=5 + i, end_column=12)
        rs.row_dimensions[5 + i].height = 30
    for col in "EFGHIJKL":
        rs.column_dimensions[col].width = 12

    out = os.path.join(DELIV, "Checklist_Rencana_Sertifikasi_GLD_EMC_IP_RF_ATEX.xlsx")
    wb.save(out)
    print("written", out)


if __name__ == "__main__":
    build_block_diagram()
    build_cert()
