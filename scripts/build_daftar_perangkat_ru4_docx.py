# -*- coding: utf-8 -*-
"""Daftar Perangkat & Kebutuhan Pemasangan GLD — RU IV Cilacap (untuk tim PT Pertamina Patra Niaga).

Output: Deliverables/Daftar_Perangkat_dan_Kebutuhan_Pemasangan_GLD_RU-IV_Cilacap.docx
"""
import os

from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from build_persiapan_instalasi_corporate_docx import (
    make_doc, set_cell_shading, set_cell_border, add_field, set_cell_margins, set_table_borders,
    NAVY, GRAY, WHITE, HEAD_SHADE, ZEBRA_SHADE, OK_BG, OK_BD, WARN_BG, WARN_BD,
)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "Deliverables", "Daftar_Perangkat_dan_Kebutuhan_Pemasangan_GLD_RU-IV_Cilacap.docx")
DOC_NO = "LGU/GLD/INSTALASI-PERANGKAT/2026-001"
REVISION = "1.2"
DOC_DATE = "2 Oktober 2026"

doc, sec = make_doc()


def p(text="", size=10.2, bold=False, italic=False, color=None, space_after=7, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    if text:
        r = para.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return para


def bullet(text, size=10):
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(3)
    r = bp.add_run(text)
    r.font.size = Pt(size)


def box(text, kind="ok", label=None):
    bg, bd = {"ok": (OK_BG, OK_BD), "warn": (WARN_BG, WARN_BD)}[kind]
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, bg)
    set_cell_border(cell, bd, sz=6)
    set_cell_margins(cell, top=110, bottom=110, start=130, end=130)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    if label:
        r = para.add_run(label + " ")
        r.font.bold = True
        r.font.size = Pt(9.6)
    r = para.add_run(text)
    r.font.size = Pt(9.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def finish(tbl, keep=True):
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    for row in tbl.rows:
        rp = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); rp.append(cs)
    if keep:
        for row in tbl.rows[:-1]:
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.keep_with_next = True


def table(headers, rows, widths, font=9.0, bold_first=True, keep=True):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, h in enumerate(headers):
        c = tbl.rows[0].cells[i]
        set_cell_shading(c, "1B2A4A")
        c.paragraphs[0].paragraph_format.space_after = Pt(1)
        r = c.paragraphs[0].add_run(h)
        r.font.bold = True; r.font.size = Pt(8.6); r.font.color.rgb = WHITE
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            if ri % 2 == 1:
                set_cell_shading(c, ZEBRA_SHADE)
            lines = str(val).split("\n")
            for li, line in enumerate(lines):
                para = c.paragraphs[0] if li == 0 else c.add_paragraph()
                para.paragraph_format.space_after = Pt(1)
                r = para.add_run(line)
                r.font.size = Pt(font)
                if bold_first and ci == 0:
                    r.font.bold = True
    for row in tbl.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    set_table_borders(tbl)
    finish(tbl, keep=keep)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ---------------------------------------------------------------- cover
lt = doc.add_table(rows=1, cols=1)
lc = lt.rows[0].cells[0]
set_cell_shading(lc, "1B2A4A")
lc.paragraphs[0].paragraph_format.space_after = Pt(2)
r = lc.paragraphs[0].add_run("PT LAPI GANESHA UTAMA")
r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = WHITE
lp = lc.add_paragraph()
lp.paragraph_format.space_after = Pt(3)
r = lp.add_run("Bekerja sama dengan Lab IoT & Fisika Institut Teknologi Bandung")
r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xC7, 0xD2, 0xE0)
doc.add_paragraph().paragraph_format.space_after = Pt(10)

title = doc.add_heading(level=0)
title.paragraph_format.space_after = Pt(4)
tr = title.add_run("Daftar Perangkat & Kebutuhan Pemasangan GLD")
tr.font.size = Pt(20); tr.font.bold = True; tr.font.color.rgb = NAVY
p("Pilot RU IV Cilacap — perangkat yang akan dipasang dan hal yang perlu disiapkan di lokasi",
  size=12, color=GRAY, space_after=12)

mt = doc.add_table(rows=0, cols=2)
mt.style = "Table Grid"
for k, v in [
    ("Nomor dokumen", DOC_NO), ("Revisi", REVISION), ("Tanggal", DOC_DATE),
    ("Ditujukan kepada", "Tim PT Pertamina Patra Niaga — RU IV Cilacap"),
    ("Disiapkan oleh", "PT LAPI Ganesha Utama, bersama Lab IoT & Fisika Institut Teknologi Bandung"),
    ("Tujuan", "Bahan persiapan pemasangan perangkat GLD Tahap 2 di RU IV Cilacap"),
]:
    row = mt.add_row().cells
    set_cell_shading(row[0], HEAD_SHADE)
    row[0].paragraphs[0].paragraph_format.space_after = Pt(2)
    r0 = row[0].paragraphs[0].add_run(k)
    r0.font.bold = True; r0.font.size = Pt(9.3); r0.font.color.rgb = GRAY
    row[1].paragraphs[0].paragraph_format.space_after = Pt(2)
    r1 = row[1].paragraphs[0].add_run(v)
    r1.font.size = Pt(9.6)
    row[0].width = Inches(1.9); row[1].width = Inches(4.6)
doc.add_paragraph().paragraph_format.space_after = Pt(8)

# header / footer
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("Daftar Perangkat & Kebutuhan Pemasangan GLD — RU IV Cilacap")
hr.font.size = Pt(8); hr.font.color.rgb = GRAY; hr.font.italic = True
fp = sec.footer.paragraphs[0]
fp.add_run(f"{DOC_NO}  ·  Rev. {REVISION}  ·  {DOC_DATE}  ·  Halaman ")
add_field(fp, "PAGE", "1")
fp.add_run(" dari ")
add_field(fp, "NUMPAGES", "1")
for run in fp.runs:
    run.font.size = Pt(8); run.font.color.rgb = GRAY

# ---------------------------------------------------------------- 1
doc.add_heading("1. Ringkasan", level=1)
p("Sistem GLD terdiri dari sensor node GLD yang mendeteksi gas dan menjalankan model AI di perangkat, "
  "Cluster Head yang meneruskan data melalui jaringan LoRa, Gateway yang menjembatani jaringan LoRa ke "
  "server, dan PC server di lokasi yang menjalankan broker data dan dashboard operator.")
p("Arah pemasangan yang sudah disepakati: instalasi permanen di lokasi non-ATEX pada perimeter Sulfur "
  "Recovery Unit (SRU), dan gas test chamber yang dioperasikan di kantor kilang (non-area proses). Pemasangan "
  "fisik dilaksanakan oleh vendor yang ditunjuk Pertamina; LGU & ITB menyiapkan perangkat, desain pemasangan, "
  "supervisi QA, serta pekerjaan elektrikal spesifik GLD dan commissioning.")

# ---------------------------------------------------------------- 2
doc.add_heading("2. Daftar Perangkat yang Akan Dipasang", level=1)
table(["Perangkat", "Jumlah", "Spesifikasi utama", "Penempatan"], [
    ["GLD — Node Sensor", "3 unit + 1 cadangan",
     "8 sensor gas MQ + AI on-device (ESP32-S3)\n24 VDC, maks 8 W (≈0,33 A)\n200 × 90 × 290 mm; 2,378 kg\n"
     "Enclosure aluminium ADC12 + stainless steel, IP66\nAntena omni 3 dBi; alarm visual/audible terintegrasi",
     "Titik pantau di perimeter SRU (outdoor)"],
    ["Cluster Head (CH)", "2 unit",
     "Relay LoRa dua radio (star + mesh)\nCatu daya solar (2 panel) + baterai 18650; maks 0,73 W\n"
     "CH besar: enclosure aluminium silinder ± Ø76 × 104 mm (+ konektor ± 15 mm)\nAntena 3 dBi (star) dan 8 dBi (mesh)",
     "Outdoor, area aman, terpapar sinar matahari"],
    ["Gateway", "1 unit",
     "Jembatan LoRa mesh ke MQTT (ESP32-S3)\n5 V via adaptor; maks 0,73 W\nEnclosure metal (dimensi dikonfirmasi tim LGU)\n"
     "Antena omni 8 dBi; uplink Wi-Fi ke PC server",
     "Unit di dalam ruangan (safe area); antena di luar/atap"],
    ["PC server site", "1 unit",
     "Minimum 4 vCore, RAM 8 GB, SSD 100 GB (disarankan 8 vCore/16 GB/250 GB)\nUbuntu Server 22.04 LTS, Docker\n"
     "NIC Gigabit; NIC kedua bila dashboard diakses dari intranet kantor",
     "Ruangan dekat Gateway"],
    ["Router jaringan field", "1 unit", "Jaringan lokal Gateway ↔ PC server (independen dari jaringan kilang)",
     "Bersama PC server"],
    ["Gas test chamber portable", "1 set",
     "Chamber akrilik dengan pompa, solenoid valve, sensor pembanding; untuk pengambilan dataset dan "
     "pelatihan model AI di lokasi", "Kantor kilang (non-area proses)"],
], [1.35, 0.95, 2.95, 1.25], keep=False)
p("Jumlah GLD, Cluster Head, dan Gateway mengikuti konfigurasi RU IV pada proposal GLD Tahap 2; jumlah dan "
  "titik final disesuaikan dengan penetapan titik pasang bersama RU IV.", size=9, italic=True, color=GRAY)

# ---------------------------------------------------------------- 3
doc.add_heading("3. Yang Perlu Disiapkan di Lokasi (Pertamina RU IV & Vendor Instalasi)", level=1)
table(["Perangkat", "Kebutuhan di lokasi", "Pihak"], [
    ["GLD (per titik)",
     "Titik catu daya 24 VDC ≥1 A per unit, termasuk kabel, PSU/adaptor AC-DC, proteksi & titik isolasi (LOTO)\n"
     "Struktur pemasangan: pipa/handrail 2\" (DN50, OD 60,3 mm) di titik pasang\n"
     "Jalur kabel dari sumber daya ke titik pasang\nAkses kerja dan clearance di depan inlet gas sensor",
     "Pertamina RU IV"],
    ["Bracket GLD",
     "Fabrikasi pelat mounting 250 × 250 mm + U-bolt M10 (2\"/DN50) sesuai gambar CAD LGU, material "
     "anti-korosi, tanpa las/bor pada struktur existing\nPemasangan fisik unit",
     "Vendor instalasi"],
    ["Cluster Head",
     "Lokasi outdoor di area aman dengan paparan matahari untuk panel surya\nTiang tegak (pipa) sebagai "
     "dudukan unit CH, bracket 2 panel surya (diklem ke tiang), dan antena di puncak tiang — konfigurasi "
     "seperti Gambar 1", "Pertamina RU IV & vendor"],
    ["Gateway",
     "Ruang indoor di area aman + stopkontak untuk adaptor 5 V\nTitik antena di luar/atap dengan jalur kabel "
     "antena ke unit", "Pertamina RU IV"],
    ["PC server",
     "Unit PC fisik sesuai spesifikasi minimum (bagian 2), ditempatkan dekat Gateway, dengan catu daya\n"
     "Akses ke intranet kantor melalui NIC kedua bila dashboard ingin dibuka dari kantor", "Pertamina RU IV"],
    ["Gas test chamber", "Ruang/meja di kantor kilang (non-area proses) + stopkontak listrik", "Pertamina RU IV"],
    ["Umum",
     "Penetapan titik pasang & dokumen klasifikasi area\nPengesahan TRA/JSA, izin kerja, izin masuk personel & barang\n"
     "Penugasan vendor instalasi\nPIC operasi, HSE, control room, dan IT selama pemasangan", "Pertamina RU IV"],
], [1.25, 4.05, 1.2], keep=False)

doc.add_heading("Contoh konfigurasi pemasangan Cluster Head", level=2)
p("Cluster Head akan dipasang dengan konfigurasi seperti foto uji di bawah: unit CH dan antena omni berada "
  "pada satu tiang, dua panel surya dipasang pada bracket baja yang diklem (U-clamp) ke tiang dengan arah "
  "hadap berlawanan, dan kabel diturunkan sepanjang tiang. Bracket ini tidak memerlukan las maupun bor "
  "pada tiang.")
ph = doc.add_table(rows=2, cols=3)
ph.alignment = WD_TABLE_ALIGNMENT.CENTER
caps = ["(a) Bracket 2 panel surya, diklem ke tiang", "(b) Tiang + antena, uji di lapangan",
        "(c) Tiang + antena, uji di atap gedung"]
for i, fn in enumerate(["ch_bracket_solar_detail.jpg", "ch_mast_antena_lapangan.jpg", "ch_mast_antena_atap.jpg"]):
    c = ph.rows[0].cells[i]
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraphs[0].add_run().add_picture(os.path.join(REPO, "scripts", "assets", "ch_mount_photos", fn),
                                          width=Inches(2.05))
    cp = ph.rows[1].cells[i].paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(caps[i]); r.font.size = Pt(8.4); r.font.color.rgb = GRAY
p("Gambar 1. Konfigurasi pemasangan Cluster Head dengan bracket panel surya (uji LGU & ITB di Bandung).",
  size=9, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
p("Drum pada foto hanya dudukan sementara untuk pengujian. Di RU IV, tiang dipasang permanen (ditanam dan "
  "dicor, atau diklem ke struktur existing) sesuai dokumen Spesifikasi Material Instalasi; material tiang "
  "dan dudukan mengikuti ketentuan kilang.", size=9, italic=True, color=GRAY)

# ---------------------------------------------------------------- 4
doc.add_heading("4. Yang Disiapkan LGU & ITB", level=1)
for t in [
    "Unit GLD (4), Cluster Head (2), Gateway (1), dan router jaringan field — firmware terpasang dan diverifikasi sebelum keberangkatan.",
    "Gambar CAD pemasangan (pelat mounting U-bolt) sebagai acuan fabrikasi vendor.",
    "Instalasi software di PC server: broker MQTT, backend, dan dashboard (installer satu paket).",
    "Gas test chamber portable beserta perlengkapannya.",
    "Supervisi teknis/QA selama pemasangan, terminasi elektrikal spesifik GLD, energize, dan commissioning bersama HSE/control room.",
]:
    bullet(t)

# ---------------------------------------------------------------- 5
doc.add_heading("5. Urutan Pemasangan", level=1)
table(["Tahap", "Kegiatan", "Pihak utama"], [
    ["1", "Penetapan titik pasang, klasifikasi area, dan izin kerja", "Pertamina RU IV"],
    ["2", "Penyiapan catu daya 24 VDC, jalur kabel, ruang Gateway & PC server", "Pertamina RU IV"],
    ["3", "Fabrikasi dan pemasangan bracket", "Vendor instalasi (supervisi LGU)"],
    ["4", "Pemasangan GLD, Cluster Head, Gateway, antena", "Vendor instalasi (supervisi LGU)"],
    ["5", "Instalasi PC server & router, konfigurasi jaringan field", "LGU & ITB"],
    ["6", "Terminasi elektrikal GLD, energize, uji komunikasi & alarm, commissioning", "LGU & ITB bersama HSE/control room"],
], [0.6, 4.2, 1.7])

box("Detail lengkap tersedia pada dokumen Pembagian Persiapan Instalasi, Daftar Persiapan Vendor, dan Draf "
    "Permintaan Penyediaan Material Instalasi RU IV Cilacap yang telah dibagikan sebelumnya.", kind="ok",
    label="Rujukan.")

doc.save(OUT)
print("written", OUT)
