# -*- coding: utf-8 -*-
"""Daftar tindak lanjut revisi dossier ATEX/IECEx GLD V2.

Disusun dari "Panduan Revisi ATEX/IECEx" dibandingkan dengan dossier yang sudah
dikirim ke GTS (IECEx_ATEX_Certification_Information_Requirements_Rev25092026_formatted).
Fokus utama: daftar gambar teknik & foto yang perlu disiapkan, lalu perbaikan
kontradiksi dan item dokumen non-gambar.

Output: Deliverables/Daftar_Tindak_Lanjut_Revisi_ATEX_GLD.docx
"""
import os

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "Deliverables", "Daftar_Tindak_Lanjut_Revisi_ATEX_GLD.docx")

NAVY = RGBColor(0x1B, 0x33, 0x5F)
INK = RGBColor(0x26, 0x23, 0x21)
GRAY = RGBColor(0x55, 0x5B, 0x66)
HEAD_SHADE = "EAEFF9"
NOTE_SHADE = "F3F6FC"
WARN_SHADE = "FBEEE4"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10)
st.font.color.rgb = INK

sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width, sec.page_height = sec.page_height, sec.page_width
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(1.6))
sec.top_margin = Cm(1.5)
sec.bottom_margin = Cm(1.5)

# footer nomor halaman
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Daftar Tindak Lanjut Revisi Dossier ATEX/IECEx — GLD V2  ·  Halaman ")
r.font.size = Pt(8)
r.font.color.rgb = GRAY
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run()
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY
    if kind:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), kind)
        run._r.append(fc)
    else:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = text
        run._r.append(it)


def shade(cell, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


def p(text="", size=10, bold=False, italic=False, color=None, space_after=6, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color or INK
    return para


def h(text):
    return p(text, size=13, bold=True, color=NAVY, space_after=4)


def note_box(text, fill=NOTE_SHADE):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    shade(cell, fill)
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table(headers, rows, widths, small_cols=()):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for i, txt in enumerate(headers):
        shade(hdr.cells[i], HEAD_SHADE)
        hdr.cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
        rr = hdr.cells[i].paragraphs[0].add_run(txt)
        rr.font.bold = True
        rr.font.size = Pt(8.5)
        rr.font.color.rgb = NAVY
    for vals in rows:
        row = tbl.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for idx, val in enumerate(vals):
            cell = row.cells[idx]
            lines = val if isinstance(val, list) else [val]
            for li, line in enumerate(lines):
                para = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.space_before = Pt(2 if li == 0 else 0)
                run = para.add_run(("• " if isinstance(val, list) else "") + line)
                run.font.size = Pt(8.5 if idx in small_cols else 9)
                if idx == 0:
                    run.font.bold = True
                if idx in small_cols:
                    run.font.color.rgb = GRAY
    tbl.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl._tbl.tblPr.append(layout)
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    for row in tbl.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


# ---------------------------------------------------------------- judul
p("PT LAPI GANESHA UTAMA  ·  GAS LEAK DETECTION TAHAP 2", size=9, bold=True, color=GRAY, space_after=2)
p("Daftar Tindak Lanjut Revisi Dossier ATEX/IECEx — GLD V2", size=18, bold=True, color=NAVY, space_after=2)
p("Fokus: gambar teknik & foto yang perlu disiapkan, perbaikan kontradiksi, dan item dokumen yang masih kosong",
  size=10.5, color=GRAY, space_after=8)

table(
    ["Dasar penyusunan", "Keterangan"],
    [
        ["Acuan", "Panduan Revisi ATEX/IECEx (catatan reviewer atas draf checklist ExCB)"],
        ["Dokumen yang dibandingkan",
         "IECEx/ATEX Certification Information Requirements — GLD V2, Rev 25 Sep 2026 (formatted, 60 hlm), "
         "versi yang sudah dikirim ke GTS"],
        ["Tanggal", "7 Oktober 2026"],
        ["Sifat dokumen", "Internal — daftar kerja; bukan berkas yang dikirim ke GTS/ExCB"],
    ],
    widths=(2.2, 7.6),
)

note_box(
    "Cara pakai: Bagian A dan B adalah daftar gambar dan foto yang disiapkan tim. Kolom \"Sudah ada\" menunjukkan "
    "bahan yang sudah tersedia di repo dan tinggal dirapikan; kolom \"Cek\" dicentang setelah gambar/foto selesai. "
    "Bagian C wajib dibereskan sebelum dossier dikirim ulang, karena berupa pernyataan yang saling bertentangan "
    "di dokumen yang sama. Bagian D adalah data non-gambar yang masih kosong."
)

# ---------------------------------------------------------------- ketentuan umum gambar
h("Ketentuan umum untuk setiap gambar teknik")
table(
    ["Wajib ada di setiap lembar", "Keterangan"],
    [
        ["Title block", "Nomor gambar, revisi, tanggal, judul, nomor lembar (mis. 2/5), skala, nama pembuat & pemeriksa"],
        ["Satuan & toleransi", "Satuan mm; toleransi umum (mis. ISO 2768-m) + toleransi khusus pada dimensi Ex-critical"],
        ["Material", "Grade material dan perlakuan permukaan tertulis di gambar (bukan hanya di tabel teks)"],
        ["Konsistensi", "Gambar menjadi sumber yang menentukan; nilai di tabel teks dossier harus sama dengan gambar"],
        ["Konfigurasi", "Hanya konfigurasi produksi 24 VDC; elemen yang sudah tidak dipakai (battery case, posisi "
                        "antena lama) dihapus atau ditandai \"superseded\""],
        ["Penomoran usulan", "GLD-M-xxx untuk mekanis, GLD-E-xxx untuk elektrik (boleh diganti sistem Galaksi)"],
    ],
    widths=(2.2, 7.6),
)

# ---------------------------------------------------------------- A. gambar
doc.add_page_break()
h("A. Daftar gambar teknik yang perlu disiapkan (butir 6.a)")
p("Urutan prioritas: G-03 (enclosure/flame path) dan G-07 (skematik) paling menentukan penilaian Ex d.",
  size=9, color=GRAY, space_after=4)

GAMBAR = [
    ["G-01", "Assembly drawing (gambar rakitan keseluruhan)",
     ["Body & cover casing", "Front SS mesh + plat penahan", "DC fan", "8 sensor MQ", "Motherboard PCB",
      "Modul alarm LED/buzzer", "Antena", "Cable entry", "Grounding stud", "Bracket/U-bolt plate",
      "Gasket/seal", "Semua sambungan berulir antar bagian", "Balon nomor + daftar part"],
     "Exploded view di dossier (hlm. 25–26, 33) — belum ada title block, dimensi, daftar part", "Tinggi", ""],
    ["G-02", "Section / internal arrangement drawing (penampang jalur gas)",
     ["Alur: udara luar → front mesh → fan → sensing chamber → mesh MQ → elemen sensor",
      "Hubungan sensing chamber dengan PCB",
      "Tegas: PCB satu volume dengan sensor, ATAU ada partisi (\"PCB plate\") — sebutkan materialnya",
      "Panah arah aliran udara"],
     "Belum ada", "Tinggi", ""],
    ["G-03", "Enclosure structure drawing (parameter Ex d)",
     ["Material + grade (ADC12?) & perlakuan permukaan", "Tebal dinding & tebal cover",
      "Dimensi body", "Tipe ulir, pitch, panjang engagement (jumlah ulir penuh)",
      "Panjang joint / flame path & celah maksimum (gap) + toleransi", "Volume bebas internal (cm³)",
      "Lokasi gasket & permukaan seal", "Dudukan mesh", "Ulir cable entry", "Penetrasi antena & modul alarm"],
     "CAD \"ATEX CASING v2\" (.step) & \"GLD ATEX CASE v3\" (PDF) — baru envelope eksterior; parameter flame path "
     "belum ada (perlu dari mitra casing)", "Kritis", ""],
    ["G-04", "Front mesh assembly drawing",
     ["Material mesh (SS304/316)", "Diameter luar & tebal", "Ukuran mesh / pori", "Jumlah lapis",
      "Ring/plat penahan & cara fiksasi", "Jarak mesh ke fan", "Orientasi", "Gasket bila ada",
      "Part number / sertifikat bila ada"],
     "Foto \"Stainless Wire-Mesh Plate\" (hlm. 20) — tanpa dimensi", "Tinggi", ""],
    ["G-05", "MQ sensor arrangement drawing",
     ["Posisi ke-8 MQ + tipe tiap posisi (MQ-2, 3B, 4, 5, 6, 7B, 8, 135)", "Jarak antar sensor",
      "Orientasi", "Posisi relatif terhadap fan", "Tinggi sensor terhadap front cover", "Posisi mesh bawaan MQ"],
     "Layout PCB sensor board (EasyEDA) — perlu dijadikan gambar berlabel tipe MQ", "Sedang", ""],
    ["G-06", "DC fan mounting drawing",
     ["Model/merk fan", "Dimensi", "Sekrup pemasangan", "Jarak ke front mesh", "Arah aliran (panah)",
      "Clearance ke casing", "Konektor", "Cara fiksasi"],
     "Belum ada; model fan belum tercatat sama sekali", "Tinggi", ""],
    ["G-07", "Skematik listrik (bukan diagram blok)",
     ["Input 24 VDC + proteksi (PTC, TVS SMBJ33A, P-MOSFET)", "Rail 5 V & 3,3 V", "Suplai heater MQ",
      "Rangkaian analog MQ & ADC ADS1256", "ESP32-S3", "LoRa E22", "Driver fan", "Driver LED/buzzer",
      "RS-485/Modbus", "Grounding",
      "Tiap lembar: no. dokumen, revisi, tanggal, judul, nomor lembar"],
     "Ada: 9 lembar skematik EasyEDA (scripts/assets/cert_doc_schematics/01–09) — perlu title block & revisi", "Kritis", ""],
    ["G-08", "PCB layout",
     ["Top & bottom layer", "Penempatan komponen", "Dimensi board & lubang mounting", "Konektor",
      "Area daya / arus besar", "Interface sensor", "Grounding",
      "Satu set per board: Main Board, Sensor Board (+ Alarm/Terminal board bila terpisah)"],
     "Ada: PCB layout copper (10-pcb-layout.png) + PDF PCB EasyEDA — perlu per-board & per-layer", "Sedang", ""],
    ["G-09", "Terminal & input 24 VDC drawing",
     ["Penandaan terminal & pin (+24 V, 0 V, PE)", "Model konektor", "Ukuran kabel",
      "Cable gland/entry yang dipakai", "Polaritas", "Torsi sekrup terminal"],
     "Ilustrasi \"24VDC POWER IN\" (hlm. 37, 44) — masih tertulis \"BATTERY OR 24VDC\"", "Tinggi", ""],
    ["G-10", "Grounding / bonding drawing",
     ["Posisi grounding stud", "Ukuran baut", "Urutan washer–lug–mur", "Material hardware",
      "Bonding ke enclosure & antar bagian logam", "Simbol PE", "Torsi", "Batas resistansi kontinuitas"],
     "Hanya foto stud (hlm. 45)", "Tinggi", ""],
    ["G-11", "Cable-entry drawing",
     ["Lokasi entry", "Bentuk & ukuran ulir (PG13.5 atau M20×1,5 — harus diputuskan)", "Adapter bila ada",
      "Gland + seal", "Panjang engagement", "Blanking plug bila ada entry tak terpakai"],
     "Referensi BP18-1Z (M20×1,5) — tidak cocok dengan tabel dossier (PG13.5)", "Kritis", ""],
    ["G-12", "Antenna interface drawing",
     ["SMA bulkhead", "Lubang pemasangan", "Mur/washer", "Penetrasi enclosure", "Kabel internal coax/U.FL",
      "Seal", "Orientasi antena (posisi atas base enclosure — versi terbaru)"],
     "Belum ada; foto antena lama sudah tidak sesuai", "Sedang", ""],
    ["G-13", "Alarm module drawing",
     ["Housing", "Gasket", "Cara fiksasi ke enclosure", "Wiring", "Interface ke enclosure utama", "Material"],
     "Foto LED & buzzer + rubber gasket (hlm. 15–16, 36)", "Sedang", ""],
]
table(["No", "Gambar", "Isi wajib", "Sudah ada", "Prioritas", "Cek"],
      GAMBAR, widths=(0.5, 1.7, 3.9, 2.4, 0.7, 0.5), small_cols=(3,))

# ---------------------------------------------------------------- B. foto
doc.add_page_break()
h("B. Daftar foto produk yang perlu disiapkan (butir 2.4)")
p("Format caption: \"Figure 4-n. GLD V2 – <judul>.\" lalu satu kalimat yang menyebut apa yang terlihat. "
  "Jangan menampilkan marking Ex final seolah sudah tersertifikasi, dan jangan menyebut mesh MQ \"flame-arresting\".",
  size=9, color=GRAY, space_after=4)

FOTO = [
    ["F-01", "Tampak luar keseluruhan — depan, belakang, kiri, kanan, atas, bawah",
     "Bentuk enclosure, posisi antena, cable entry, modul alarm, bracket, grounding stud, inlet gas",
     "Sebagian (depan/kanan/belakang, hlm. 11) — lengkapi 6 sisi unit rakitan", ""],
    ["F-02", "Area nameplate", "Lokasi pelat; bila belum final beri keterangan \"marking draft\"", "Belum", ""],
    ["F-03", "Close-up area sensing depan", "Front SS mesh, cover, posisi fan di belakangnya",
     "Ada foto cover + mesh (hlm. 12)", ""],
    ["F-04", "Bagian sensing internal (cover dibuka)", "8 MQ, posisi terhadap fan, ruang sensing chamber",
     "Ada (hlm. 22) — perlu sudut yang menunjukkan fan", ""],
    ["F-05", "Close-up sensor MQ (1–2 buah)", "Mesh stainless bawaan sensor", "Ada (hlm. 17)", ""],
    ["F-06", "Motherboard — sisi komponen & sisi bawah", "Konektor, bagian daya, MCU, LoRa, ADC, interface sensor",
     "Sebagian (di dalam casing, hlm. 21)", ""],
    ["F-07", "Input daya & terminal (close-up)", "Terminal 24 VDC, gland/konektor, proteksi input, jalur grounding",
     "Belum (foto close-up)", ""],
    ["F-08", "Grounding / bonding point", "Stud eksternal & sambungannya ke enclosure", "Belum (close-up)", ""],
    ["F-09", "Semua penetrasi enclosure", "Cable gland, SMA antena, interface alarm, blanking plug",
     "Belum", ""],
    ["F-10", "Modul alarm", "LED/buzzer, sambungan mekanis, gasket, koneksi ke enclosure",
     "Ada (hlm. 15–16)", ""],
    ["F-11", "Enclosure terbuka (susunan internal utuh)",
     "Apakah fan, MQ, PCB, terminal dalam satu volume atau terpisah", "Belum", ""],
    ["F-12", "Urutan perakitan (foto nyata)", "Cover, mesh, fan, PCB, sensor array, gasket, enclosure",
     "Ada versi ilustrasi (hlm. 33–41) — foto nyata belum", ""],
    ["F-13", "Antena posisi terbaru", "Antena di atas base enclosure (ganti foto posisi samping yang lama)",
     "Belum — foto lama superseded", ""],
]
table(["No", "Foto", "Harus terlihat", "Status di dossier", "Cek"],
      FOTO, widths=(0.5, 2.6, 3.6, 2.5, 0.5), small_cols=(3,))

# ---------------------------------------------------------------- C. kontradiksi
doc.add_page_break()
h("C. Kontradiksi yang wajib dibetulkan sebelum dossier dikirim ulang")
KONTRA = [
    ["C-1", "Marking 6.h & 6.i memuat II 2 G D, Ex tb IIIC T135°C Db, Tamb −40…+85 °C (disalin dari flame detector "
            "pembanding)",
     "Ganti: II 2G Ex db IIC T4 Gb, Tamb −20 °C ≤ Ta ≤ +60 °C (usulan); hapus kategori debu D & Ex tb",
     "hlm. 57–58"],
    ["C-2", "Tamb berbeda di 3 tempat: −20/+85 (§2, §3.9, g.b), −20/+60 (§5), Ta max +85 (6.f.2)",
     "Seragamkan −20…+60 °C. Pada +85 °C kondisi fan mati ≈139 °C → melewati batas T4", "hlm. 6, 9, 22, 52–53"],
    ["C-3", "Cable entry PG13.5 di tabel vs referensi M20×1,5", "Putuskan satu, sesuaikan dengan gambar G-11",
     "§2, §3.9"],
    ["C-4", "Protection concept \"To be confirmed\" / Ex [db/ib], sementara form GTS dicentang Ex d",
     "Tulis Ex d (flameproof enclosure, usulan pemohon)", "§2"],
    ["C-5", "e.3.2: PCB dipisah \"PCB plate\" vs IS-03: PCB & MQ satu volume",
     "Pastikan konstruksi nyata, tuangkan di G-02, samakan e.2 & e.3", "e.2, e.3.2"],
    ["C-6", "Battery case & \"POWER IN BATTERY OR 24VDC\" masih tampil di gambar",
     "Putuskan baterai dikeluarkan dari produk yang disertifikasi; hapus dari gambar & BOM (IS-17 ditutup)",
     "hlm. 35, 44"],
    ["C-7", "Material enclosure ditulis umum \"aluminum alloy + stainless steel\"",
     "Tulis grade (ADC12 die-cast) → proses die casting BERLAKU di 6.d, bukan N.A.", "§2, 6.d"],
    ["C-8", "Daftar alat uji sampel mencantumkan \"alarm load / relay simulator\"",
     "Cek apakah GLD punya output relay; bila tidak, hapus atau ganti", "Bagian 3"],
]
table(["No", "Masalah", "Tindakan", "Lokasi"], KONTRA, widths=(0.5, 4.0, 4.2, 1.1), small_cols=(3,))

# ---------------------------------------------------------------- D. non-gambar
h("D. Item dokumen non-gambar yang masih kosong / belum lengkap")
NONG = [
    ["6.b", "BOM Ex-critical mekanis EX-01…EX-16",
     ["Enclosure, cover, front mesh, gasket/O-ring, cable gland, blanking plug, SMA bulkhead, antena, "
      "grounding stud, fastener, insulator internal, laminasi PCB (Tg/CTI/UL94), potting/adhesive",
      "Model DC fan (V/I, suhu maks, sertifikasi) — belum ada sama sekali",
      "Buzzer belum masuk BOM; MPN konektor VIN kosong; konfirmasi AO4407 sebagai driver heater MQ",
      "Kolom status seragam: Final / Open / Pending certificate / TBC / N.A."],
     "Mitra casing (mekanis) · LGU (elektrik)", ""],
    ["6.c", "Tabel material non-logam MAT-01…MAT-08 + datasheet supplier",
     ["Gasket cover, seal gland, FR-4, housing konektor, isolasi kabel, housing antena, impeller fan, adhesive",
      "Properti: rentang suhu, ageing, flame rating/UL94, CTI, Tg, tahan kimia/UV, antistatik"],
     "Mitra casing · LGU", ""],
    ["6.d", "Uraian proses manufaktur (kosong total di dossier)",
     ["Machining & inspeksi dimensi; surface treatment; pemasangan mesh, fan, MQ (cegah tipe tertukar)",
      "PCB assembly/AOI; cable entry; grounding (uji kontinuitas); gasket; penutupan akhir (torsi, engagement)",
      "Kontrol versi firmware & model AI; routine test akhir; tabel proses N.A. (las, potting)"],
     "Mitra casing · LGU", ""],
    ["6.e", "Penutupan IS-01…IS-17 (semua masih Open)",
     ["Dokumentasi pabrikan mesh MQ & front mesh (tanpa klaim flame-arresting bila tak ada sertifikat)",
      "Data fan; daya RF maks (22 dBm + antena 3 dBi); risiko elektrostatis non-logam; kontinuitas bonding"],
     "LGU · mitra casing", ""],
    ["6.f", "Hasil uji suhu (tabel 6.f.11)",
     ["Masukkan data 8 thermocouple yang sudah ada (≈114 °C ekstrapolasi +60 °C; enclosure ≈79 °C)",
      "Ukur TC-3/4/5 (DC/DC, inductor, MOSFET); kondisi alarm aktif & LoRa TX maks; mesh tersumbat; fault elektrik",
      "Catat metadata: Ta uji, tegangan, arus, mode, status fan/alarm/RF, waktu stabil, kalibrasi termokopel"],
     "LGU", ""],
    ["6.g", "Nilai TBD di instruksi pemakaian",
     ["Toleransi input 24 VDC, arus inrush, rating fuse/MCB upstream",
      "Torsi gland, terminal, grounding, SMA; clearance minimum depan inlet",
      "Batas kontinuitas grounding; waktu warm-up sensor; suhu/kelembapan simpan; cara membersihkan mesh"],
     "LGU · mitra casing", ""],
    ["6.i", "Sertifikat komponen Ex",
     ["Enclosure (bila memakai enclosure bersertifikat)", "Cable gland — Ex d barrier gland untuk IIC",
      "Blanking plug", "Fan / modul alarm bila diklaim Ex"],
     "Mitra casing", ""],
    ["1.3–1.5", "Basic information",
     ["Bagan organisasi sampai fungsi QA & produksi (bukan hanya direksi/komisaris)",
      "ISO 9001: masih Open (audit Okt 2026)", "Nama & alamat penyedia perakitan PCB"],
     "PT Galaksi · LGU", ""],
]
table(["Butir", "Item", "Yang perlu dilengkapi", "Sumber / PIC", "Cek"],
      NONG, widths=(0.6, 1.9, 5.1, 1.7, 0.5), small_cols=(3,))

# ---------------------------------------------------------------- urutan
h("Urutan kerja yang disarankan")
table(
    ["Langkah", "Kegiatan"],
    [
        ["1", "Bereskan C-1…C-4 (marking, Tamb, cable entry, Ex d) — cepat & paling berisiko bila dibaca GTS"],
        ["2", "Putuskan C-5 & C-6 (partisi PCB, status baterai) — menentukan isi G-02, BOM, dan IS-03/IS-17"],
        ["3", "Minta paket gambar Ex d + data material ke mitra casing (G-03, G-04, G-11, 6.b mekanis, 6.c, 6.d, 6.i)"],
        ["4", "Siapkan gambar yang bisa dibuat internal: G-01, G-02, G-05, G-06, G-09, G-10, G-12, G-13"],
        ["5", "Rapikan G-07 & G-08 dari file EasyEDA yang sudah ada (title block, nomor, revisi)"],
        ["6", "Ambil foto F-01…F-13 sekaligus dalam satu sesi (unit rakitan, lalu dibongkar bertahap)"],
        ["7", "Lengkapi uji suhu (6.f) lalu isi tabel hasil"],
    ],
    widths=(0.8, 9.0),
)

p("Dokumen kerja internal. Nilai yang belum ada jangan diisi perkiraan — tandai Open/TBC sampai ada bukti.",
  size=8.5, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print("written", OUT)
