# -*- coding: utf-8 -*-
"""Formulir isian data dasar sertifikasi ATEX/IECEx (Bagian 1 checklist ExCB).

Dokumen internal berbahasa Indonesia untuk diisi fungsi legal/administrasi &
mutu LGU. Hasil isiannya nanti ditranskrip ke Bagian 1 dokumen submission
berbahasa Inggris (`Dokumen_Teknis_Sertifikasi_GLD_IECEx_ATEX`).

Output: Deliverables/Formulir_Isian_Data_Dasar_Sertifikasi_ATEX_GLD.docx
"""
import os

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "Deliverables", "Formulir_Isian_Data_Dasar_Sertifikasi_ATEX_GLD.docx")

NAVY = RGBColor(0x1B, 0x33, 0x5F)
BLUE = RGBColor(0x2B, 0x5F, 0xCB)
INK = RGBColor(0x26, 0x23, 0x21)
GRAY = RGBColor(0x55, 0x5B, 0x66)
HEAD_SHADE = "EAEFF9"
FILL_SHADE = "FFFFFF"
NOTE_SHADE = "F3F6FC"
LINE_GRAY = "C9D2E3"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.font.color.rgb = INK

sec = doc.sections[0]
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.8)


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def p(text="", size=10.5, bold=False, italic=False, color=None, space_after=8, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color or INK
    return para


def note_box(text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, NOTE_SHADE)
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = GRAY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def form_table(rows, widths=(0.4, 2.5, 2.3, 2.3)):
    """rows: (no, 'Data yang diminta', 'Keterangan/format') — kolom isian dikosongkan."""
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for i, h in enumerate(["NO", "DATA YANG DIMINTA", "ISIAN", "KETERANGAN / FORMAT"]):
        set_cell_shading(hdr.cells[i], HEAD_SHADE)
        hdr.cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
        r = hdr.cells[i].paragraphs[0].add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = NAVY
    for no, label, hint in rows:
        row = tbl.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for idx, val in enumerate([no, label, "", hint]):
            cells[idx].paragraphs[0].paragraph_format.space_after = Pt(2)
            run = cells[idx].paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)
            if idx == 1:
                run.font.bold = True
            if idx == 3:
                run.font.color.rgb = GRAY
                run.font.size = Pt(9)
        set_cell_shading(cells[2], FILL_SHADE)
        # beri ruang tulis pada kolom isian
        cells[2].paragraphs[0].paragraph_format.space_before = Pt(6)
        cells[2].paragraphs[0].paragraph_format.space_after = Pt(6)
    tbl.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl._tbl.tblPr.append(layout)
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    for row in tbl.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


# ---------------------------------------------------------------- kop
lt = doc.add_paragraph()
lt.paragraph_format.space_after = Pt(0)
r = lt.add_run("PT LAPI GANESHA UTAMA")
r.font.size = Pt(13)
r.font.bold = True
r.font.color.rgb = NAVY
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(14)
r = sub.add_run("Bekerja sama dengan Institut Teknologi Bandung — Lab IoT & Lab Fisika")
r.font.size = Pt(9)
r.font.color.rgb = GRAY

p("FORMULIR ISIAN — INTERNAL", size=8.5, bold=True, color=BLUE, space_after=2)
p("Data Dasar Sertifikasi ATEX / IECEx (Bagian 1)", size=18, bold=True, color=NAVY, space_after=4)
p("Gas Leak Detector (GLD) V2 — Node Sensor", size=11, color=GRAY, space_after=12)

meta = doc.add_table(rows=0, cols=2)
meta.style = "Table Grid"
for k, v in [
    ("Dokumen terkait", "Dokumen Teknis Sertifikasi GLD IECEx/ATEX — No. LGU/GLD/IECEX-TDF/2026-001, Rev. 0.2"),
    ("Dasar", "IECEx/ATEX Certification Information Requirements (ExCB), Bagian 1: Basic Information"),
    ("Tujuan", "Mengumpulkan data legalitas, organisasi, fasilitas produksi, dan sistem mutu yang belum tersedia"),
    ("Diisi oleh", "Fungsi Legal/Administrasi dan fungsi Mutu — dibantu mitra casing untuk butir 4"),
    ("Batas waktu", "……………………………………"),
]:
    row = meta.add_row().cells
    set_cell_shading(row[0], HEAD_SHADE)
    for idx, val in enumerate([k, v]):
        row[idx].paragraphs[0].paragraph_format.space_after = Pt(2)
        run = row[idx].paragraphs[0].add_run(val)
        run.font.size = Pt(9.5)
        run.font.bold = (idx == 0)
        if idx == 0:
            run.font.color.rgb = NAVY
    row[0].width = Inches(1.5)
    row[1].width = Inches(5.0)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

note_box(
    "Cara pakai: isi kolom ISIAN, lalu lampirkan berkas pendukung sesuai kolom keterangan. Isian akan "
    "ditranskrip ke Bagian 1 dokumen submission berbahasa Inggris yang dikirim ke lembaga sertifikasi (ExCB). "
    "Kosongkan bila memang belum ada — jangan diisi perkiraan, karena data ini menjadi bagian dari berkas "
    "resmi pendaftaran."
)

# ---------------------------------------------------------------- 1
p("1. Formulir aplikasi ExCB", size=12.5, bold=True, color=NAVY, space_after=4)
p("Template formulir diterbitkan oleh ExCB dan belum diterima. Data teknis untuk mengisinya sudah tersedia di "
  "dokumen submission; yang perlu dipastikan hanya butir-butir berikut.", size=9.5, color=GRAY, space_after=6)
form_table([
    ("1.1", "Nama & kontak ExCB yang dituju", "Lembaga sertifikasi yang akan menerima pengajuan"),
    ("1.2", "Template formulir aplikasi sudah diminta?", "Tanggal permintaan dan kanal (email/PIC)"),
    ("1.3", "Jalur sertifikasi yang diminta", "IECEx CoC, ATEX EU-type examination, atau keduanya"),
    ("1.4", "Negara/pasar tujuan", "Menentukan skema & penandaan yang berlaku"),
])

# ---------------------------------------------------------------- 2
p("2. Legalitas perusahaan", size=12.5, bold=True, color=NAVY, space_after=4)
form_table([
    ("2.1", "Nama badan hukum lengkap", "Sesuai akta — sudah diketahui: PT LAPI Ganesha Utama (mohon konfirmasi penulisan resmi)"),
    ("2.2", "Bentuk badan hukum & status kepemilikan", "Sesuai akta pendirian"),
    ("2.3", "Nomor Induk Berusaha (NIB)", "Lampirkan salinan"),
    ("2.4", "NPWP", "Lampirkan salinan"),
    ("2.5", "Akta pendirian & perubahan terakhir", "Nomor, tanggal, notaris, dan pengesahan kementerian"),
    ("2.6", "Alamat domisili terdaftar", "Sesuai yang tercantum pada NIB"),
    ("2.7", "Klasifikasi bidang usaha (KBLI) yang relevan", "Harus konsisten dengan kegiatan manufaktur perangkat"),
    ("2.8", "Terjemahan bahasa Inggris berkas legal", "Tanyakan ke ExCB apakah perlu terjemahan tersumpah"),
])

# ---------------------------------------------------------------- 3
p("3. Organisasi & kontak", size=12.5, bold=True, color=NAVY, space_after=4)
form_table([
    ("3.1", "Bagan organisasi", "Harus memperlihatkan fungsi desain, produksi, dan mutu beserta garis pelaporan"),
    ("3.2", "Penandatangan aplikasi", "Nama, jabatan, dan dasar kewenangan"),
    ("3.3", "PIC proyek sertifikasi", "Nama, jabatan, email, telepon — satu pintu korespondensi ExCB"),
    ("3.4", "PIC teknis produk", "Yang menjawab pertanyaan teknis atas technical file"),
    ("3.5", "PIC mutu", "Pendamping asesmen sistem mutu (butir 5)"),
    ("3.6", "PIC mitra ITB", "Lab IoT / Lab Fisika ITB — nama & kontak"),
    ("3.7", "Alamat korespondensi & bahasa kerja", "Bahasa Inggris diasumsikan kecuali ditentukan lain"),
])

# ---------------------------------------------------------------- 4
p("4. Fasilitas produksi", size=12.5, bold=True, color=NAVY, space_after=4)
p("Butir 4.3–4.5 perlu diisi bersama mitra casing eksternal; butir ini juga menjadi prasyarat datasheet "
  "material dan deskripsi proses manufaktur di dokumen submission.", size=9.5, color=GRAY, space_after=6)
form_table([
    ("4.1", "Alamat lokasi perakitan & pengujian akhir", "Saat ini tahap prototipe di laboratorium — sebutkan alamat lab"),
    ("4.2", "Penyedia fabrikasi & assembly PCB", "Nama perusahaan + alamat; perlu konfirmasi tertulis"),
    ("4.3", "Mitra casing/enclosure: identitas & alamat pabrik", "Diisi oleh mitra"),
    ("4.4", "Mitra casing: kemampuan proses", "Permesinan, perlakuan permukaan, pengelasan, potting — bila ada"),
    ("4.5", "Profil fasilitas produksi", "Luas area, peralatan utama, jumlah personel, kapasitas per bulan"),
    ("4.6", "Pemeriksaan barang masuk & QC akhir", "Bentuk pemeriksaan dan pencatatannya"),
    ("4.7", "Rencana lokasi produksi serial", "Bila berbeda dari lokasi prototipe"),
])

# ---------------------------------------------------------------- 5
p("5. Sistem mutu", size=12.5, bold=True, color=NAVY, space_after=4)
form_table([
    ("5.1", "Apakah perusahaan memiliki sertifikat ISO 9001?", "Ya / tidak — bila ya: nomor, ruang lingkup, lembaga penerbit, masa berlaku"),
    ("5.2", "Manual mutu", "Lampirkan bila ada"),
    ("5.3", "Daftar dokumen prosedur", "Cukup daftar/indeks pada tahap ini"),
    ("5.4", "Pengendalian gambar & perubahan desain", "Bagaimana revisi gambar dikendalikan di produksi"),
    ("5.5", "Format catatan uji rutin produksi", "Untuk uji rutin yang diwajibkan setelah metode proteksi ditetapkan"),
])
note_box(
    "Catatan: baik IECEx maupun ATEX mensyaratkan bukti sistem mutu yang mencakup produksi Ex sebelum "
    "sertifikat terbit — melalui Quality Assessment Report (IECEx) atau modul penjaminan mutu produksi/"
    "verifikasi produk (ATEX) — terlepas dari ada tidaknya ISO 9001. Sertifikat ISO 9001 umumnya memperpendek "
    "asesmen tersebut, bukan menggantikannya. Jalur yang berlaku dikonfirmasi ke ExCB."
)

# ---------------------------------------------------------------- lampiran
doc.add_page_break()
p("Daftar lampiran", size=12.5, bold=True, color=NAVY, space_after=6)
att = doc.add_table(rows=1, cols=3)
att.style = "Table Grid"
att.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["NO", "BERKAS", "STATUS (ADA / BELUM ADA)"]):
    set_cell_shading(att.rows[0].cells[i], HEAD_SHADE)
    att.rows[0].cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
    r = att.rows[0].cells[i].paragraphs[0].add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = NAVY
for i, nama in enumerate([
    "Salinan NIB",
    "Salinan NPWP",
    "Akta pendirian dan perubahan terakhir",
    "Bagan organisasi",
    "Sertifikat ISO 9001 (bila ada)",
    "Manual mutu (bila ada)",
    "Daftar dokumen prosedur (bila ada)",
    "Profil fasilitas produksi / company profile",
    "Surat konfirmasi penyedia fabrikasi PCB",
    "Data mitra casing (identitas, alamat, kemampuan proses)",
], 1):
    row = att.add_row()
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    cells = row.cells
    for idx, val in enumerate([str(i), nama, ""]):
        cells[idx].paragraphs[0].paragraph_format.space_after = Pt(4)
        cells[idx].paragraphs[0].paragraph_format.space_before = Pt(4)
        run = cells[idx].paragraphs[0].add_run(val)
        run.font.size = Pt(9.5)
    cells[0].width = Inches(0.4)
    cells[1].width = Inches(4.4)
    cells[2].width = Inches(2.2)
doc.add_paragraph().paragraph_format.space_after = Pt(16)

p("Pengesahan isian", size=12.5, bold=True, color=NAVY, space_after=8)
sig = doc.add_table(rows=2, cols=2)
sig.style = "Table Grid"
for i, (peran, ket) in enumerate([
    ("Diisi oleh", "Nama, jabatan, tanggal, tanda tangan"),
    ("Diperiksa oleh", "Nama, jabatan, tanggal, tanda tangan"),
]):
    set_cell_shading(sig.rows[0].cells[i], HEAD_SHADE)
    sig.rows[0].cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
    r = sig.rows[0].cells[i].paragraphs[0].add_run(peran)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = NAVY
    cell = sig.rows[1].cells[i]
    cell.paragraphs[0].paragraph_format.space_after = Pt(46)
    run = cell.paragraphs[0].add_run(ket)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY
for row in sig.rows:
    for c in row.cells:
        c.width = Inches(3.5)

doc.add_paragraph().paragraph_format.space_after = Pt(8)
p("Dokumen internal — bukan berkas yang dikirim ke ExCB. Isian pada formulir ini ditranskrip ke Bagian 1 "
  "dokumen submission berbahasa Inggris sebelum pengajuan.",
  size=8.5, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print("written", OUT)
