# -*- coding: utf-8 -*-
"""Draf isian Application Form for ATEX (Shanghai Global Testing Services /
GTS, versi A0) yang diterima user 23 Sep 2026 -- Sumber Dokumen/input form
sertifikasi/GTS_ATEX_Application_Form_A0.pdf.

Ini dokumen INTERNAL (bukan berkas yang dikirim ke GTS): mengisi field yang
datanya sudah tersedia dari Dokumen_Teknis_Sertifikasi_GLD_IECEx_ATEX, dan
menandai jujur field yang masih perlu keputusan/konfirmasi (terutama Method
of Protection dan identitas Applicant) alih-alih menebak.

Output: Deliverables/Draf_Isian_Formulir_Aplikasi_ATEX_GTS_GLD.docx
"""
import os

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "Deliverables", "Draf_Isian_Formulir_Aplikasi_ATEX_GTS_GLD.docx")

NAVY = RGBColor(0x1B, 0x33, 0x5F)
BLUE = RGBColor(0x2B, 0x5F, 0xCB)
INK = RGBColor(0x26, 0x23, 0x21)
GRAY = RGBColor(0x55, 0x5B, 0x66)
GOOD = RGBColor(0x14, 0x91, 0x74)
WARN = RGBColor(0xB5, 0x6A, 0x1E)
GAP = RGBColor(0xB0, 0x2A, 0x37)
HEAD_SHADE = "EAEFF9"
FILL_SHADE = "FFFFFF"
NOTE_SHADE = "F3F6FC"
WARN_SHADE = "FCEFDD"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.font.color.rgb = INK

sec = doc.sections[0]
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)
sec.top_margin = Cm(1.6)
sec.bottom_margin = Cm(1.6)


def set_cell_shading(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def fixed_layout(tbl, widths):
    tbl.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl._tbl.tblPr.append(layout)
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    for row in tbl.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


def p(text="", size=10.5, bold=False, italic=False, color=None, space_after=8, space_before=0, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color or INK
    return para


def note_box(text, shade=NOTE_SHADE, text_color=GRAY, label=None, label_color=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, shade)
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    cell.paragraphs[0].paragraph_format.space_before = Pt(2)
    if label:
        r0 = cell.paragraphs[0].add_run(label + "  ")
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = label_color or text_color
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = text_color
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def field_table(rows, widths=(1.9, 3.0, 2.4)):
    """rows: (label EN/ZH, value, status_text/status_color) -> 3 kolom."""
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for i, h in enumerate(["FIELD (FORM GTS)", "ISIAN DRAF", "STATUS"]):
        set_cell_shading(hdr.cells[i], HEAD_SHADE)
        hdr.cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
        r = hdr.cells[i].paragraphs[0].add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = NAVY
    for label, value, status in rows:
        row = tbl.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        cells[0].paragraphs[0].paragraph_format.space_after = Pt(2)
        r = cells[0].paragraphs[0].add_run(label)
        r.font.bold = True
        r.font.size = Pt(9.5)
        set_cell_shading(cells[1], FILL_SHADE)
        cells[1].paragraphs[0].paragraph_format.space_after = Pt(2)
        r = cells[1].paragraphs[0].add_run(value)
        r.font.size = Pt(9.5)
        status_label, status_color = status
        cells[2].paragraphs[0].paragraph_format.space_after = Pt(2)
        r = cells[2].paragraphs[0].add_run(status_label)
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = status_color
    fixed_layout(tbl, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def checkbox_list(options, checked_labels, note_map=None):
    note_map = note_map or {}
    for opt in options:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        mark = "☒" if opt in checked_labels else "☐"
        r = para.add_run(mark + "  " + opt)
        r.font.size = Pt(9.8)
        if opt in checked_labels:
            r.font.bold = True
            r.font.color.rgb = NAVY
        if opt in note_map:
            r2 = para.add_run("  — " + note_map[opt])
            r2.font.size = Pt(8.7)
            r2.font.italic = True
            r2.font.color.rgb = GRAY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


AVAILABLE = ("Tersedia", GOOD)
PARTIAL = ("Sebagian", WARN)
TBD = ("Perlu keputusan", GAP)
DRAFT = ("Draf — cek ulang", WARN)

# ---------------------------------------------------------------- kop
lt = doc.add_paragraph()
lt.paragraph_format.space_after = Pt(0)
r = lt.add_run("PT LAPI GANESHA UTAMA")
r.font.size = Pt(13)
r.font.bold = True
r.font.color.rgb = NAVY
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(14)
r = sub.add_run("Bekerja sama dengan Institut Teknologi Bandung — Lab IoT & Lab Fisika")
r.font.size = Pt(9)
r.font.color.rgb = GRAY

p("DRAF INTERNAL — BELUM DIKIRIM", size=8.5, bold=True, color=BLUE, space_after=2)
p("Isian Application Form for ATEX (GTS, Rev. A0)", size=17.5, bold=True, color=NAVY, space_after=4)
p("Gas Leak Detector (GLD) V2 — Node Sensor", size=11, color=GRAY, space_after=12)

meta = doc.add_table(rows=0, cols=2)
meta.style = "Table Grid"
for k, v in [
    ("Formulir sumber", "Application Form for ATEX, versi A0 — Shanghai Global Testing Services Co., Ltd. "
                          "(GTS), diterima 23 Sep 2026"),
    ("Berkas asli", "Sumber Dokumen/input form sertifikasi/GTS_ATEX_Application_Form_A0.pdf"),
    ("Dokumen acuan data", "Dokumen Teknis Sertifikasi GLD IECEx/ATEX — No. LGU/GLD/IECEX-TDF/2026-001"),
    ("Tujuan", "Draf isian internal agar transkrip ke formulir resmi GTS tinggal salin — BUKAN berkas "
                "yang dikirim ke GTS"),
    ("Status", "Applicant sudah diputuskan (LGU). Sebagian field lain masih “Perlu keputusan”/data belum ada — lihat highlight"),
]:
    row = meta.add_row().cells
    set_cell_shading(row[0], HEAD_SHADE)
    for idx, val in enumerate([k, v]):
        row[idx].paragraphs[0].paragraph_format.space_after = Pt(2)
        run = row[idx].paragraphs[0].add_run(val)
        run.font.size = Pt(9.5)
        run.font.bold = (idx == 0)
        if idx == 0:
            run.font.color.rgb = NAVY
    row[0].width = Inches(1.6)
    row[1].width = Inches(5.0)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

note_box(
    "Formulir GTS memisahkan dua entitas: “Applicant” (pemohon/委托方) dan “Manufacturer” "
    "(制造商). Dikonfirmasi user (23 Sep 2026): “Applicant” = PT LAPI Ganesha Utama (LGU) — "
    "konsisten dengan peran LGU sebagai design and development authority di dokumen sertifikasi "
    "(dec:130–131); PT Galaksi Megatama Indonesia tetap “Manufacturer”. Alamat resmi, PIC "
    "kontak, dan penandatangan LGU untuk pengajuan ini BELUM ada di berkas manapun di repo — masih "
    "perlu dilengkapi sebelum formulir bisa ditranskrip penuh.",
    shade=NOTE_SHADE, text_color=INK, label="✓ Diputuskan:", label_color=GOOD,
)

# ============================================================ 1. Applicant contact
p("1. Applicant Contact Information — Info Kontak Pemohon", size=12.5, bold=True,
  color=NAVY, space_after=4)
field_table([
    ("Applicant company name — 委托方公司名称", "PT LAPI Ganesha Utama (LGU)", AVAILABLE),
    ("Applicant company address", "Belum ada di berkas manapun — perlu alamat resmi terdaftar LGU", TBD),
    ("Contact person", "Belum ditentukan — lihat Formulir_Isian_Data_Dasar butir 3.3 (PIC proyek sertifikasi)", TBD),
    ("Position", "Belum ditentukan", TBD),
    ("Mobile / Tel", "Belum ditentukan", TBD),
    ("E-mail", "Belum ditentukan", TBD),
    ("Report & invoice mailing address", "Diasumsikan alamat LGU (Applicant) — sama seperti alamat Applicant di atas, cek ulang", DRAFT),
], widths=(2.0, 3.5, 1.8))

# ============================================================ 2. Data bilingual
p("2. Data Utama (akan menjadi acuan sertifikat)", size=12.5, bold=True, color=NAVY, space_after=4)
p("Kolom “Chinese” pada formulir asli tidak diisi di sini — belum ada terjemahan Mandarin resmi "
  "untuk data produk; tanyakan ke GTS apakah wajib disediakan pemohon atau diterjemahkan oleh lab.",
  size=9, italic=True, color=GRAY, space_after=6)
field_table([
    ("Applicant 申请公司", "PT LAPI Ganesha Utama (LGU)", AVAILABLE),
    ("Add 申请公司地址", "= baris Applicant company address di atas (belum ada data)", TBD),
    ("Manufacturer 制造商", "PT Galaksi Megatama Indonesia", AVAILABLE),
    ("Add 制造商地址", "Plaza Summarecon Bekasi, Jl. Bulevar Ahmad Yani Kav. K.01, Level 7, "
     "Harapanmulya Village, Medansatria Sub-district, Bekasi City, West Java 17143, Indonesia", AVAILABLE),
    ("Product 产品名称", "Gas Leak Detector (GLD) — Node Sensor", AVAILABLE),
    ("Model No 全部型号", "GLD V2 (satu model saat ini; belum ada varian lain)", AVAILABLE),
    ("Main model 主测型号", "GLD V2", AVAILABLE),
    ("Marking Code 防爆标识", "Belum bisa diisi — tergantung Method of Protection yang belum "
     "diputuskan (lihat bagian Method of Protection). Rekomendasi tim sejauh ini (bukan keputusan ExCB): "
     "II 2G Ex [d/e/i — TBD] IIC T4 Gb.", TBD),
    ("Rating 技术参数", "24 VDC, arus maks. ≈0,33 A, daya maks. 7,995 W (≈8 W) — "
     "konfigurasi produksi (bukan varian baterai R&D)", AVAILABLE),
    ("Product size and weight 产品尺寸及重量", "200 × 90 × 290 mm; 2,3 kg", AVAILABLE),
], widths=(2.0, 4.4, 0.9))

p("Short description of product(s) / 产品的简要描述 — draf:", size=9.8, bold=True,
  color=NAVY, space_after=3)
note_box(
    "“A fixed (stationary), wall/pipe-mounted wireless sensor node for continuous flammable and toxic gas "
    "leak detection in oil & gas refinery processing areas. The unit samples ambient air via an internal "
    "MQ-series sensor array, classifies gas presence on-device using an embedded AI model, and reports "
    "readings and alarms over a LoRa wireless mesh network to a central monitoring system. Powered by a "
    "continuous 24 VDC site supply.”",
    shade=NOTE_SHADE, text_color=INK, label="Draf (cek ulang):", label_color=WARN,
)

# ============================================================ Method of protection
p("Method of Protection — Belum diputuskan (penghambat utama, lihat dec:98)", size=12.5, bold=True,
  color=GAP, space_after=4)
note_box(
    "Ini bukan sekadar formulir kosong — memilih salah satu opsi ini adalah keputusan rekayasa yang "
    "menentukan seluruh scope pengujian (flame-path, komponen bersertifikat Ex, dst.). Belum ada keputusan "
    "resmi. Kandidat yang relevan berdasarkan riset internal (bukan rekomendasi final): “Ex d” "
    "(flameproof enclosure — cocok untuk casing die-cast aluminium ADC12 yang sudah ada) dan/atau "
    "“Ex i” (intrinsic safety — relevan karena catu daya rendah, 24VDC/~8W) untuk sirkuit "
    "internal. Jangan mencentang opsi manapun sebelum ada keputusan tim engineering.",
    shade=WARN_SHADE, text_color=WARN, label="⚠",
)
checkbox_list([
    "Gas – Flameproof enclosures \"d\"",
    "Gas and Dust – Pressurized enclosures \"p\"",
    "Gas – Powder filling \"q\"",
    "Gas – Oil immersion \"o\"",
    "Gas – Increased safety \"e\"",
    "Gas and Dust – Intrinsic safety \"i\"",
    "Gas – Type of protection \"n\"",
    "Gas – Optical radiation \"op\"",
    "Gas and Dust – Encapsulation \"m\"",
    "Dust – Protection by enclosures \"t\"",
    "Safety-, Control-, Regulation Device (EN 50495)",
], checked_labels=[])

# ============================================================ Group / subgroup / level
p("Group Equipment, Explosion Subgroup, Level of Protection", size=12.5, bold=True, color=NAVY, space_after=4)
p("Ditandai berdasarkan rekomendasi tim yang sudah tercatat di Dokumen Teknis Sertifikasi §2.5 — "
  "“Group II” konsisten dengan penggunaan non-tambang (refinery) dan relatif tidak ambigu; "
  "Subgroup/Level tetap berstatus rekomendasi tim, bukan keputusan ExCB.", size=9, italic=True, color=GRAY,
  space_after=6)
checkbox_list([
    "Group I (mining)",
    "Group II (Non-mining equipment)",
    "Group III (Dust, IEC only)",
], checked_labels=["Group II (Non-mining equipment)"],
    note_map={"Group II (Non-mining equipment)": "sesuai penggunaan refinery/non-tambang — relatif pasti"})
checkbox_list([
    "Subgroup II", "Subgroup IIA", "Subgroup IIB", "Subgroup IIC",
], checked_labels=["Subgroup IIC"],
    note_map={"Subgroup IIC": "rekomendasi tim (H₂ termasuk gas target) — BUKAN keputusan ExCB, lihat §2.5"})
checkbox_list([
    "Category 1 G/Ga (Zone 0)", "Category 2 G/Gb (Zone 1)", "Category 3 G/Gc (Zone 2)",
], checked_labels=["Category 2 G/Gb (Zone 1)"],
    note_map={"Category 2 G/Gb (Zone 1)": "rekomendasi tim — BUKAN keputusan ExCB, lihat §2.5"})

# ============================================================ IP / classification / misc
p("IP Protection, Classification of Installation and Use, dan lainnya", size=12.5, bold=True, color=NAVY,
  space_after=4)
field_table([
    ("IP Protection", "IP66", AVAILABLE),
    ("Classification of installation and use", "Stationary (dipasang tetap via bracket/U-bolt ke struktur pipa "
     "existing — bukan portable maupun hand-held)", AVAILABLE),
    ("Intended use limit for a certain kind of atmosphere", "Flammable/toxic gas atmospheres di area proses "
     "kilang minyak & gas; gas target sesuai model AI (Bagian 2.3.b) — draf, cek ulang cakupan resmi", DRAFT),
    ("Ambient pressure range", "Tidak ada data pengujian — disarankan centang “No specified standard, "
     "lab to recommend” (无指定标准，由实验室推荐)", PARTIAL),
    ("Sample Return", "Belum diputuskan — perlu keputusan admin/logistik (retur berbayar / diambil sendiri "
     "/ dimusnahkan)", TBD),
], widths=(2.3, 4.2, 0.8))

# ============================================================ page break -> documents submitted
doc.add_page_break()
p("3. Documents Submitted — Status Kesiapan Berkas", size=12.5, bold=True, color=NAVY, space_after=4)
p("Dicocokkan terhadap isi Dokumen Teknis Sertifikasi GLD IECEx/ATEX (Bagian 2.6) per 23 Sep 2026.", size=9,
  italic=True, color=GRAY, space_after=6)
field_table([
    ("User Instructions (English)", "Tersedia — Instruction_Manual_GLD.docx (Rev 0.8)", AVAILABLE),
    ("Wiring/Circuit diagram", "Tersedia — diagram blok skematik 9 lembar penuh (§2.6.a)", AVAILABLE),
    ("CDF (critical component list)", "Sebagian — BOM elektronik lengkap ada (EasyEDA/JLCPCB), tapi daftar "
     "komponen kritis Ex-safety-relevant (enclosure/gasket/cable gland/baterai) belum disusun terpisah", PARTIAL),
    ("Drawings (structure)", "Sebagian — gambar CAD berdimensi casing v3 & bracket ada, tapi gambar struktur "
     "resmi dari mitra casing (PT Galaksi) dengan parameter flame-path belum ada", PARTIAL),
    ("Specification", "Tersedia — Dokumen_Teknis_Sertifikasi_GLD_IECEx_ATEX itu sendiri berfungsi sbg spec", AVAILABLE),
    ("Rating Label", "Belum ada — desain nameplate/marking belum dibuat (§2.6.h)", ("Belum tersedia", GAP)),
    ("Declaration of Difference Between Models", "Tidak relevan saat ini — hanya 1 model (GLD V2), tidak ada "
     "varian yang perlu dideklarasikan bedanya", ("N/A", GRAY)),
    ("Copy of Critical Component Certificate", "Belum ada — belum ada komponen Ex-certified yang teridentifikasi "
     "dalam BOM (§2.6.i)", ("Belum tersedia", GAP)),
], widths=(2.3, 4.3, 0.7))

p("Pengesahan isian", size=12.5, bold=True, color=NAVY, space_after=8, space_before=10)
sig = doc.add_table(rows=2, cols=2)
sig.style = "Table Grid"
for i, (peran, ket) in enumerate([
    ("Diisi oleh", "Nama, jabatan, tanggal, tanda tangan"),
    ("Diperiksa oleh", "Nama, jabatan, tanggal, tanda tangan"),
]):
    set_cell_shading(sig.rows[0].cells[i], HEAD_SHADE)
    sig.rows[0].cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
    r = sig.rows[0].cells[i].paragraphs[0].add_run(peran)
    r.font.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = NAVY
    cell = sig.rows[1].cells[i]
    cell.paragraphs[0].paragraph_format.space_after = Pt(46)
    run = cell.paragraphs[0].add_run(ket)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRAY
for row in sig.rows:
    for c in row.cells:
        c.width = Inches(3.4)

doc.add_paragraph().paragraph_format.space_after = Pt(8)
p("Dokumen internal — draf isian, bukan berkas yang dikirim ke GTS. Field bertanda “Perlu keputusan” "
  "harus diselesaikan sebelum transkrip ke formulir resmi.", size=8.5, italic=True, color=GRAY,
  align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print("written", OUT)
