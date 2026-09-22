# -*- coding: utf-8 -*-
"""Build Deliverables/Instruction_Manual_GLD.docx from the canonical HTML source.

Walks the HTML body (BeautifulSoup) and re-renders it as a corporate-format
Word document (python-docx), matching the visual language already used for
Dokumen_Teknis_Sertifikasi_GLD_IECEx_ATEX.docx (same helper patterns).
"""
import os
import re
import base64
import binascii

from bs4 import BeautifulSoup, NavigableString, Tag
from PIL import Image
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_HTML = os.path.join(REPO, "Deliverables", "Instruction_Manual_GLD.html")
OUT_PATH = os.path.join(REPO, "Deliverables", "Instruction_Manual_GLD.docx")
TMP_IMG_DIR = os.path.join(REPO, "scripts", "assets", "_tmp_instr_manual_imgs")
os.makedirs(TMP_IMG_DIR, exist_ok=True)

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
GRAY = RGBColor(0x5B, 0x5B, 0x5B)
INK = RGBColor(0x1F, 0x1F, 0x1F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2C, 0x6B, 0x33)
AMBER = RGBColor(0x8A, 0x6D, 0x1E)
RED = RGBColor(0xB2, 0x3A, 0x3A)

LINE_GRAY = "D8D8D8"
HEAD_SHADE = "EFEFEC"
ZEBRA_SHADE = "FAFAF8"
WARN_BG, WARN_BD = "FFF6E5", "C98A12"
DANGER_BG, DANGER_BD = "FBE9E9", "B23A3A"
OK_BG, OK_BD = "EAF4EC", "3C7A44"
NOTE_BG, NOTE_BD = "EFF3F7", "7C93AE"

STATUS_COLOR = {
    "status-final": GREEN,
    "status-partial": AMBER,
    "status-pending": RED,
}

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.25

for i, sz, col in [(1, 16, NAVY), (2, 13, NAVY), (3, 11, NAVY)]:
    st = doc.styles[f"Heading {i}"]
    st.font.name = "Calibri"
    st.font.size = Pt(sz)
    st.font.bold = True
    st.font.color.rgb = col
    st.paragraph_format.space_before = Pt(16 if i == 1 else 12)
    st.paragraph_format.space_after = Pt(6)

sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(1.8)

# ---------- low-level helpers ----------

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def set_cell_border(cell, color=LINE_GRAY, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)

def add_field(paragraph, field_code, fallback_text="Right-click and choose Update Field."):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = field_code
    fld_sep = OxmlElement("w:fldChar"); fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t"); fld_text.text = fallback_text
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    r_el = run._r
    r_el.append(fld_begin); r_el.append(instr)
    r2 = paragraph.add_run(); r2._r.append(fld_sep)
    r3 = paragraph.add_run(); r3._r.append(fld_text)
    r4 = paragraph.add_run(); r4._r.append(fld_end)

def p(text="", size=10.5, bold=False, italic=False, color=None, space_after=8, align=None, style=None):
    para = doc.add_paragraph(style=style)
    if align is not None:
        para.alignment = align
    para.paragraph_format.space_after = Pt(space_after)
    if text:
        run = para.add_run(text)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = color
    return para

def add_picture_fit(paragraph, path, max_w_in, max_h_in=None):
    with Image.open(path) as im:
        w_px, h_px = im.size
    aspect = w_px / h_px
    w_in = max_w_in
    h_in = w_in / aspect
    if max_h_in and h_in > max_h_in:
        h_in = max_h_in
        w_in = h_in * aspect
    run = paragraph.add_run()
    run.add_picture(path, width=Inches(w_in), height=Inches(h_in))

_img_counter = [0]

def save_data_uri(src):
    """Decode a data:image/...;base64,XXXX URI to a temp file, return its path."""
    m = re.match(r"data:image/(\w+);base64,(.*)", src, re.S)
    if not m:
        return None
    ext, b64 = m.group(1), m.group(2)
    ext = "jpg" if ext == "jpeg" else ext
    _img_counter[0] += 1
    path = os.path.join(TMP_IMG_DIR, f"img_{_img_counter[0]}.{ext}")
    try:
        data = base64.b64decode(b64)
    except binascii.Error:
        return None
    with open(path, "wb") as f:
        f.write(data)
    return path

# ---------- inline run walker (bold / status-colored spans inside <p>/<li>/<td>) ----------

def add_inline(paragraph, node, size=10.5, base_bold=False):
    if isinstance(node, NavigableString):
        text = str(node)
        if text:
            run = paragraph.add_run(text)
            run.font.size = Pt(size)
            run.font.bold = base_bold
        return
    if not isinstance(node, Tag):
        return
    if node.name in ("b", "strong"):
        for c in node.children:
            add_inline(paragraph, c, size=size, base_bold=True)
    elif node.name == "span" and node.get("class"):
        cls = node.get("class")[0]
        color = STATUS_COLOR.get(cls)
        run = paragraph.add_run(node.get_text())
        run.font.size = Pt(size)
        run.font.bold = True
        if color:
            run.font.color.rgb = color
    elif node.name == "br":
        paragraph.add_run().add_break()
    else:
        for c in node.children:
            add_inline(paragraph, c, size=size, base_bold=base_bold)

def add_rich_paragraph(tag, size=10.5, space_after=8, style=None):
    para = doc.add_paragraph(style=style)
    para.paragraph_format.space_after = Pt(space_after)
    for c in tag.children:
        add_inline(para, c, size=size)
    return para

# ---------- block-level renderers ----------

def render_ul(ul_tag, size=10.0):
    for li in ul_tag.find_all("li", recursive=False):
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(3)
        for c in li.children:
            add_inline(bp, c, size=size)

def render_table(table_tag, col_widths=None):
    rows = table_tag.find_all("tr", recursive=False)
    if not rows:
        return
    ncols = len(rows[0].find_all(["th", "td"], recursive=False))
    tbl = doc.add_table(rows=0, cols=ncols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, tr in enumerate(rows):
        cells_src = tr.find_all(["th", "td"], recursive=False)
        row_cells = tbl.add_row().cells
        is_header = ri == 0 and tr.find("th") is not None
        for ci, cell_src in enumerate(cells_src):
            if ci >= ncols:
                break
            cell = row_cells[ci]
            cell.paragraphs[0].paragraph_format.space_after = Pt(2)
            if is_header:
                set_cell_shading(cell, HEAD_SHADE)
                run = cell.paragraphs[0].add_run(cell_src.get_text().strip())
                run.font.bold = True
                run.font.size = Pt(9)
                run.font.color.rgb = GRAY
            else:
                if ri % 2 == 0:
                    set_cell_shading(cell, ZEBRA_SHADE)
                for c in cell_src.children:
                    add_inline(cell.paragraphs[0], c, size=9.5)
    if col_widths:
        for row in tbl.rows:
            for i, w in enumerate(col_widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl

BOX_STYLE = {
    "warning": (DANGER_BG, DANGER_BD),
    "caution": (WARN_BG, WARN_BD),
    "note":    (NOTE_BG, NOTE_BD),
    "status":  (OK_BG, OK_BD),
}

def render_box(box_tag):
    classes = box_tag.get("class", [])
    kind = next((c for c in classes if c in BOX_STYLE), "note")
    bg, bd = BOX_STYLE[kind]
    lbl_tag = box_tag.find("span", class_="lbl")
    label = lbl_tag.get_text().strip() if lbl_tag else kind.upper()

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, bg)
    set_cell_border(cell, bd, sz=6)
    cell.paragraphs[0].paragraph_format.space_after = Pt(4)
    lr = cell.paragraphs[0].add_run(label.upper())
    lr.font.bold = True
    lr.font.size = Pt(9)
    lr.font.color.rgb = INK

    body_added = False
    for child in box_tag.children:
        if not isinstance(child, Tag):
            continue
        if child.name == "span" and "lbl" in (child.get("class") or []):
            continue
        if child.name == "ul":
            for li in child.find_all("li", recursive=False):
                lp = cell.add_paragraph()
                lp.paragraph_format.space_after = Pt(3)
                dash = lp.add_run("• ")
                dash.font.size = Pt(10); dash.font.bold = True
                for c in li.children:
                    add_inline(lp, c, size=10)
            body_added = True
        elif child.name == "p":
            bp = cell.add_paragraph()
            bp.paragraph_format.space_after = Pt(2)
            for c in child.children:
                add_inline(bp, c, size=10)
            body_added = True
    if not body_added:
        # plain text content directly inside the box (rare)
        text = box_tag.get_text().replace(label, "", 1).strip()
        if text:
            bp = cell.add_paragraph()
            bp.add_run(text).font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def render_figure(fig_tag, full_width=True):
    img = fig_tag.find("img")
    figcap = fig_tag.find("figcaption")
    if img and img.get("src", "").startswith("data:image"):
        path = save_data_uri(img["src"])
        if path:
            fp = doc.add_paragraph()
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_picture_fit(fp, path, max_w_in=6.2 if full_width else 3.0,
                             max_h_in=None if full_width else 2.6)
    if figcap:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(10)
        for c in figcap.children:
            add_inline(cp, c, size=9)
        for run in cp.runs:
            run.font.color.rgb = GRAY
            run.font.italic = True

# ============================================================
# PARSE SOURCE HTML
# ============================================================
soup = BeautifulSoup(open(SRC_HTML, encoding="utf-8").read(), "html.parser")
wrap = soup.find("main", class_="wrap")

cover = wrap.find("div", class_="cover")
letterhead = wrap.find("div", class_="letterhead")

# ---------- letterhead band ----------
lt = doc.add_table(rows=1, cols=1)
lc = lt.rows[0].cells[0]
set_cell_shading(lc, "1B2A4A")
lc.paragraphs[0].paragraph_format.space_after = Pt(2)
org_tag = letterhead.find("div", class_="org")
org_name = org_tag.find("b").get_text() if org_tag else "PT LAPI GANESHA UTAMA"
org_sub = org_tag.find("span").get_text() if org_tag and org_tag.find("span") else ""
lr = lc.paragraphs[0].add_run(org_name)
lr.font.bold = True; lr.font.size = Pt(14); lr.font.color.rgb = WHITE
if org_sub:
    lp2 = lc.add_paragraph()
    lp2.paragraph_format.space_after = Pt(3)
    lr2 = lp2.add_run(org_sub)
    lr2.font.size = Pt(9.5); lr2.font.color.rgb = RGBColor(0xC7, 0xD2, 0xE0)
doc_tag = letterhead.find("div", class_="doc")
doc.add_paragraph().paragraph_format.space_after = Pt(6)
if doc_tag:
    p(doc_tag.get_text(" ", strip=True), size=9, color=GRAY, space_after=14)

# ---------- title / subtitle ----------
h1 = cover.find("h1")
title = doc.add_heading(level=0)
title.paragraph_format.space_after = Pt(4)
tr = title.add_run(h1.get_text(" ", strip=True).replace("\n", " "))
tr.font.size = Pt(20); tr.font.bold = True; tr.font.color.rgb = NAVY
subtitle = cover.find("p", class_="subtitle")
if subtitle:
    p(subtitle.get_text(strip=True), size=12, color=GRAY, space_after=12)

# ---------- document control table ----------
docctl = cover.find("table", class_="docctl")
p("Document Control", size=11.5, bold=True, color=NAVY, space_after=4)
mt = doc.add_table(rows=0, cols=2)
mt.style = "Table Grid"
for tr_ in docctl.find_all("tr"):
    th, td = tr_.find("th"), tr_.find("td")
    row = mt.add_row().cells
    set_cell_shading(row[0], HEAD_SHADE)
    row[0].paragraphs[0].paragraph_format.space_after = Pt(2)
    r = row[0].paragraphs[0].add_run(th.get_text(strip=True))
    r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = GRAY
    row[1].paragraphs[0].paragraph_format.space_after = Pt(2)
    for c in td.children:
        add_inline(row[1].paragraphs[0], c, size=10)
    row[0].width = Inches(2.0); row[1].width = Inches(4.5)
doc.add_paragraph().paragraph_format.space_after = Pt(10)

# ---------- scope banner ----------
scope = cover.find("div", class_="scope")
if scope:
    tbl = doc.add_table(rows=1, cols=1)
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, WARN_BG)
    set_cell_border(cell, WARN_BD, sz=6)
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    for c in scope.children:
        add_inline(cell.paragraphs[0], c, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

# ---------- revision history ----------
revtable = cover.find("table", class_="revtable")
p("Revision History", size=11.5, bold=True, color=NAVY, space_after=4)
render_table(revtable, col_widths=[0.9, 1.1, 4.5])

doc.add_page_break()

# ---------- TOC ----------
for _toc_style, _toc_size in (("TOC 1", 10.5),):
    try:
        _st = doc.styles[_toc_style]
    except KeyError:
        from docx.enum.style import WD_STYLE_TYPE
        _st = doc.styles.add_style(_toc_style, WD_STYLE_TYPE.PARAGRAPH)
    _st.font.size = Pt(_toc_size)
    _st.font.name = "Calibri"
    _st.paragraph_format.space_before = Pt(0)
    _st.paragraph_format.space_after = Pt(4)

p("Table of Contents", size=13, bold=True, color=NAVY, space_after=6)
toc_para = doc.add_paragraph()
add_field(toc_para, 'TOC \\o "1-2" \\h \\z \\u',
          fallback_text="Right-click and choose Update Field to generate the table of contents.")

doc.add_page_break()

# ---------- header / footer ----------
header = sec.header
header.is_linked_to_previous = False
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("GLD Instruction Manual  |  Working Draft — Confidential")
hr.font.size = Pt(8); hr.font.color.rgb = GRAY; hr.font.italic = True

footer = sec.footer
footer.is_linked_to_previous = False
fp = footer.paragraphs[0]
fr1 = fp.add_run("LGU/GLD/INSTR-MAN/2026-001  ·  Rev. 0.8")
fr1.font.size = Pt(8); fr1.font.color.rgb = GRAY
tab_stops = fp.paragraph_format.tab_stops
tab_stops.add_tab_stop(Inches(6.5), alignment=2)
fp.add_run("\t")
add_field(fp, "PAGE", fallback_text="1")
fp.add_run(" of ")
add_field(fp, "NUMPAGES", fallback_text="1")
for run in fp.runs:
    run.font.size = Pt(8); run.font.color.rgb = GRAY

# ============================================================
# BODY (sections 1..17)
# ============================================================
section_div = wrap.find("div", class_="section")
body_nodes = [n for n in section_div.children if isinstance(n, Tag)]
for node in body_nodes:
    if node.name == "footer":
        continue
    if node.name == "h2":
        heading_text = node.get_text(" ", strip=True)
        heading_text = re.sub(r"^\s*(\d+)\.\s*", r"\1. ", heading_text)
        doc.add_heading(heading_text, level=1)
        continue
    if node.name == "h3":
        doc.add_heading(node.get_text(" ", strip=True), level=2)
    elif node.name == "p":
        add_rich_paragraph(node, size=10.5, space_after=8)
    elif node.name == "ul":
        render_ul(node)
    elif node.name == "table" and "spec" in (node.get("class") or []):
        render_table(node)
    elif node.name == "div" and "box" in (node.get("class") or []):
        render_box(node)
    elif node.name == "div" and "photos" in (node.get("class") or []):
        for fig in node.find_all("figure", recursive=False):
            render_figure(fig, full_width=False)
    elif node.name == "figure" and "cadfig" in (node.get("class") or []):
        render_figure(node, full_width=True)
    # other node types (stray divs, etc.) are skipped deliberately

# ---------- closing note ----------
doc.add_paragraph().paragraph_format.space_before = Pt(10)
foot = doc.add_paragraph()
r = foot.add_run(
    "This document is a working draft, structured from the instruction manual of a comparable certified "
    "diffusion-type gas detector (New Cosmos KD-12/KD-12R) and populated with GLD's own confirmed design data. "
    "It is not yet an approved operational manual, and GLD is not yet ATEX/IECEx certified. Fields marked "
    "“Pending confirmation” must not be relied upon for procurement, installation, or certification "
    "purposes without further verification."
)
r.font.size = Pt(9); r.font.color.rgb = GRAY; r.font.italic = True

doc.save(OUT_PATH)
print("written", OUT_PATH)
