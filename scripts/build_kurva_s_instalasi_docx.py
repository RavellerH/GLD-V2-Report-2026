from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
import datetime as dt

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "Deliverables"
OUT_DOCX = OUT_DIR / "Kurva_S_Persiapan_Instalasi_RU-IV_Cilacap.docx"
CHART_PNG = ROOT / "scripts" / "assets" / "kurva_s_instalasi_chart.png"


def build_chart():
    CHART_PNG.parent.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial", "DejaVu Sans"],
        "font.size": 12,
        "axes.edgecolor": "#000000",
        "text.color": "#000000",
        "axes.labelcolor": "#000000",
        "xtick.color": "#000000",
        "ytick.color": "#000000",
    })

    d0 = dt.date(2026, 8, 10)  # kunjungan/survey terakhir RU IV Cilacap (9-10 Agustus 2026)
    d1 = dt.date(2026, 9, 18)

    fig, ax = plt.subplots(figsize=(9.6, 4.7), dpi=200)

    # actual line (solid black) - LGU & ITB readiness only (what was asked)
    ax.plot([d0, d1], [0, 100], color="#000000", linewidth=2.6, marker="o",
             markersize=7, markerfacecolor="#000000", zorder=5)

    # 95% threshold reference line (disebutkan Pertamina utk syarat agenda meeting RU IV)
    ax.axhline(95, color="#000000", linewidth=1.1, linestyle=(0, (6, 4)), alpha=0.65, zorder=2)
    ax.text(d0 + dt.timedelta(days=1), 91, "95% — ambang meeting lanjutan RU IV (disebut Pertamina)",
            ha="left", va="top", fontsize=9.2, color="#000000", fontweight="bold")

    # data point labels
    ax.annotate("9–10 Agustus 2026\nsurvey lokasi RU IV Cilacap\n(kunjungan terakhir)\n0%", xy=(d0, 0), xytext=(0, -42),
                textcoords="offset points", ha="center", va="top", fontsize=9.3, color="#262626")
    ax.annotate("100%", xy=(d1, 100), xytext=(-16, 4), textcoords="offset points",
                ha="right", va="center", fontsize=14, fontweight="bold", color="#000000")
    ax.annotate("18 Sep 2026\n(hari ini)", xy=(d1, 100), xytext=(-6, -26), textcoords="offset points",
                ha="right", va="top", fontsize=9.3, color="#262626")

    ax.set_xlim(d0 - dt.timedelta(days=2), d1 + dt.timedelta(days=2))
    ax.set_ylim(-10, 112)
    ax.set_yticks([0, 25, 50, 75, 100])
    ax.set_yticklabels(["0%", "25%", "50%", "75%", "100%"])
    ax.xaxis.set_major_locator(mdates.DayLocator(interval=7))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%d %b"))
    ax.grid(axis="y", color="#D9D9D9", linewidth=0.8, zorder=0)
    for spine in ("top", "right"):
        ax.spines[spine].set_visible(False)
    for spine in ("left", "bottom"):
        ax.spines[spine].set_color("#000000")
        ax.spines[spine].set_linewidth(1.1)
    ax.tick_params(colors="#000000", labelsize=9.5)

    fig.tight_layout(pad=1.2)
    fig.savefig(CHART_PNG, facecolor="white", bbox_inches="tight")
    plt.close(fig)
    return CHART_PNG

BLACK = "000000"
DARK = "262626"
MID = "808080"
LINE = "BFBFBF"
ZEBRA = "F2F2F2"
HEADFILL = "1A1A1A"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_run_font(run, name="Arial", size=9.5, bold=None, italic=None, color=DARK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def format_cell(cell, bold=False, color=DARK, size=8.9, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_run_font(run, size=8, color=MID)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run2 = paragraph.add_run(" dari ")
    set_run_font(run2, size=8, color=MID)
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "begin")
    instr2 = OxmlElement("w:instrText")
    instr2.set(qn("xml:space"), "preserve")
    instr2.text = " NUMPAGES "
    fld4 = OxmlElement("w:fldChar")
    fld4.set(qn("w:fldCharType"), "end")
    run2._r.append(fld3)
    run2._r.append(instr2)
    run2._r.append(fld4)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    for level, size in ((1, 14), (2, 11.5)):
        style = styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    if level == 1:
        p.paragraph_format.border_bottom = None
        pPr = p._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "3")
        bottom.set(qn("w:color"), BLACK)
        pbdr.append(bottom)
        pPr.append(pbdr)
    r = p.add_run(text)
    set_run_font(r, size=14 if level == 1 else 11.5, bold=True, color=BLACK)
    return p


def add_paragraph(doc, text="", bold_lead=None, align=None, size=9.8, color=DARK, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.22
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=size, bold=True, color=color)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=size, color=color)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color)
    return p


def add_table(doc, headers, rows, widths, font_size=8.7, mark_col=None, section_rows=None):
    """rows: list of tuples. section_rows: set of row indices (0-based, within `rows`) that
    are full-width category header rows (single string in rows[i][0])."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = header
        set_cell_shading(cell, HEADFILL)
        cell.width = Inches(widths[idx])
        format_cell(cell, bold=True, color=WHITE, size=8.3, align=WD_ALIGN_PARAGRAPH.CENTER)

    body_idx = 0
    for values in rows:
        row = table.add_row()
        if section_rows and body_idx in section_rows:
            merged = row.cells[0]
            for c in row.cells[1:]:
                merged = merged.merge(c)
            merged.text = str(values[0])
            set_cell_shading(merged, "D9D9D9")
            format_cell(merged, bold=True, size=8.4, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT)
            body_idx += 1
            continue
        zebra = (body_idx % 2 == 1)
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = str(value)
            cell.width = Inches(widths[idx])
            if zebra:
                set_cell_shading(cell, ZEBRA)
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or idx == mark_col else WD_ALIGN_PARAGRAPH.LEFT
            bold = (idx == mark_col)
            format_cell(cell, size=font_size, align=align, bold=bold)
        body_idx += 1
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_document_control(doc):
    rows = [
        ("Nomor dokumen", "LGU/GLD/KURVA-INSTALASI/2026-001"),
        ("Revisi", "1.0"),
        ("Tanggal", "18 September 2026"),
        ("Status", "Untuk koordinasi PT Pertamina Patra Niaga — RU IV Cilacap"),
        ("Disiapkan oleh", "LAPI Ganesha Utama bersama Lab IoT/Instrumentation and Computation ITB"),
        ("Ditujukan kepada", "PT Pertamina Patra Niaga — RU IV Cilacap"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, (label, value) in enumerate(rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        table.rows[i].cells[0].width = Inches(1.7)
        table.rows[i].cells[1].width = Inches(4.8)
        set_cell_shading(table.rows[i].cells[0], "E8E8E8")
        format_cell(table.rows[i].cells[0], bold=True, size=9)
        format_cell(table.rows[i].cells[1], size=9)


def kpi_row(doc, label, value, sub):
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color=LINE, size="5")
    c0 = table.rows[0].cells[0]
    c0.width = Inches(1.55)
    c0.text = label
    set_cell_shading(c0, "E8E8E8")
    format_cell(c0, bold=True, size=7.6, align=WD_ALIGN_PARAGRAPH.CENTER, color=DARK)
    c1 = table.rows[1].cells[0]
    c1.width = Inches(1.55)
    c1.text = value
    format_cell(c1, bold=True, size=17, align=WD_ALIGN_PARAGRAPH.CENTER, color=BLACK)
    return table


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    header = section.header.paragraphs[0]
    header.text = "GLD TAHAP 2   |   PERSIAPAN INSTALASI RU IV CILACAP"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=7.8, bold=True, color=DARK)
    add_page_number(section.footer.paragraphs[0])

    # Cover block
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(50)
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PT LAPI GANESHA UTAMA")
    set_run_font(r, size=10.5, bold=True, color=MID)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("KURVA-S PERSIAPAN INSTALASI")
    set_run_font(r, size=23, bold=True, color=BLACK)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(20)
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("RU IV Cilacap")
    set_run_font(r2, size=17, bold=False, color=DARK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("Status kesiapan sebelum mobilisasi & pemasangan fisik — LGU & ITB, Pertamina RU IV, Vendor Instalasi")
    set_run_font(r, size=10, italic=True, color=MID)

    add_document_control(doc)
    doc.add_page_break()

    # 1 Ringkasan
    add_heading(doc, "1  Ringkasan Kesiapan", 1)
    add_paragraph(
        doc,
        "Cakupan dokumen. Dokumen ini melaporkan status kesiapan sebelum mobilisasi dan pemasangan fisik "
        "GLD di RU IV Cilacap, dihitung dari checklist kerja bersama tiga pihak (LGU & ITB, Pertamina RU IV, "
        "Vendor Instalasi). Ini bukan Kurva-S proyek GLD Tahap 2 secara keseluruhan (9 bulan) — cakupannya "
        "khusus tahap persiapan instalasi di RU IV Cilacap.",
        bold_lead="Cakupan dokumen. ",
    )

    tbl = doc.add_table(rows=1, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, (lab, val, sub) in enumerate([
        ("LGU & ITB", "100%", "15 dari 15 item"),
        ("PERTAMINA RU IV", "0%", "0 dari 14 item"),
        ("VENDOR INSTALASI", "0%", "0 dari 11 item"),
        ("KESELURUHAN", "38%", "15 dari 40 item"),
    ]):
        cell = tbl.rows[0].cells[i]
        cell.width = Inches(1.75)
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(lab + "\n")
        set_run_font(r1, size=7.6, bold=True, color=MID)
        r2 = p1.add_run(val + "\n")
        set_run_font(r2, size=19, bold=True, color=BLACK)
        r3 = p1.add_run(sub)
        set_run_font(r3, size=7.8, color=MID)
        set_cell_shading(cell, "F2F2F2" if i < 3 else "E8E8E8")
        set_cell_margins(cell, top=140, bottom=140)
    set_table_borders(tbl, color=LINE, size="5")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Warning box - permit clarification
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = box.rows[0].cells[0]
    set_cell_shading(cell, "EDEDED")
    set_cell_margins(cell, top=110, bottom=110, start=130, end=130)
    p = cell.paragraphs[0]
    r1 = p.add_run("CATATAN PENTING — ")
    set_run_font(r1, size=9, bold=True, color=BLACK)
    r2 = p.add_run(
        "angka 100% pada kolom LGU & ITB mencakup desain, unit, dokumen kerja, dan kesiapan teknis — "
    )
    set_run_font(r2, size=9, color=DARK)
    r3 = p.add_run("bukan perizinan. ")
    set_run_font(r3, size=9, bold=True, color=BLACK)
    r4 = p.add_run(
        "Seluruh proses perizinan (pengesahan TRA/JSA, izin kerja, izin masuk personel, dan surat masuk "
        "barang ke area RU IV) sepenuhnya berada pada wewenang dan tanggung jawab Pertamina RU IV, "
        "tercakup dalam angka 0/14 pada kolom Pertamina RU IV — bukan bagian dari 15 item LGU."
    )
    set_run_font(r4, size=9, color=DARK)
    tbl_pr = box._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "10")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BLACK)
        borders.append(node)
    tbl_pr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 2 Tren
    add_heading(doc, "2  Tren Persiapan", 1)
    add_paragraph(
        doc,
        "Metodologi. Kurva ini melacak progres LGU & ITB secara khusus — sesuai pertanyaan yang diajukan "
        "Pertamina ('dari sisi LGU persiapannya sudah berapa persen?'), bukan angka gabungan tiga pihak. "
        "Titik awal diambil dari kunjungan/survey lapangan terakhir ke RU IV Cilacap (9–10 Agustus 2026) "
        "— belum ada kunjungan lain setelahnya — dibandingkan dengan snapshot kesiapan hari ini "
        "(18 September 2026).",
        bold_lead="Metodologi. ",
    )

    chart_path = build_chart()
    doc.add_picture(str(chart_path), width=Inches(6.7))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run("Grafik 1. Kurva-S kesiapan LGU & ITB, persiapan instalasi RU IV Cilacap")
    set_run_font(r, size=8.3, italic=True, color=MID)

    add_table(
        doc,
        ["Tanggal", "Peristiwa", "Progres LGU & ITB"],
        [
            ("9–10 Agustus 2026", "Survey lokasi RU IV Cilacap (kunjungan lapangan terakhir)", "0%"),
            ("18 September 2026 (hari ini)", "Seluruh 15 item tanggung jawab LGU & ITB selesai/siap", "100%"),
        ],
        [1.7, 3.85, 1.25],
        font_size=9,
    )
    add_paragraph(
        doc,
        "* Ambang yang disebutkan Pertamina pada percakapan 18 September 2026 (>95% agar meeting lanjutan "
        "bersama RU IV dapat diagendakan) sudah terlampaui dari sisi kesiapan LGU & ITB. Ini tidak mencakup "
        "kesiapan Pertamina RU IV dan Vendor Instalasi, yang saat dokumen ini disusun masih 0% (lihat Bagian 1 "
        "dan 3) — perizinan dan penugasan resmi tetap berada di luar kendali LGU.",
        size=8.6, color=MID, space_after=10,
    )

    doc.add_page_break()

    # 3 Breakdown
    add_heading(doc, "3  Breakdown per Kategori dan Pihak", 1)
    add_paragraph(
        doc,
        "Jumlah item persiapan yang sudah selesai/siap dibandingkan total item pada kategori tersebut, per "
        "pihak. Rincian penuh per item ada pada Bagian 4."
    )
    add_table(
        doc,
        ["Kategori", "LGU & ITB", "Pertamina RU IV", "Vendor Instalasi"],
        [
            ("1. Lokasi & klasifikasi area", "2 / 2", "0 / 2", "0 / 2"),
            ("2. Dokumen, izin & K3", "2 / 2", "0 / 4", "0 / 3"),
            ("3. Perangkat, mekanik & bracket", "3 / 3", "0 / 2", "0 / 3"),
            ("4. Catu daya & kabel", "2 / 2", "0 / 2", "0 / 2"),
            ("5. Jaringan & server", "3 / 3", "0 / 2", "—"),
            ("6. Koordinasi & commissioning hari-H", "3 / 3", "0 / 2", "0 / 1"),
            ("Total", "15 / 15", "0 / 14", "0 / 11"),
        ],
        [2.9, 1.1, 1.35, 1.35],
        font_size=9,
    )

    doc.add_page_break()

    # 4 Rincian per item
    add_heading(doc, "4  Rincian per Item (40 Item)", 1)
    add_paragraph(
        doc,
        "Metodologi. Item ditandai SELESAI apabila status tertulis SELESAI / TERSEDIA / SIAP / SUDAH; selain "
        "itu ditandai BELUM. Sumber: “Pembagian Persiapan Instalasi RU IV Cilacap” (17 September 2026).",
        bold_lead="Metodologi. ",
    )

    ITEMS = [
        ("1. LOKASI & KLASIFIKASI AREA",),
        ("1", "Basis desain mounting (pelat mounting U-bolt 2″/DN50), non-invasive — tanpa bor/las", "LGU", "SELESAI", "SELESAI"),
        ("2", "Gambar CAD (STEP/OBJ) diserahkan ke vendor sebagai acuan fabrikasi", "LGU", "SELESAI", "SELESAI"),
        ("3", "Dokumen klasifikasi area tertulis resmi (perimeter SRU)", "RU IV", "Dalam proses", "BELUM"),
        ("4", "Penetapan titik pasang detail per unit, rute kabel & tinggi kerja", "RU IV", "Survey ada, detail belum dikunci", "BELUM"),
        ("5", "Verifikasi lapangan & shop drawing titik pasang", "Vendor", "Menunggu penugasan", "BELUM"),
        ("6", "Konfirmasi orientasi bracket (tiang vertikal vs handrail horizontal)", "Vendor", "Belum dikonfirmasi", "BELUM"),
        ("2. DOKUMEN, IZIN & KESELAMATAN KERJA",),
        ("7", "Draft JSA/HSE sebagai acuan awal (belum menggantikan TRA/JSA resmi)", "LGU", "Tersedia (draft)", "SELESAI*"),
        ("8", "Checklist kesiapan instalasi lengkap sebagai referensi kerja bersama", "LGU", "Tersedia", "SELESAI"),
        ("9", "Pengesahan TRA/JSA spesifik titik pasang", "RU IV", "Belum disahkan", "BELUM"),
        ("10", "Izin kerja, izin masuk personel & surat masuk barang", "RU IV", "Wajib sebelum mobilisasi", "BELUM"),
        ("11", "Penugasan resmi vendor pelaksana & konfirmasi kompetensi kerja tinggi", "RU IV", "Perlu penugasan", "BELUM"),
        ("12", "Daftar kontak PIC (operasi/HSE/control room/IT) & jalur stop-work", "RU IV", "Perlu diisi", "BELUM"),
        ("13", "Personel bersertifikasi kerja di ketinggian & metode kerja aman", "Vendor", "Menunggu penugasan resmi", "BELUM"),
        ("14", "Toolbox meeting, APD, LOTO, gas test, koordinasi control room", "Vendor", "Perlu rencana kerja", "BELUM"),
        ("15", "Partisipasi dalam pengesahan TRA/JSA sebagai pelaksana fisik", "Vendor", "Menunggu", "BELUM"),
        ("3. PERANGKAT, MEKANIK & BRACKET",),
        ("16", "Unit GLD (3 RU IV + 1 cadangan), firmware & AI on-device", "LGU", "Siap, verifikasi pra-mobilisasi", "SELESAI*"),
        ("17", "Spesifikasi bracket/U-bolt/fastener sudah diberikan ke vendor", "LGU", "Selesai", "SELESAI"),
        ("18", "Supervisi teknis/QA selama pemasangan fisik", "LGU", "Siap dilaksanakan", "SELESAI"),
        ("19", "Menyetujui metode mounting non-invasive pada struktur existing", "RU IV", "Menunggu konfirmasi", "BELUM"),
        ("20", "Menyediakan akses ke struktur existing titik pasang", "RU IV", "Tergantung survey detail", "BELUM"),
        ("21", "Fabrikasi & pemasangan fisik bracket", "Vendor", "Menunggu penugasan", "BELUM"),
        ("22", "Menyediakan alat kerja, kunci, dan fastener sesuai spesifikasi", "Vendor", "Detail fastener terbuka", "BELUM"),
        ("23", "Melaksanakan pemasangan fisik GLD & Cluster Head", "Vendor", "Belum dapat dimulai", "BELUM"),
        ("4. CATU DAYA & KABEL",),
        ("24", "Spesifikasi catu daya 24VDC ≥ 1A per unit GLD sudah ditetapkan", "LGU", "Selesai", "SELESAI"),
        ("25", "Terminasi elektrikal spesifik-GLD, energize, dan commissioning", "LGU", "Siap dilaksanakan", "SELESAI"),
        ("26", "Menyediakan titik catu daya 24VDC di lokasi pasang", "RU IV", "Perlu konfirmasi titik", "BELUM"),
        ("27", "Menyediakan jalur kabel dari sumber daya ke titik pasang", "RU IV", "Menunggu konfirmasi", "BELUM"),
        ("28", "Instalasi fisik jalur kabel, terminasi mekanik & identifikasi kabel", "Vendor", "Menunggu penugasan", "BELUM"),
        ("29", "Menerapkan LOTO selama pekerjaan instalasi kabel & catu daya", "Vendor", "Menunggu rencana kerja", "BELUM"),
        ("5. JARINGAN & SERVER",),
        ("30", "Setup PC server, MQTT broker, adapter MQTT→backend, dan dashboard", "LGU", "Siap dikonfigurasi", "SELESAI"),
        ("31", "Router LGU sendiri di lokasi — jaringan field berjalan independen", "LGU", "Sudah dirancang", "SELESAI"),
        ("32", "Uji lokal & acceptance test sebelum serah terima", "LGU", "Checklist hari-H disiapkan", "SELESAI"),
        ("33", "Menyediakan PC dedicated fisik dekat Gateway", "RU IV", "Menunggu pengadaan", "BELUM"),
        ("34", "Review keamanan IT bila dashboard dibuka ke intranet kantor (NIC kedua)", "RU IV", "Review formal terbuka", "BELUM"),
        ("6. KOORDINASI & COMMISSIONING HARI-H",),
        ("35", "Memimpin energize & uji fungsional bersama HSE/control room", "LGU", "Siap dilaksanakan", "SELESAI"),
        ("36", "Menyusun kriteria acceptance test (telemetri, alarm/clear, rollback)", "LGU", "Sudah disusun", "SELESAI"),
        ("37", "Menyerahkan log commissioning, konfigurasi akhir & foto titik pasang", "LGU", "Siap dilaksanakan", "SELESAI"),
        ("38", "Menyediakan PIC operasi/HSE/control room saat hari-H", "RU IV", "Perlu diisi", "BELUM"),
        ("39", "Menyetujui rencana energize & menandatangani berita acara commissioning", "RU IV", "Menunggu jadwal", "BELUM"),
        ("40", "Hadir sebagai pelaksana fisik hari-H, dikoordinasikan LGU/HSE", "Vendor", "Menunggu penugasan", "BELUM"),
    ]
    section_rows = {i for i, row in enumerate(ITEMS) if len(row) == 1}
    add_table(
        doc,
        ["No", "Item", "Pihak", "Status", "Ket."],
        ITEMS,
        [0.35, 3.55, 0.65, 1.7, 0.65],
        font_size=8.1,
        mark_col=4,
        section_rows=section_rows,
    )
    add_paragraph(
        doc,
        "* Item 7 (draft JSA) dan 16 (unit GLD) ditandai selesai dari sisi LGU, namun masing-masing masih "
        "menyisakan syarat di pihak lain: Item 7 masih berstatus draft, menunggu pengesahan resmi RU IV "
        "(Item 9); Item 16 masih memerlukan verifikasi pra-mobilisasi (cek serial number, firmware, kondisi "
        "fisik) sebelum keberangkatan.",
        size=8.4, color=MID,
    )

    doc.add_page_break()

    # 5 Langkah berikutnya
    add_heading(doc, "5  Langkah Berikutnya", 1)
    add_paragraph(
        doc,
        "Pengesahan TRA/JSA spesifik titik pasang, penugasan resmi vendor instalatur oleh Pertamina RU IV, "
        "dan penyediaan infrastruktur di lokasi (titik catu daya, PC server, akses struktur) — ketiganya "
        "berada di sisi Pertamina RU IV dan vendor. Dari sisi LGU, seluruh item persiapan yang menjadi "
        "tanggung jawab kami sudah selesai. Kami mengusulkan meeting online bersama vendor instalatur dan "
        "RU IV untuk menyepakati jadwal dan menutup item-item ini bersama."
    )

    add_heading(doc, "6  Referensi Dokumen", 1)
    refs = [
        "Pembagian Persiapan Instalasi RU IV Cilacap, 17 September 2026 (Rev 1.3).",
        "Checklist Kesiapan Instalasi RU IV Cilacap, 10 September 2026 (Rev 1.0).",
        "Daftar Persiapan Vendor Instalasi RU IV Cilacap, 17 September 2026.",
        "Draf Permintaan Penyediaan Material Instalasi RU IV Cilacap, 17 September 2026.",
    ]
    for i, ref in enumerate(refs, 1):
        add_paragraph(doc, f"{i}. {ref}", size=8.8)

    add_paragraph(
        doc,
        "Dokumen ini bukan izin kerja, sertifikat, atau pengganti persetujuan HSE/operasi Pertamina. "
        "Instalasi permanen di area berbahaya tetap menunggu gate keselamatan dan sertifikasi terkait.",
        size=8.2, color=MID, space_after=0,
    )

    doc.core_properties.title = "Kurva-S Persiapan Instalasi RU IV Cilacap"
    doc.core_properties.subject = "Status kesiapan pra-mobilisasi instalasi GLD RU IV Cilacap"
    doc.core_properties.author = "LAPI Ganesha Utama bersama Lab IoT/Instrumentation and Computation ITB"
    doc.core_properties.keywords = "GLD, RU IV Cilacap, Instalasi, Kurva-S, Persiapan"
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_document()
