# -*- coding: utf-8 -*-
"""Lampiran tangkapan layar dashboard RU (aplikasi GLD) untuk dimasukkan ke laporan.

Sumber gambar: Deliverables/Demo_Dashboard_GLD_DigitalizationDay/screenshots/semua_halaman/terang/
(tema terang dipakai karena lebih jelas saat dicetak).

Output: Deliverables/Lampiran_Tangkapan_Layar_Dashboard_GLD.docx
PDF: python3 scripts/docx_to_pdf_with_toc.py <docx>
"""
import os

from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from build_persiapan_instalasi_corporate_docx import (
    make_doc, set_cell_shading, add_field, add_picture_fit,
    NAVY, GRAY, WHITE, HEAD_SHADE,
)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(REPO, "Deliverables", "Demo_Dashboard_GLD_DigitalizationDay", "screenshots", "semua_halaman", "terang")
OUT = os.path.join(REPO, "Deliverables", "Lampiran_Tangkapan_Layar_Dashboard_GLD.docx")

DOC_NO = "LGU/GLD/LAMP-DASH/2026-001"
REVISION = "1.0"
DOC_DATE = "19 Oktober 2026"

# (file, judul, keterangan)
PAGES = [
    ("00-login.png", "Halaman Login",
     "Akses dashboard dibatasi dengan akun per RU. Tampilan memakai identitas PT Pertamina Patra Niaga dan "
     "dapat dipindah antara tema terang dan gelap."),
    ("02-overview.png", "Overview",
     "Ringkasan kondisi RU: jumlah perangkat online/offline, alarm aktif, rata-rata baterai, status hasil AI "
     "tiap sensor GLD, dan linimasa alarm terbaru. Saat ada alarm, banner merah di bagian atas menampilkan "
     "sensor, jenis gas, tingkat keyakinan, area, koordinat, dan jalur Cluster Head yang dilalui."),
    ("01-map-view.png", "Map View — posisi perangkat",
     "Posisi Gateway, Cluster Head, dan sensor GLD di area RU. Sensor yang alarm ditandai merah dengan "
     "keterangan jenis gas dan area. Panel kanan menampilkan tren deteksi gas 24 jam per sensor."),
    ("01-map-view-topologi.png", "Map View — topologi jaringan",
     "Struktur jaringan LoRa mesh dari Gateway ke Cluster Head dan sensor. Sambungan yang terputus atau "
     "sensor yang alarm ditandai merah."),
    ("03-devices.png", "Devices",
     "Daftar seluruh perangkat (Gateway, Cluster Head, sensor GLD) dengan area pemasangan, hasil gas "
     "terakhir, baterai, sinyal, kesehatan, dan status online. Perangkat dapat didaftarkan dan diubah dari "
     "halaman ini."),
    ("04-unit-layout.png", "Unit Layout",
     "Sensor GLD dikelompokkan per Cluster Head beserta status (Safe / Leak Detected), jenis gas, dan "
     "tingkat keyakinan hasil AI."),
    ("05-alerts.png", "Alerts — konsol alarm",
     "Daftar alarm aktif dan riwayat alarm yang sudah ditangani. Setiap alarm memuat lokasi (area, "
     "koordinat, Cluster Head), peta skematis, serta tombol Peta, Acknowledge, dan Resolve. Tindakan "
     "operator tercatat beserta nama akunnya."),
    ("06-events.png", "Events — log kejadian",
     "Catatan seluruh kejadian sistem: alarm GLD, pelampauan ambang, perubahan status perangkat, "
     "keamanan data, dan tindakan operator. Dapat difilter per jenis/status dan diekspor ke CSV atau PDF."),
    ("07-analytics.png", "Analytics",
     "Tren keyakinan deteksi gas per jam, jumlah pembacaan berisiko, heatmap deteksi 7 hari per jam, "
     "sensor dengan risiko tertinggi, serta kondisi armada (online/offline, sebaran baterai, kualitas "
     "jaringan). Periode dapat dipilih 24 jam atau 7 hari."),
    ("08-settings.png", "Settings",
     "Pengaturan ambang risiko (MIDDLE ≥ 0,70, HIGH ≥ 0,80), laju penyegaran data, dan manajemen akun "
     "pengguna per RU."),
]

doc, sec = make_doc()


def p(text="", size=10.2, bold=False, italic=False, color=None, space_after=7, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    if text:
        r = para.add_run(text)
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return para


def bullet(text, size=10):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_after = Pt(3)
    r = para.add_run(text)
    r.font.size = Pt(size)


# ---------------------------------------------------------------- letterhead
lt = doc.add_table(rows=1, cols=1)
lc = lt.rows[0].cells[0]
set_cell_shading(lc, "1B2A4A")
lc.paragraphs[0].paragraph_format.space_after = Pt(2)
r = lc.paragraphs[0].add_run("PT LAPI GANESHA UTAMA")
r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = WHITE
lp = lc.add_paragraph()
lp.paragraph_format.space_after = Pt(3)
r = lp.add_run("Bekerja sama dengan Lab IoT & Fisika Institut Teknologi Bandung")
r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xC7, 0xD2, 0xE0)
doc.add_paragraph().paragraph_format.space_after = Pt(14)

p("LAMPIRAN", size=10, bold=True, color=GRAY, space_after=4)
title = doc.add_heading(level=0)
title.paragraph_format.space_after = Pt(4)
tr = title.add_run("Tangkapan Layar Dashboard Pemantauan GLD")
tr.font.size = Pt(22); tr.font.bold = True; tr.font.color.rgb = NAVY
p("Aplikasi dashboard RU — Gas Leak Detection (GLD) Tahap 2", size=12, color=GRAY, space_after=16)

mt = doc.add_table(rows=0, cols=2)
mt.style = "Table Grid"
for k, v in [
    ("Nomor dokumen", DOC_NO),
    ("Revisi", REVISION),
    ("Tanggal", DOC_DATE),
    ("Proyek", "Gas Leak Detection (GLD) Tahap 2 — PT Pertamina Patra Niaga (pilot RU IV Cilacap)"),
    ("Isi", "Tangkapan layar setiap halaman pada menu navigasi dashboard RU"),
]:
    row = mt.add_row().cells
    set_cell_shading(row[0], HEAD_SHADE)
    row[0].paragraphs[0].paragraph_format.space_after = Pt(2)
    r0 = row[0].paragraphs[0].add_run(k)
    r0.font.bold = True; r0.font.size = Pt(9.3); r0.font.color.rgb = GRAY
    row[1].paragraphs[0].paragraph_format.space_after = Pt(2)
    r1 = row[1].paragraphs[0].add_run(v)
    r1.font.size = Pt(9.6)
    row[0].width = Inches(1.9); row[1].width = Inches(4.6)
doc.add_paragraph().paragraph_format.space_after = Pt(10)

p("Keterangan", size=11, bold=True, color=NAVY, space_after=4)
for t in [
    "Tangkapan layar diambil dari aplikasi dashboard RU yang berjalan dengan DATA SIMULASI (label "
    "“DEMO — DATA SIMULASI” tampil di setiap halaman), bukan data lapangan. Skenario yang ditampilkan: "
    "alarm LPG pada sensor RU4 GLD 2.",
    "Posisi perangkat memakai koordinat contoh; koordinat aktual RU IV Cilacap diisi setelah instalasi.",
    "Latar peta (peta jalan/citra satelit) tidak ikut tertangkap karena pengambilan gambar dilakukan tanpa "
    "akses ke server peta; saat dashboard dipakai dengan koneksi internet atau tile peta lokal, latar peta "
    "tampil di belakang posisi perangkat.",
    "Nilai sinyal −100 dBm pada Gateway dan sensor GLD adalah nilai pengganti karena data sinyal dari "
    "perangkat tersebut belum dikirim ke server, bukan hasil pengukuran.",
]:
    bullet(t, size=9.8)

p("Daftar Gambar", size=11, bold=True, color=NAVY, space_after=4)
for i, (_, ttl, _) in enumerate(PAGES, 1):
    p(f"Gambar {i}. {ttl}", size=10, space_after=2)

# ---------------------------------------------------------------- header / footer
header = sec.header
header.is_linked_to_previous = False
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("Lampiran — Tangkapan Layar Dashboard Pemantauan GLD")
hr.font.size = Pt(8); hr.font.color.rgb = GRAY; hr.font.italic = True
fp = sec.footer.paragraphs[0]
fr = fp.add_run(f"{DOC_NO}  ·  Rev. {REVISION}  ·  {DOC_DATE}")
fr.font.size = Pt(8); fr.font.color.rgb = GRAY
fp.add_run("  ·  Halaman ")
add_field(fp, "PAGE", "1")
fp.add_run(" dari ")
add_field(fp, "NUMPAGES", "1")
for run in fp.runs:
    run.font.size = Pt(8); run.font.color.rgb = GRAY

# ---------------------------------------------------------------- one page per screenshot
for i, (fname, ttl, desc) in enumerate(PAGES, 1):
    doc.add_page_break()
    doc.add_heading(f"Gambar {i}. {ttl}", level=1)
    p(desc, size=10.2, space_after=8)
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_picture_fit(pic, os.path.join(SHOTS, fname), 6.5, 7.6)

doc.save(OUT)
print(OUT)
