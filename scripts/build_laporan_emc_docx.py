# -*- coding: utf-8 -*-
"""Laporan Teknis Pengujian EMC (pre-scan EN 55032 Kelas A) Node GLD & CH — rev 1.0 (format LGU).

Sumber: draf tim "Laporan_Teknis_Pengujian_EMC" (12 hlm, disimpan di
Sumber Dokumen/data mentah test EMC di telkom Test House okt 20206/Laporan_Teknis_Pengujian_EMC_draf_tim.pdf).
Isi bab 1-4 draf tim dipertahankan (teks, grafik, tabel titik ukur diparse dari PDF draf).
Ditambahkan: kontrol dokumen, ringkasan hasil, identitas EUT & kondisi uji, analisis temuan,
checklist EMC, rencana tindak lanjut, lembar pengesahan.

Run:  python3 scripts/build_laporan_emc_docx.py
PDF:  python3 scripts/docx_to_pdf_with_toc.py Deliverables/Laporan_Teknis_Pengujian_EMC_GLD_CH.docx
"""
import os
import re
import subprocess
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_spesifikasi_material_instalasi_docx import (  # noqa: E402
    make_doc, para, bullets, heading, table, fix_widths, set_cell_shading, banner, GRAY,
)
import build_laporan_uji_lab_docx as ul  # noqa: E402
import build_spesifikasi_material_instalasi_docx as sm  # noqa: E402
from docx.shared import Pt, Inches  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from PIL import Image  # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(REPO, "Sumber Dokumen", "data mentah test EMC di telkom Test House okt 20206",
                   "Laporan_Teknis_Pengujian_EMC_draf_tim.pdf")
ASSET = os.path.join(REPO, "scripts", "assets", "emc_telkom_okt2026")
OUT = os.path.join(REPO, "Deliverables", "Laporan_Teknis_Pengujian_EMC_GLD_CH.docx")
DOC_NO = "LGU/GLD/EMC/2026-001"
ul.REV = "1.0"
ul.DATE = "28 Oktober 2026"
sm.STATUS_STYLE.update({"Selesai": (sm.GREEN, "EAF4EC"), "Sebagian": (sm.AMBER, "FFF3DC"),
                        "Belum": (sm.RED, "FBE6E4")})

FREQ = re.compile(r"^\s*(\d+,\d+)\s+(kHz|MHz|GHz)\s+(.*?)\s+(Di (?:bawah|atas) batas)\s*$")


def parse_tables():
    """Ambil 12 tabel titik ukur dari PDF draf tim -> {nomor_tabel: [baris]}."""
    txt = subprocess.run(["pdftotext", "-layout", SRC, "-"], capture_output=True, text=True,
                         check=True).stdout
    tables, cur = {}, None
    for line in txt.splitlines():
        m = re.match(r"\s*Tabel (\d+) Hasil ", line)
        if m:
            cur = int(m.group(1)); tables[cur] = []
            continue
        f = FREQ.match(line)
        if f and cur:
            vals = f.group(3).split()
            tables[cur].append([f"{f.group(1)} {f.group(2)}"] + vals + [f.group(4)])
    assert len(tables) == 12 and all(len(v) == 18 for v in tables.values()), \
        {k: len(v) for k, v in tables.items()}
    return tables


def fig(doc, name, caption, max_w=6.3, max_h=3.3):
    path = os.path.join(ASSET, name)
    w, h = Image.open(path).size
    scale = min(max_w / w, max_h / h)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(w * scale))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(8)
    r = c.add_run(caption)
    r.font.size = Pt(8.8); r.font.italic = True; r.font.color.rgb = GRAY


def caption(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.bold = True; r.font.size = Pt(9.4)


def points(doc, n, title, rows, radiated):
    caption(doc, f"Tabel {n}. {title}")
    if radiated:
        head = ["Frekuensi", "Level (dBµV/m)", "Batas (dBµV/m)", "Margin (dB)", "Pol.", "Tinggi (cm)",
                "Sudut (°)", "Interpretasi"]
        widths = [1.05, 0.85, 0.85, 0.75, 0.45, 0.65, 0.6, 1.0]
    else:
        head = ["Frekuensi", "Level (dBµV)", "Batas (dBµV)", "Margin (dB)", "Koreksi (dB)", "Interpretasi"]
        widths = [1.2, 1.0, 1.0, 1.0, 1.0, 1.2]
    body = []
    for r in rows:
        mark = "Di atas batas*" if r[-1] == "Di atas batas" else r[-1]
        body.append(r[:-1] + [mark])
    t = table(doc, head, body, widths, size=8.2)
    from docx.oxml import OxmlElement
    trPr = t.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:tblHeader"))  # header tabel diulang bila tabel pindah halaman


def build():
    T = parse_tables()
    doc, sec = make_doc()
    ul.front(
        doc, sec,
        "Laporan Teknis Pengujian EMC Node GLD dan Cluster Head",
        "Pre-scan emisi elektromagnetik EN 55032 Kelas A di Telkom Test House — hasil, analisis temuan, "
        "checklist kesiapan EMC, dan rencana tindak lanjut",
        DOC_NO, "Laporan Teknis Pengujian EMC GLD & CH",
        [("Perangkat uji", "Node sensor GLD V2 (LoRa 920 MHz, catu 24 VDC) dan Cluster Head besar "
                           "(LoRa 921 MHz, catu baterai)"),
         ("Tanggal pengujian", "15 Oktober 2026 (tanggal pada grafik keluaran perangkat lunak ukur)"),
         ("Laboratorium", "Telkom Test House — ruang anechoic (radiasi) dan meja uji konduksi"),
         ("Sifat pengujian", "Pre-scan / pre-compliance. Bukan laporan uji resmi laboratorium terakreditasi "
                             "dan bukan sertifikat")],
    )
    # override lokasi uji bawaan front(): tulis ulang baris lokasi pada tabel kontrol dokumen
    for t in doc.tables:
        for row in t.rows:
            if row.cells[0].text == "Lokasi uji":
                row.cells[1].paragraphs[0].runs[0].text = ("Telkom Test House (pengujian EMC); analisis oleh "
                                                           "PT LAPI Ganesha Utama & Lab IoT ITB")
            if row.cells[0].text == "Dokumen terkait":
                row.cells[1].paragraphs[0].runs[0].text = (
                    "Draf tim \"Laporan Pengujian Kompatibilitas Elektromagnetik (EMC) Node GLD dan CH\" (12 hlm); "
                    "data mentah keluaran software ukur Telkom Test House (7 berkas); Checklist & Rencana Uji "
                    "Sertifikasi GLD (EMC · IP · RF · Monitor ATEX), berkas Excel pendamping")

    # ------------------------------------------------------------ riwayat revisi
    heading(doc, "Riwayat revisi")
    table(doc, ["Rev.", "Tanggal", "Uraian"], [
        ["Draf", "Okt 2026", "Draf tim: hasil emisi konduksi & radiasi GLD dan CH (bab 1–4 laporan ini)"],
        ["1.0", "28 Okt 2026", "Format laporan LGU; ditambah ringkasan hasil, identitas perangkat & kondisi uji, "
                               "analisis temuan (bab 6), checklist EMC (bab 7), rencana tindak lanjut (bab 8), "
                               "lembar pengesahan. Angka hasil ukur tidak diubah."],
    ], [0.6, 1.1, 5.2], size=9)

    # ------------------------------------------------------------ ringkasan
    heading(doc, "Ringkasan hasil")
    para(doc, "Pengujian ini adalah pre-scan emisi terhadap batas EN 55032 Kelas A (peralatan lingkungan industri). "
              "Dari 216 titik ukur yang dipilih instrumen, 212 titik berada di bawah batas. Empat titik yang "
              "melampaui batas seluruhnya berada pada frekuensi kerja LoRa (920 MHz GLD, 921 MHz CH), yaitu sinyal "
              "pancar yang memang disengaja, bukan emisi liar. Penilaian sinyal pancar dilakukan dengan standar "
              "radio (lihat bab 6), sehingga perlu konfirmasi laboratorium bahwa pita frekuensi kerja dikecualikan "
              "dari penilaian EN 55032.")
    table(doc, ["Perangkat", "Uji", "Detektor", "Titik", "Margin terkecil (di luar pita LoRa)", "Status pre-scan"], [
        ["GLD", "Konduksi AC — L1", "QP / AV", "18 / 18", "6,03 dB @ 323,9 kHz (QP)", "Di bawah batas"],
        ["GLD", "Konduksi AC — L2", "QP / AV", "18 / 18", "6,25 dB @ 323,9 kHz (QP)", "Di bawah batas"],
        ["GLD", "Radiasi 30 MHz–1 GHz", "QP", "18", "0,51 dB @ 117,9 MHz (H, 300 cm)", "Di bawah batas — margin kritis"],
        ["GLD", "Radiasi 1–6 GHz", "Peak / AV", "18 / 18", "33,97 dB @ 1,8129 GHz (AV)", "Di bawah batas"],
        ["CH besar", "Konduksi port DC", "QP / AV", "18 / 18", "42,41 dB @ 150 kHz (QP)", "Di bawah batas"],
        ["CH besar", "Radiasi 30 MHz–1 GHz", "QP", "18", "3,28 dB @ 978,9 MHz", "Di bawah batas — perlu perhatian"],
        ["CH besar", "Radiasi 1–6 GHz", "Peak / AV", "18 / 18", "5,56 dB @ 1,8425 GHz (AV, harmonik ke-2)",
         "Di bawah batas — perlu perhatian"],
    ], [0.75, 1.35, 0.75, 0.6, 2.05, 1.4], size=8.6)
    para(doc, "Titik di atas batas (frekuensi kerja LoRa): GLD 919,654 / 920,045 / 920,387 MHz (−1,56 / −36,29 / "
              "−1,51 dB) dan CH 920,973 MHz (−55,47 dB).", size=9.4, color=GRAY)
    bullets(doc, [
        "Hasil pre-scan positif: emisi konduksi dan emisi radiasi di luar pita LoRa berada di bawah batas Kelas A.",
        "Risiko utama: margin GLD hanya 0,51–0,95 dB di 117–127 MHz. Margin ini lebih kecil dari ketidakpastian "
        "ukur radiasi yang lazim (±4–5 dB), sehingga pada uji resmi titik ini dapat melampaui batas. Mitigasi "
        "disarankan sebelum uji resmi.",
        "Cakupan belum lengkap: uji imunitas, port catu 24 VDC GLD, Gateway, CH kecil, dan uji standar radio "
        "belum dilakukan. Daftar lengkap ada di checklist bab 7.",
    ])
    doc.add_page_break()

    # ------------------------------------------------------------ bab 1 (teks tim)
    heading(doc, "1. Pengujian EMC")
    para(doc, "Kompatibilitas elektromagnetik atau electromagnetic compatibility (EMC) adalah kemampuan suatu perangkat "
              "untuk beroperasi tanpa menimbulkan gangguan elektromagnetik yang berlebihan terhadap perangkat lain di "
              "sekitarnya. Pengujian EMC perlu dilakukan karena rangkaian digital, catu daya, dan pemancar radio pada "
              "node menghasilkan sinyal berfrekuensi tinggi yang dapat merambat melalui kabel atau terpancar ke udara. "
              "Pengujian ini mengukur besar emisi tersebut dan membandingkannya dengan batas EN 55032 Kelas A, yaitu "
              "batas emisi bagi peralatan yang digunakan di lingkungan industri.")
    para(doc, "Parameter yang diuji adalah emisi konduksi dan emisi radiasi. Emisi konduksi adalah tegangan gangguan "
              "radiofrekuensi yang merambat melalui kabel catu daya, diukur pada 150 kHz sampai 30 MHz dalam dBµV. "
              "Emisi radiasi adalah kuat medan listrik yang dipancarkan perangkat ke udara, diukur pada 30 MHz sampai "
              "1 GHz serta 1 GHz sampai 6 GHz dalam dBµV/m. Hasil ukur dibandingkan dengan batas EN 55032 Kelas A, dan "
              "selisihnya disebut margin, yaitu batas dikurangi level terukur pada saat scanning. Margin positif berarti "
              "level berada di bawah batas, sedangkan margin negatif berarti level melampaui batas.")
    para(doc, "Pada pengujian konduksi, port daya perangkat dihubungkan ke analyzer, kemudian spektrum dipindai pada "
              "150 kHz sampai 30 MHz. Pada pengujian radiasi, perangkat diletakkan di atas meja putar di dalam ruang "
              "anechoic dan medan yang dipancarkan diterima oleh antena ukur. Polarisasi antena, tinggi antena, dan "
              "sudut meja putar diubah untuk memperoleh level maksimum pada setiap frekuensi.")
    para(doc, "Node GLD beroperasi pada frekuensi 920 MHz dan node CH beroperasi pada frekuensi 921 MHz. Pada pengujian "
              "konduksi node GLD disuplai oleh PSU 24 V sehingga emisi konduksi diukur pada kedua konduktor masukan "
              "daya PSU, yaitu L1 dan L2 (sisi AC), sedangkan node CH disuplai oleh baterai sehingga emisi konduksi "
              "diukur pada port catu daya arus searah (DC).")

    heading(doc, "1.1 Identitas perangkat uji dan kondisi pengujian", level=2)
    table(doc, ["Parameter", "Node sensor GLD", "Cluster Head (besar)"], [
        ["Versi", "GLD V2 (prototipe)", "CH V2 (prototipe)"],
        ["Mikrokontroler / modul radio", "ESP32-S3-WROOM-1U / EBYTE E22-900MM22S",
         "ESP32-S3-WROOM-1U / EBYTE E22-900MM22S"],
        ["Frekuensi kerja LoRa", "920 MHz (star)", "920 MHz (star) & 921 MHz (mesh); terukur 921 MHz"],
        ["Daya pancar modul maks. / antena", "22 dBm / omni 3 dBi", "22 dBm / omni 3 dBi (star), 8 dBi (mesh)"],
        ["Catu daya saat uji", "PSU 220 VAC → 24 VDC (konduksi diukur di sisi AC)", "Baterai (konduksi di port DC)"],
        ["Konsumsi daya maks.", "7,995 W @ 24 VDC", "0,73 W"],
        ["Nomor seri, versi firmware, mode operasi, laju kirim LoRa, panjang & tata letak kabel, model PSU",
         "Tidak tercatat di draf — wajib dicatat pada uji berikutnya",
         "Tidak tercatat di draf — wajib dicatat pada uji berikutnya"],
    ], [2.0, 2.45, 2.45], size=8.8)
    para(doc, "Sumber data perangkat: lembar parameter spesifikasi EMC tim. Jarak ukur radiasi 3 m diturunkan dari "
              "nilai batas yang dipakai software ukur (50/57 dBµV/m), belum tertulis eksplisit di draf.",
         size=9, color=GRAY)

    heading(doc, "1.2 Batas yang digunakan", level=2)
    table(doc, ["Uji", "Rentang", "Batas Kelas A (sesuai software ukur)"], [
        ["Konduksi port AC (GLD)", "0,15–0,5 MHz / 0,5–30 MHz", "QP 79 / 73 dBµV; AV 66 / 60 dBµV"],
        ["Konduksi port DC (CH)", "0,15–0,5 MHz / 0,5–30 MHz", "QP 97→87 / 87 dBµV; AV 84→74 / 74 dBµV"],
        ["Radiasi (3 m)", "30–230 MHz / 230–1000 MHz", "QP 50 / 57 dBµV/m"],
        ["Radiasi (3 m)", "1–3 GHz / 3–6 GHz", "Peak 76 / 80 dBµV/m; AV 56 / 60 dBµV/m"],
    ], [1.8, 2.0, 3.1], size=8.8)
    doc.add_page_break()

    # ------------------------------------------------------------ bab 2 GLD
    heading(doc, "2. Hasil Pengujian Node GLD")
    heading(doc, "2.1 Emisi konduksi jalur L1", level=2)
    fig(doc, "fig1.png", "Gambar 1. Grafik emisi konduksi node GLD pada jalur L1")
    para(doc, "Pada Gambar 1, sumbu horizontal menyatakan frekuensi dan sumbu vertikal menyatakan level gangguan dalam "
              "dBµV. Kurva merah adalah hasil detektor Quasi Peak, kurva kuning adalah hasil detektor Average, dan garis "
              "horizontal adalah batas EN 55032 Kelas A untuk masing-masing detektor. Kedua kurva berada di bawah batas, "
              "artinya gangguan radiofrekuensi yang tersalur melalui kabel catu daya tidak melampaui nilai ambang batas "
              "sehingga tidak berpotensi mengganggu perangkat lain yang terhubung pada jaringan daya yang sama.")
    points(doc, 1, "Hasil Quasi Peak emisi konduksi node GLD pada jalur L1", T[1], False)
    points(doc, 2, "Hasil Average emisi konduksi node GLD pada jalur L1", T[2], False)
    para(doc, "Tabel 1 dan Tabel 2 memuat nilai pada titik ukur yang dipilih instrumen: level menunjukkan amplitudo "
              "terukur, batas menunjukkan nilai acuan, margin menunjukkan selisih batas dan level, dan koreksi "
              "menunjukkan faktor koreksi sistem ukur. Seluruh margin bernilai positif, sehingga semua titik berada di "
              "bawah batas.")
    heading(doc, "2.2 Emisi konduksi jalur L2", level=2)
    fig(doc, "fig2.png", "Gambar 2. Grafik emisi konduksi node GLD pada jalur L2")
    para(doc, "Pada Gambar 2, sumbu, kurva, dan garis batas sama seperti Gambar 1. Kedua kurva berada di bawah batas, "
              "artinya gangguan radiofrekuensi yang tersalur melalui jalur L2 juga tidak melampaui nilai ambang batas, "
              "sehingga kedua jalur daya node GLD memenuhi batas emisi konduksi.")
    points(doc, 3, "Hasil Quasi Peak emisi konduksi node GLD pada jalur L2", T[3], False)
    points(doc, 4, "Hasil Average emisi konduksi node GLD pada jalur L2", T[4], False)
    para(doc, "Tabel 3 dan Tabel 4 memuat nilai Quasi Peak dan Average jalur L2 dengan susunan kolom yang sama seperti "
              "jalur L1. Seluruh margin bernilai positif, sehingga semua titik berada di bawah batas.")
    heading(doc, "2.3 Emisi radiasi 30 MHz sampai 1 GHz", level=2)
    fig(doc, "fig3.png", "Gambar 3. Grafik emisi radiasi node GLD 30 MHz sampai 1 GHz")
    para(doc, "Pada Gambar 3, sumbu vertikal menyatakan kuat medan listrik dalam dBµV/m, kurva merah adalah hasil "
              "detektor Quasi Peak, dan garis berjenjang adalah batas EN 55032 Kelas A. Puncak sempit pada sekitar "
              "920 MHz melampaui batas karena merupakan frekuensi operasi node GLD, yaitu sinyal yang dipancarkan "
              "pemancar radio node untuk komunikasi dan bukan gangguan yang tidak diinginkan. Di luar pita tersebut, "
              "emisi radiasi node GLD berada di bawah batas, namun tiga titik pada 117–127 MHz hanya bermargin "
              "0,51–0,95 dB (dibahas pada bab 6).")
    points(doc, 5, "Hasil Quasi Peak emisi radiasi node GLD 30 MHz sampai 1 GHz", T[5], True)
    para(doc, "Tabel 5 memuat kuat medan, batas, margin, dan konfigurasi pengukuran saat level maksimum diperoleh: "
              "polarisasi menunjukkan orientasi antena (H horizontal, V vertikal), tinggi menunjukkan posisi vertikal "
              "antena, dan sudut menunjukkan posisi putar meja uji. Tiga titik di sekitar 920 MHz bermargin negatif "
              "karena merupakan sinyal pancar node (ditandai *), sedangkan titik lainnya bermargin positif.")
    heading(doc, "2.4 Emisi radiasi 1 GHz sampai 6 GHz", level=2)
    fig(doc, "fig4.png", "Gambar 4. Grafik emisi radiasi node GLD 1 GHz sampai 6 GHz")
    para(doc, "Pada Gambar 4, kurva merah adalah hasil detektor Max Peak, kurva kuning adalah hasil detektor Average, "
              "dan garis berjenjang adalah batas EN 55032 Kelas A untuk masing-masing detektor. Kedua kurva berada "
              "jauh di bawah batas, artinya emisi radiasi node GLD pada rentang ini tidak melampaui nilai ambang batas "
              "dan tidak ada komponen yang mendekati batas.")
    points(doc, 6, "Hasil Max Peak emisi radiasi node GLD 1 GHz sampai 6 GHz", T[6], True)
    points(doc, 7, "Hasil Average emisi radiasi node GLD 1 GHz sampai 6 GHz", T[7], True)
    para(doc, "Tabel 6 memuat titik Max Peak dan Tabel 7 memuat titik Average beserta konfigurasi antena saat level "
              "maksimum diperoleh. Seluruh margin bernilai positif, sehingga semua titik berada di bawah batas.")
    doc.add_page_break()

    # ------------------------------------------------------------ bab 3 CH
    heading(doc, "3. Hasil Pengujian Node CH")
    heading(doc, "3.1 Emisi konduksi port DC", level=2)
    fig(doc, "fig5.png", "Gambar 5. Grafik emisi konduksi node CH pada port DC")
    para(doc, "Pada Gambar 5, sumbu horizontal menyatakan frekuensi dan sumbu vertikal menyatakan level gangguan pada "
              "port DC dalam dBµV; kurva merah adalah hasil detektor Quasi Peak, kurva kuning adalah hasil detektor "
              "Average, dan garis horizontal adalah batas masing-masing detektor. Kedua kurva berada di bawah batas "
              "pada seluruh rentang ukur dengan selisih yang besar, artinya gangguan radiofrekuensi pada port catu daya "
              "DC node CH jauh di bawah nilai ambang batas sehingga tidak berpotensi mengganggu perangkat lain yang "
              "terhubung pada sumber daya yang sama.")
    points(doc, 8, "Hasil Quasi Peak emisi konduksi port DC node CH", T[8], False)
    points(doc, 9, "Hasil Average emisi konduksi port DC node CH", T[9], False)
    para(doc, "Tabel 8 dan Tabel 9 memuat nilai Quasi Peak dan Average pada titik ukur yang dipilih instrumen dengan "
              "susunan kolom yang sama seperti tabel node GLD. Seluruh margin bernilai positif, sehingga semua titik "
              "berada di bawah batas.")
    heading(doc, "3.2 Emisi radiasi 30 MHz sampai 1 GHz", level=2)
    fig(doc, "fig6.png", "Gambar 6. Grafik emisi radiasi node CH 30 MHz sampai 1 GHz")
    para(doc, "Gambar 6 menyajikan kuat medan dengan sumbu, detektor, dan batas yang sama seperti Gambar 3. Puncak "
              "sempit pada sekitar 921 MHz melampaui batas karena berimpit dengan frekuensi operasi node CH, yaitu "
              "sinyal yang sengaja dipancarkan pemancar radio node agar dapat diterima perangkat lain. Batas EN 55032 "
              "ditetapkan untuk membatasi gangguan yang tidak disengaja di luar frekuensi kerja node.")
    points(doc, 10, "Hasil Quasi Peak emisi radiasi node CH 30 MHz sampai 1 GHz", T[10], True)
    para(doc, "Tabel 10 memuat kuat medan, batas, margin, dan konfigurasi pengukuran saat level maksimum diperoleh. "
              "Hanya titik 920,973 MHz yang bermargin negatif karena merupakan sinyal pancar node (ditandai *), "
              "sedangkan titik lainnya bermargin positif. Titik 978,880 MHz bermargin 3,28 dB (dibahas pada bab 6).")
    heading(doc, "3.3 Emisi radiasi 1 GHz sampai 6 GHz", level=2)
    fig(doc, "fig7.png", "Gambar 7. Grafik emisi radiasi node CH 1 GHz sampai 6 GHz")
    para(doc, "Pada Gambar 7, kurva merah adalah hasil detektor Max Peak, kurva kuning adalah hasil detektor Average, "
              "dan garis berjenjang adalah batas untuk masing-masing detektor. Kedua kurva berada di bawah batas, "
              "tetapi puncak pada sekitar 1,84 GHz mendekati batas Average (margin 5,56 dB). Frekuensi ini hampir tepat "
              "dua kali frekuensi operasi node CH, yaitu harmonik kedua sinyal pancar yang timbul akibat "
              "ketidaklinearan penguat pemancar. Kondisi ini masih di bawah batas EN 55032 Kelas A.")
    points(doc, 11, "Hasil Max Peak emisi radiasi node CH 1 GHz sampai 6 GHz", T[11], True)
    points(doc, 12, "Hasil Average emisi radiasi node CH 1 GHz sampai 6 GHz", T[12], True)
    para(doc, "Tabel 11 memuat titik Max Peak dan Tabel 12 memuat titik Average beserta konfigurasi antena saat level "
              "maksimum diperoleh. Seluruh margin bernilai positif, sehingga semua titik berada di bawah batas.")
    para(doc, "* Titik pada frekuensi kerja LoRa (sinyal pancar yang disengaja).", size=9, color=GRAY)

    # ------------------------------------------------------------ bab 4 dokumentasi
    heading(doc, "4. Dokumentasi Pelaksanaan Pengujian")
    para(doc, "Dokumentasi berikut menunjukkan ruang anechoic, sistem penempatan antena, dan konfigurasi meja pengukuran "
              "emisi radiasi dan konduksi di Telkom Test House.")
    fig(doc, "fig8_anechoic.jpg", "Gambar 8. Ruang anechoic dan area penempatan perangkat uji radiasi", max_h=3.0)
    fig(doc, "fig9_antena.jpg", "Gambar 9. Sistem antena dan mekanisme positioning", max_h=2.6)
    fig(doc, "fig10_konduksi.jpg", "Gambar 10. Konfigurasi meja pengujian emisi konduksi", max_h=2.6)

    # ------------------------------------------------------------ bab 5 kesimpulan
    heading(doc, "5. Kesimpulan Hasil Pre-scan")
    bullets(doc, [
        "Emisi konduksi GLD (sisi AC PSU, L1 & L2) dan CH (port DC) berada di bawah batas EN 55032 Kelas A, dengan "
        "margin terkecil 6,03 dB (GLD, 323,9 kHz QP).",
        "Emisi radiasi GLD dan CH di luar pita LoRa berada di bawah batas. Margin GLD pada 117–127 MHz sangat kecil "
        "(0,51–0,95 dB) dan menjadi titik risiko utama untuk uji resmi.",
        "Empat titik di atas batas adalah sinyal pancar LoRa sendiri (920/921 MHz). Titik ini dinilai dengan standar "
        "radio, bukan EN 55032, dengan syarat dikonfirmasi oleh laboratorium.",
        "Hasil ini bersifat pre-scan. Untuk klaim kepatuhan EMC masih diperlukan uji imunitas, uji port 24 VDC GLD, "
        "dan laporan resmi dari laboratorium terakreditasi.",
    ])
    doc.add_page_break()

    # ------------------------------------------------------------ bab 6 analisis
    heading(doc, "6. Analisis Temuan dan Catatan Teknis")
    table(doc, ["No", "Temuan", "Dampak", "Rekomendasi"], [
        ["E-1", "GLD radiasi QP 117,40 / 117,89 / 126,82 MHz: 49,16 / 49,49 / 49,05 dBµV/m vs batas 50 dBµV/m "
                "(margin 0,84 / 0,51 / 0,95 dB), polarisasi H, antena 300 cm.",
         "Margin di bawah ketidakpastian ukur radiasi (±4–5 dB). Pada uji resmi berisiko melampaui batas.",
         "Cari sumber: harmonik clock/osilator, regulator switching, atau kabel (24 VDC, fan) yang bertindak sebagai "
         "antena. Uji ferit pada kabel daya & fan, filter input, perbaikan grounding enclosure. Ulang pre-scan; "
         "target margin ≥ 6 dB."],
        ["E-2", "Sinyal pancar LoRa melampaui batas EN 55032 (GLD 93,29 dBµV/m @ 920,045 MHz; CH 112,47 dBµV/m "
                "@ 920,973 MHz).",
         "Wajar untuk pemancar, tetapi harus dinyatakan resmi sebagai pita kerja yang dikecualikan.",
         "Minta laboratorium menerapkan exclusion band sesuai EN 301 489-1/-3, atau ulang ukur dengan pemancar dalam "
         "mode idle/receive. Kinerja pemancar dinilai pada uji radio (rencana RF)."],
        ["E-3", "CH harmonik ke-2 1,8425 GHz: AV 50,44 dBµV/m (margin 5,56 dB), Peak 52,21 dBµV/m. CH 978,88 MHz "
                "QP margin 3,28 dB.",
         "Indikasi EIRP harmonik ≈ −45 dBm (AV, konversi 3 m), di bawah batas spurious radio −30 dBm. Margin EMC "
         "masih tipis.",
         "Ukur spurious secara konduksi di port antena saat uji RF. Pertimbangkan filter low-pass/harmonik pada "
         "jalur antena CH."],
        ["E-4", "Konduksi GLD diukur pada sisi AC PSU (L1/L2), bukan pada port 24 VDC GLD.",
         "Hasil mewakili gabungan PSU + GLD. Batas produk GLD yang disertifikasi adalah masukan 24 VDC saja; PSU di "
         "lapangan disediakan Pertamina di junction box.",
         "Tambahkan uji konduksi port DC 24 VDC GLD dengan AAN/LISN DC. Catat merek & model PSU yang dipakai."],
        ["E-5", "Level fundamental GLD terukur jauh lebih rendah dari CH (93,3 vs 112,5 dBµV/m).",
         "Detektor QP tidak tepat untuk mengukur daya pancar LoRa (chirp, siklus kirim pendek). Angka ini tidak "
         "boleh dipakai sebagai ukuran daya pancar.",
         "Daya pancar & EIRP diukur pada uji RF (konduksi di port antena + gain antena)."],
        ["E-6", "Draf tidak mencatat nomor seri, versi firmware, mode operasi, laju kirim, tata letak kabel, daftar "
                "alat ukur & kalibrasi, ketidakpastian ukur, jarak ukur.",
         "Laporan belum dapat dipakai sebagai bukti formal; pengulangan uji sulit dibandingkan.",
         "Gunakan lembar data uji standar (checklist bab 7) pada setiap sesi uji."],
        ["E-7", "Belum diuji: imunitas (ESD, radiated RF, EFT, surge, conducted RF, dip/interupsi), Gateway, CH kecil.",
         "Kepatuhan EMC belum dapat diklaim. Lingkungan kilang memiliki gangguan tinggi (motor, inverter, switching).",
         "Rencanakan uji imunitas EN 61000-6-2 / EN 55035; untuk alat deteksi gas pertimbangkan EN 50270. Uji "
         "Gateway & CH kecil."],
    ], [0.45, 2.1, 2.05, 2.3], size=8.4)

    # ------------------------------------------------------------ bab 7 checklist
    heading(doc, "7. Checklist Kesiapan EMC")
    para(doc, "Status per 28 Oktober 2026. Versi yang dapat diperbarui ada di berkas Excel pendamping "
              "(sheet \"Checklist EMC\").", size=9.4, color=GRAY)
    CK = [
        ["EMC-01", "Pre-scan emisi radiasi GLD 30 MHz–1 GHz", "Selesai", "Margin kritis 0,51 dB (E-1)"],
        ["EMC-02", "Pre-scan emisi radiasi GLD 1–6 GHz", "Selesai", "Margin ≥ 33,97 dB"],
        ["EMC-03", "Pre-scan konduksi GLD sisi AC PSU (L1/L2)", "Selesai", "Margin ≥ 6,03 dB"],
        ["EMC-04", "Pre-scan radiasi CH besar 30 MHz–6 GHz", "Selesai", "Margin 3,28 dB & 5,56 dB (E-3)"],
        ["EMC-05", "Pre-scan konduksi port DC CH besar", "Selesai", "Margin ≥ 42,41 dB"],
        ["EMC-06", "Konduksi port 24 VDC GLD (AAN/LISN DC)", "Belum", "E-4"],
        ["EMC-07", "Mitigasi & re-scan GLD 117–127 MHz (target margin ≥ 6 dB)", "Belum", "E-1"],
        ["EMC-08", "Konfirmasi lab: exclusion band LoRa / ulang ukur mode idle", "Belum", "E-2"],
        ["EMC-09", "Pre-scan Gateway (5 VDC adaptor + Wi-Fi)", "Belum", "Perangkat di gedung"],
        ["EMC-10", "Pre-scan CH kecil", "Belum", "E-7"],
        ["EMC-11", "Imunitas ESD (IEC 61000-4-2)", "Belum", "E-7"],
        ["EMC-12", "Imunitas medan RF radiasi (IEC 61000-4-3)", "Sebagian", "Uji interferensi co-channel LoRa 17 Sep "
                                                                         "(informal), bukan IEC 61000-4-3"],
        ["EMC-13", "Imunitas EFT/burst (IEC 61000-4-4) & surge (IEC 61000-4-5) port 24 VDC", "Belum", "E-7"],
        ["EMC-14", "Imunitas RF konduksi (IEC 61000-4-6) & dip/interupsi DC (IEC 61000-4-29)", "Belum", "E-7"],
        ["EMC-15", "Kriteria kinerja saat uji imunitas (alarm tidak palsu, tidak hilang)", "Belum",
         "Ditetapkan sebelum uji"],
        ["EMC-16", "Identitas EUT tercatat (S/N, firmware, mode, kabel, PSU)", "Belum", "E-6"],
        ["EMC-17", "Pilih lab terakreditasi & standar uji resmi", "Belum", "Koordinasi dengan agen/lab"],
        ["EMC-18", "Laporan uji resmi laboratorium", "Belum", "—"],
    ]
    table(doc, ["ID", "Item", "Status", "Catatan"], CK, [0.75, 3.35, 0.8, 2.0], size=8.6, status_col=2)

    # ------------------------------------------------------------ bab 8 tindak lanjut
    heading(doc, "8. Rencana Tindak Lanjut")
    table(doc, ["Urutan", "Kegiatan", "PIC usulan"], [
        ["1", "Lengkapi identitas EUT & kondisi uji sesi 15 Oktober (S/N, firmware, mode, PSU) dari catatan tim",
         "Lab IoT ITB"],
        ["2", "Investigasi sumber emisi 117–127 MHz di lab (near-field probe); siapkan opsi ferit/filter/grounding",
         "Lab IoT ITB"],
        ["3", "Konfirmasi ke Telkom Test House perlakuan pita LoRa (exclusion band) dan jarak ukur", "LGU"],
        ["4", "Re-scan GLD (radiasi 30 MHz–1 GHz + konduksi port 24 VDC) setelah mitigasi", "Lab IoT ITB · LGU"],
        ["5", "Pre-scan Gateway dan CH kecil", "Lab IoT ITB"],
        ["6", "Susun rencana uji imunitas (standar, level uji, kriteria kinerja) dan pilih lab", "LGU"],
        ["7", "Gabungkan dengan uji radio (RF) agar satu sesi lab mencakup EN 301 489 & uji pemancar", "LGU"],
    ], [0.7, 4.7, 1.5], size=8.8)

    ul.signoff(doc, 9)
    doc.save(OUT)
    print("written", OUT)


if __name__ == "__main__":
    build()
