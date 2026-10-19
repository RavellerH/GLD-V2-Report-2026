"""Dokumen Pemenuhan Klausul Termin — SP-001/KPI43400/2026-S0 (Sertifikasi).

Permintaan Pak Tresnandi (19 Okt 2026): klausul termin dibuatkan dokumen pemenuhannya untuk
diajukan tanda tangan ke Pak Pindoan. Klausul:
  "Tersedia prototype GLD yang telah diuji dan diperbaiki desainnya. Dilampirkan foto
   perangkatnya, dan foto-foto waktu kami kunjungan ke ITB."

Output: Paket LGU/03_Pemenuhan_Klausul_Termin_SP-001_Sertifikasi/*.docx + .pdf
Foto kunjungan Pertamina ke ITB dibaca dari FOTO_KUNJUNGAN (bila folder ada); bila belum ada,
Lampiran 2 berisi bingkai kosong untuk ditempel foto.

Pakai: python scripts/build_pemenuhan_termin1_sertifikasi_sp001.py
"""
import glob
import io
import os
import subprocess

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageOps

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT_DIR = os.path.join(ROOT, "Paket LGU", "03_Pemenuhan_Klausul_Termin_SP-001_Sertifikasi")
NAMA = "Dokumen_Pemenuhan_Klausul_Termin_SP-001_Sertifikasi_GLD"
LOGO = os.path.join(HERE, "assets", "lgu_cover")
FOTO = os.path.join(HERE, "assets", "assembly_photos_0710")
FOTO_LAMA = os.path.join(ROOT, "Paket Pertamina", "09_Laporan_FAT_GLD_Tahap2", "Lampiran_06_Foto_Unit_GLD_Terakit.jpg")
PCB_ITERASI = os.path.join(HERE, "assets", "pcb_iterasi", "gld_pcb_v1_v5.png")
FOTO_KUNJUNGAN = os.path.join(ROOT, "Sumber Dokumen", "Foto_Kunjungan_Pertamina_Lab_IoT_ITB_6Agustus2026")
FONT = "Arial"
NAVY = RGBColor(0x26, 0x23, 0x21)

NO_SP = "SP-001/KPI43400/2026-S0"
NO_DOK = "LGU-GLD-SP001-T1-2026-001"
REV = "1.0"
TGL = "Oktober 2026"
TEMPAT = "Bandung, ........ Oktober 2026"
LGU_TTD = [("Team Leader,", "Dr. Maman Budiman"), ("Direktur Utama,", "Ir. Harry Fardiman")]
PPN_JAB = "Manager Domestic Product Content & Digitalization"
PPN_NAMA = "Agustinus Pindoan Panjaitan"


# ---------------------------------------------------------------- helpers
def run(p, text, size=10, bold=False, italic=False, underline=False, color=None):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.bold, r.italic, r.underline = bold, italic, underline
    if color is not None:
        r.font.color.rgb = color
    return r


def para(container, text="", size=10, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, before=0,
         italic=False, underline=False):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing = 1.15
    for i, line in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run(p, line, size, bold, italic, underline)
    return p


def rich(container, parts, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, style=None):
    """parts: list of (text, bold)."""
    p = container.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    for t, b in parts:
        run(p, t, size, b)
    return p


def heading(doc, text, size=12):
    p = para(doc, text, size, True, WD_ALIGN_PARAGRAPH.LEFT, after=6, before=10)
    p.paragraph_format.keep_with_next = True
    return p


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), hexcolor)
    tcPr.append(s)


def borders(tbl, sz="6"):
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{e}")
        x.set(qn("w:val"), "single"); x.set(qn("w:sz"), sz); x.set(qn("w:color"), "000000")
        b.append(x)
    tbl._tbl.tblPr.append(b)


def fixed(tbl, widths):
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
    tbl._tbl.tblPr.append(lay)
    tbl.autofit = False
    grid = tbl._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w.twips)))
    for row in tbl.rows:
        for c, w in zip(row.cells, widths):
            c.width = w


def ctext(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    for i, line in enumerate(text.split("\n")):
        if i:
            p.add_run().add_break()
        run(p, line, size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def img_stream(path, max_px=1400):
    im = ImageOps.exif_transpose(Image.open(path)).convert("RGB")
    im.thumbnail((max_px, max_px))
    b = io.BytesIO()
    im.save(b, "JPEG", quality=82)
    b.seek(0)
    return b


def foto_grid(doc, items, cols=2, w=Cm(7.0), h=Cm(4.4)):
    """items: list of (path_or_None, caption). Foto dipaskan dalam kotak w x h."""
    t = doc.add_table(rows=0, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed(t, [Cm(8.2)] * cols)
    for i in range(0, len(items), cols):
        cells = t.add_row().cells
        for c, (path, cap) in zip(cells, items[i:i + cols]):
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if path:
                im = ImageOps.exif_transpose(Image.open(path))
                ar = im.width / im.height
                if ar >= w / h:
                    p.add_run().add_picture(img_stream(path), width=w)
                else:
                    p.add_run().add_picture(img_stream(path), height=h)
            else:
                box = c.add_table(rows=1, cols=1)
                borders(box, "8")
                bc = box.cell(0, 0)
                bc.width = w
                box.rows[0].height = h
                ctext(bc, "[ tempel foto kunjungan ]", 9, align=WD_ALIGN_PARAGRAPH.CENTER)
            cp = c.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after = Pt(8)
            run(cp, cap, 8.5, italic=True)
    return t


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, f"{NO_DOK}  ·  Rev. {REV}  ·  Halaman ", 8)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        r = p.add_run()
        r.font.size = Pt(8); r.font.name = FONT
        if kind:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r._r.append(f)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._r.append(it)


# ---------------------------------------------------------------- isi
def kop(doc):
    t = doc.add_table(rows=2, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    fixed(t, [Cm(3.6), Cm(9.6), Cm(3.6)])
    c = t.cell(0, 0); c.merge(t.cell(1, 0))
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraphs[0].add_run().add_picture(os.path.join(LOGO, "logo_lgu.png"), height=Cm(1.3))
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    ctext(t.cell(0, 1), "DOKUMEN PEMENUHAN KETENTUAN TERMIN 1", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER)
    ctext(t.cell(1, 1), f"Program Sertifikasi Hazardous Area (ATEX/IECEx)\nGas Leak Detector (GLD) Tahap 2 — {NO_SP}",
          9, False, WD_ALIGN_PARAGRAPH.CENTER)
    c = t.cell(0, 2); c.merge(t.cell(1, 2))
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraphs[0].add_run().add_picture(os.path.join(LOGO, "logo_pertamina.png"), width=Cm(3.0))
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para(doc, "", after=2)

    info = [
        ("Nomor Surat Pesanan", f"{NO_SP} — Sertifikasi"),
        ("Nomor dokumen / Revisi", f"{NO_DOK} / Rev. {REV}"),
        ("Tanggal", TGL),
        ("Pelaksana", "PT LAPI Ganesha Utama, bersama Lab IoT & Fisika Institut Teknologi Bandung"),
        ("Ditujukan kepada", "PT Pertamina Patra Niaga"),
        ("Objek", "Prototipe Node Sensor Gas Leak Detector (GLD) V2 — sampel S/N GLD2-0x1001"),
    ]
    t = doc.add_table(rows=len(info), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    fixed(t, [Cm(4.6), Cm(12.2)])
    for i, (k, v) in enumerate(info):
        ctext(t.cell(i, 0), k, 9, True); shade(t.cell(i, 0), "EDEDED")
        ctext(t.cell(i, 1), v, 9)


PCB_ITER = [
    ("V1", "Baseline",
     "ESP32, ADC ADS1256, multiplexer TCA9548, DAC MCP4725, jalur analog INA333/OPA333/LM321 serta bridge "
     "sensor MQ berada pada satu motherboard. ULN2003 menggerakkan buzzer, alarm 5 V, dan kipas DC; "
     "TPL5010 mengatur watchdog/wake-up.",
     "Membuktikan fungsi dasar (akuisisi 8 sensor, AI, LoRa). Menjadi acuan identifikasi masalah: heater "
     "sensor MQ (sumber panas utama) bercampur dengan elektronika digital di satu papan."),
    ("V2", "Pemisahan sensor",
     "Bridge sensor MQ, INA333, dan MCP4725 dipindah ke SensorBoardMQ terpisah dengan saklar daya TPS22919; "
     "LM321 diganti OPA320; ditambah PCF8574 (ekspander I/O) dan SHT40 (suhu & kelembapan internal). Pin "
     "alarm 24 V; pemilih catu 24 V/baterai/5 V tanpa jumper.",
     "Sumber panas (heater MQ) dipisahkan dari logika dan dayanya dapat diputus terkendali, mendukung "
     "pengendalian suhu permukaan (kelas suhu T4) dan pembatasan energi saat gangguan. SHT40 memberi "
     "pemantauan suhu/kelembapan di dalam enclosure. Pemilih catu tanpa jumper mengurangi risiko salah "
     "konfigurasi di lapangan."),
    ("V3", "Perbaikan temuan V2",
     "Jalur catu +5VA disambungkan dan diperbaiki; soket kipas DC 5 V dikembalikan. Konfigurasi sensor "
     "module, OPA320, ADS1256, TCA9548, PCF8574, dan SHT40 dipertahankan.",
     "Menutup temuan uji pada V2 (siklus uji → koreksi → uji ulang). Kipas DC diperlukan untuk aliran gas "
     "ke sensor dan pendinginan; kondisi kipas mati menjadi kasus terburuk pada uji termal 30 Sep 2026."),
    ("V4", "Kesiapan mekanik & terminal",
     "Lubang baut untuk mengikat motherboard ke PCB cover; PCB cover terpisah sebagai modul terminal "
     "(24V +/−, ALARM −/+, RS485 A/B, FAN 5V/0V, BAT −/+) yang sekaligus menutup motherboard. Versi ini "
     "dipakai pada unit sampel S/N GLD2-0x1001.",
     "Komponen internal terikat kuat (tahan getaran, jarak bebas terjaga) sebagaimana dituntut konstruksi "
     "Ex d. Kabel lapangan hanya menyentuh terminal berlabel, terpisah dari elektronika, sehingga diagram "
     "pengawatan dan instruksi instalasi untuk dossier dapat ditetapkan jelas."),
    ("V5", "Modularisasi",
     "MotherBoardAdapter hanya memuat ESP32 dan watchdog TPL5010; fungsi akuisisi (ADS1256, TCA9548, OPA320, "
     "PCF8574, SHT40, +5VA) dipindah ke Sensor PCB melalui konektor 16-pin PHD2.0.",
     "Jalur analog/akuisisi terpisah dari prosesor, memudahkan pengujian per modul, penelusuran komponen "
     "kritis (BOM Ex-critical), dan penggantian modul tanpa mengubah papan utama. Disiapkan untuk iterasi "
     "produksi berikutnya."),
]

CASING_ITER = [
    ("Tahap 1\n(s/d Agu 2026)", "Prototipe awal",
     "Kotak junction box dengan kepala sensor silinder aluminium di bagian bawah; antena LoRa di sisi samping; "
     "beacon alarm terpisah (Foto 7–8).",
     "Membuktikan fungsi sistem lengkap di lab dan saat disaksikan Pertamina 6 Agu 2026. Menjadi dasar "
     "identifikasi kebutuhan enclosure tahan ledakan yang ringkas."),
    ("Tahap 2\nATEX Casing v2\n(31 Agu 2026)", "Desain CAD berdimensi",
     "Model solid STEP dengan dimensi terkendali (leher Ø75 mm, tinggi 191,51 mm) dan pelat pemasangan "
     "U-bolt 250 × 250 mm untuk pipa 2″/DN50, baut M10.",
     "Gambar teknik berdimensi menjadi syarat dossier (2.6.a). Metode pemasangan dengan klem U-bolt tanpa "
     "las/bor sesuai ketentuan keselamatan kilang."),
    ("Tahap 3\nGLD ATEX Case v3\n(8 Sep 2026)", "Desain cadangan ATEX",
     "Kotak sensor dengan kipas DC, cakram mesh stainless steel, tutup dengan kaca transparan, lug pemasangan "
     "dinding, dan boss cable entry; set 7 lembar gambar termasuk exploded view.",
     "Disiapkan sebagai alternatif bila desain utama tidak lolos uji Ex d; mesh stainless steel pada "
     "jalur masuk gas dirancang sebagai penahan api (flame arrestor)."),
    ("Tahap 4\nUnit sampel\n(Okt 2026)", "Konfigurasi sertifikasi",
     "Base die-cast aluminium ADC12 (dengan mitra casing PT Galaksi Megatama Indonesia); tutup berulir dengan "
     "mesh stainless steel di muka depan; O-ring karet pada sambungan tutup; antena dipindah ke muka atas "
     "base melalui bulkhead SMA; blanking plug pada entry yang tidak dipakai; sekrup grounding eksternal; "
     "entry kabel M20×1,5; posisi nameplate ditetapkan.",
     "Satu enclosure logam die-cast untuk konsep Ex d. Panjang ulir leher terukur ±14,8 mm, di atas batas "
     "minimum 8 mm IEC 60079-1 untuk volume >100 cm³. O-ring mendukung target IP66; entry tak terpakai "
     "ditutup; grounding eksternal tersedia. Unit ini diuji termal 30 Sep 2026 dan diukur dimensinya 7 Okt 2026."),
]

CASING_FOTO = [
    (FOTO_LAMA, "Tahap 1 — prototipe awal"),
    (os.path.join(HERE, "assets", "mech_drawings", "casing_v2_front.jpg"), "Tahap 2 — ATEX Casing v2"),
    (os.path.join(HERE, "assets", "mech_drawings", "case_v3-1.jpg"), "Tahap 3 — GLD ATEX Case v3"),
    (os.path.join(FOTO, "IMG_14.38.01.jpg"), "Tahap 4 — unit sampel"),
]


def iterasi_tabel(doc, judul, rows):
    para(doc, judul, 9, True, WD_ALIGN_PARAGRAPH.LEFT, before=6, after=3)
    t = doc.add_table(rows=1, cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, h in enumerate(["Versi", "Fokus", "Perubahan yang dilakukan", "Peningkatan menuju sertifikasi"]):
        ctext(t.cell(0, i), h, 9, True, WD_ALIGN_PARAGRAPH.CENTER); shade(t.cell(0, i), "D9E2F3")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            ctext(cells[i], v, 8.5, i < 2, WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    fixed(t, [Cm(2.3), Cm(2.4), Cm(6.0), Cm(6.1)])


def isi(doc):
    heading(doc, "1. Pendahuluan")
    para(doc, f"Dokumen ini disampaikan oleh PT LAPI Ganesha Utama kepada PT Pertamina Patra Niaga sebagai "
              f"bukti pemenuhan ketentuan pembayaran Termin 1 pada Surat Pesanan {NO_SP} untuk pekerjaan "
              f"Sertifikasi Hazardous Area (ATEX/IECEx) Gas Leak Detector (GLD) Tahap 2.")
    para(doc, "Ketentuan Termin 1 mencakup ketersediaan prototipe GLD yang telah diuji dan diperbaiki "
              "desainnya, disertai dokumentasi foto perangkat serta dokumentasi kunjungan PT Pertamina Patra "
              "Niaga ke Institut Teknologi Bandung. Bagian 2 menyajikan matriks pemenuhan setiap ketentuan, "
              "Bagian 3 menguraikan pengujian dan perbaikan desain yang telah dilaksanakan, sedangkan "
              "dokumentasi foto disajikan pada Lampiran 1 dan Lampiran 2.")
    para(doc, "Penilaian atas pemenuhan ketentuan dan persetujuan pembayaran merupakan kewenangan "
              "PT Pertamina Patra Niaga.")

    heading(doc, "2. Matriks Pemenuhan")
    rows = [
        ("1", "Prototipe GLD tersedia",
         "Prototipe Node Sensor GLD V2 telah terakit lengkap: casing die-cast aluminium dengan tutup "
         "ber-mesh stainless steel, 8 sensor gas (MQ-2/3B/4/5/6/7B/8/135), kipas DC, motherboard ESP32-S3 "
         "dengan PCB cover terminal (V4), radio LoRa, catu 24 VDC. Unit yang diajukan untuk sertifikasi "
         "dengan nomor seri GLD2-0x1001; sebanyak 4 unit GLD tersedia (3 untuk RU IV + 1 cadangan).",
         "Lampiran 1 (Foto 1–8)", "Tersedia"),
        ("2", "Prototipe telah diuji",
         "(a) Factory Acceptance Test laboratorium FAT-01 s.d. FAT-07, 7 dari 7 item lulus "
         "(Laporan FAT No. LGU/GLD/FAT/2026-001 Rev 1.4). (b) Uji termal internal 30 Sep 2026: "
         "8 kanal thermocouple, 3 kondisi (normal, beban maksimum, kipas mati), 180 menit. "
         "(c) Pengukuran dimensi enclosure dengan jangka sorong, 7 Okt 2026.",
         "Bagian 3.1; Laporan FAT", "Tersedia"),
        ("3", "Desain telah diperbaiki",
         "PCB Node Sensor GLD diiterasi V1 s.d. V5 dengan catatan perubahan per versi; casing berkembang "
         "dari ATEX Casing v2 ke GLD ATEX Case v3; antena dipindah ke muka atas base; PCB cover terminal "
         "kabel & lubang baut motherboard ditambahkan; konsep proteksi diarahkan ke Ex d (flameproof).",
         "Bagian 3.2; Gambar 1; Lampiran 1 (Foto 7–8)", "Tersedia"),
        ("4", "Dokumentasi foto perangkat",
         "Foto prototipe terkini (tampak depan, belakang, samping, bagian dalam, motherboard, tutup "
         "ber-mesh) serta foto iterasi sebelumnya sebagai pembanding perbaikan desain.",
         "Lampiran 1", "Terlampir"),
        ("5", "Dokumentasi kunjungan ke ITB",
         "Dokumentasi kunjungan PT Pertamina Patra Niaga ke Lab IoT, Instrumentation and Computations, "
         "Gedung Laboratorium Fisika Terpadu, Institut Teknologi Bandung, 6 Agustus 2026 (peninjauan "
         "prototipe dan demonstrasi integrasi sistem).",
         "Lampiran 2", "Terlampir"),
    ]
    t = doc.add_table(rows=1, cols=5)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    widths = [Cm(0.9), Cm(3.0), Cm(8.2), Cm(2.6), Cm(2.1)]
    for i, h in enumerate(["No", "Ketentuan Termin 1", "Bukti pemenuhan", "Rujukan", "Status"]):
        ctext(t.cell(0, i), h, 9, True, WD_ALIGN_PARAGRAPH.CENTER); shade(t.cell(0, i), "D9E2F3")
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            ctext(cells[i], v, 8.5, i in (1, 4),
                  WD_ALIGN_PARAGRAPH.CENTER if i in (0, 4) else WD_ALIGN_PARAGRAPH.LEFT)
    fixed(t, widths)

    heading(doc, "3. Uraian Pemenuhan")
    heading(doc, "3.1 Prototipe telah diuji", 10.5)
    tests = [
        ("Factory Acceptance Test (FAT) laboratorium",
         "Lab IoT ITB, Mei–Agustus 2026; disaksikan PT Pertamina Patra Niaga 6 Agustus 2026",
         "7/7 item lulus: pembacaan 8 kanal sensor, klasifikasi gas AI di perangkat (ESP32-S3), "
         "komunikasi LoRa, mesh & failover Cluster Head, integrasi end-to-end, alarm otomatis, serta monitoring data dan alarm di dashboard."),
        ("Uji termal internal (persiapan kelas suhu T4)",
         "30 September 2026",
         "8 kanal thermocouple; kondisi normal, beban maksimum, dan kipas mati; 180 menit. Ekstrapolasi "
         "ke suhu ambien +60 °C: titik terpanas ±114 °C, di bawah batas efektif T4 (130 °C). Titik ukur "
         "komponen daya akan dilengkapi pada uji laboratorium."),
        ("Pengukuran dimensi enclosure",
         "7 Oktober 2026",
         "Jangka sorong: dinding leher ulir 10,50 mm (pangkal) / 5,80 mm (atas), tinggi leher berulir "
         "±14,8 mm, diameter dalam ±84,7 mm, penampang O-ring 2,70 mm — data masukan dossier Ex d."),
    ]
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for i, h in enumerate(["Pengujian", "Waktu / tempat", "Hasil"]):
        ctext(t.cell(0, i), h, 9, True, WD_ALIGN_PARAGRAPH.CENTER); shade(t.cell(0, i), "D9E2F3")
    for r in tests:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            ctext(cells[i], v, 8.5, i == 0)
    fixed(t, [Cm(4.2), Cm(4.2), Cm(8.4)])

    heading(doc, "3.2 Desain telah diperbaiki", 10.5)
    para(doc, "Perbaikan desain dilakukan bertahap pada dua komponen utama: papan elektronik (PCB) Node "
              "Sensor GLD dan casing/enclosure. Tabel berikut merinci perubahan setiap versi dan kontribusinya "
              "terhadap persyaratan sertifikasi hazardous area (konsep proteksi usulan Ex d, II 2G Ex db IIC "
              "T4 Gb, mengacu IEC 60079-0 dan IEC 60079-1).")
    iterasi_tabel(doc, "Tabel 1. Iterasi PCB Node Sensor GLD dan kontribusinya menuju sertifikasi", PCB_ITER)
    if os.path.exists(PCB_ITERASI):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(img_stream(PCB_ITERASI, 2000), width=Cm(15.5))
        para(doc, "Gambar 1. Iterasi PCB Node Sensor GLD, versi V1 hingga V5", 8.5,
             align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    iterasi_tabel(doc, "Tabel 2. Iterasi casing/enclosure GLD dan kontribusinya menuju sertifikasi", CASING_ITER)
    t = doc.add_table(rows=2, cols=len(CASING_FOTO))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fixed(t, [Cm(16.8 / len(CASING_FOTO))] * len(CASING_FOTO))
    for i, (path, cap) in enumerate(CASING_FOTO):
        p = t.cell(0, i).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(img_stream(path), height=Cm(3.0))
        ctext(t.cell(1, i), cap, 8, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(doc, "Gambar 2. Perkembangan casing GLD dari prototipe awal hingga unit sampel sertifikasi", 8.5,
         align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, before=4)
    para(doc, "Catatan: penilaian kesesuaian dilakukan oleh laboratorium dan lembaga sertifikasi. Tindak "
              "lanjut berjalan: gland Ex d bersertifikat, verifikasi mesh penahan api, uji IP66, titik ukur "
              "termal komponen daya, rating PTC 24 V, proteksi polaritas terbalik, dan keputusan terminal BAT.",
         8.5, before=2)

    heading(doc, "4. Pernyataan").paragraph_format.page_break_before = True
    para(doc, f"Berdasarkan uraian dan lampiran di atas, PT LAPI Ganesha Utama menyatakan bahwa ketentuan "
              f"Termin 1 pada Surat Pesanan {NO_SP} (Sertifikasi) — tersedianya prototipe GLD yang telah "
              f"diuji dan diperbaiki desainnya, disertai foto perangkat dan foto kunjungan ke ITB — telah "
              f"dipenuhi, dan mengajukan dokumen ini kepada PT Pertamina Patra Niaga untuk diperiksa dan "
              f"disetujui sebagai dasar proses pembayaran Termin 1.")


def pengesahan(doc):
    heading(doc, "5. Lembar Pengesahan")
    para(doc, TEMPAT, align=WD_ALIGN_PARAGRAPH.LEFT, after=8)
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    fixed(t, [Cm(8.4), Cm(8.4)])
    ctext(t.cell(0, 0), "Diajukan oleh,\nPT LAPI Ganesha Utama", 9.5, True, WD_ALIGN_PARAGRAPH.CENTER)
    ctext(t.cell(0, 1), "Diperiksa dan disetujui oleh,\nPT Pertamina Patra Niaga", 9.5, True,
          WD_ALIGN_PARAGRAPH.CENTER)
    shade(t.cell(0, 0), "EDEDED"); shade(t.cell(0, 1), "EDEDED")
    row = t.add_row().cells
    inner = row[0].add_table(rows=3, cols=2)
    fixed(inner, [Cm(3.9), Cm(3.9)])
    for i, (jab, nama) in enumerate(LGU_TTD):
        ctext(inner.cell(0, i), jab, 9, align=WD_ALIGN_PARAGRAPH.CENTER)
        inner.rows[1].height = Cm(2.4)
        p = inner.cell(2, i).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p, nama, 9, True, underline=True)
    inner2 = row[1].add_table(rows=3, cols=1)
    fixed(inner2, [Cm(7.8)])
    ctext(inner2.cell(0, 0), PPN_JAB + ",", 9, align=WD_ALIGN_PARAGRAPH.CENTER)
    inner2.rows[1].height = Cm(2.4)
    p = inner2.cell(2, 0).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, PPN_NAMA, 9, True, underline=True)
    fixed(t, [Cm(8.4), Cm(8.4)])
    para(doc, "Catatan pemeriksaan PT Pertamina Patra Niaga:", 9, True, WD_ALIGN_PARAGRAPH.LEFT, before=10, after=2)
    for _ in range(3):
        para(doc, "." * 150, 8, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)


def lampiran(doc):
    doc.add_page_break()
    para(doc, "LAMPIRAN 1 — FOTO PERANGKAT PROTOTIPE GLD", 12, True, WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, "Gas Leak Detector GLD V2 · S/N GLD2-0x1001 · Lab IoT ITB, Oktober 2026", 9,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    f = lambda n: os.path.join(FOTO, n)
    items = [
        (f("IMG_14.35.32.jpg"), "Foto 1 — Tampak depan, tertutup (tutup ber-mesh stainless steel)"),
        (f("IMG_14.38.01.jpg"), "Foto 2 — Tampak belakang, antena LoRa di muka atas base"),
        (f("IMG_14.35.54.jpg"), "Foto 3 — Tampak samping, cable entry"),
        (f("IMG_14.08.05.jpg"), "Foto 4 — Bagian dalam: 8 sensor gas pada sensor board"),
        (f("IMG_14.03.58.jpg"), "Foto 5 — Motherboard dengan PCB cover terminal (V4) di dalam base"),
        (f("IMG_14.12.07.jpg"), "Foto 6 — Tutup ber-mesh dengan kipas DC dan modul sensor"),
        (FOTO_LAMA, "Foto 7 — Iterasi sebelumnya (Agustus 2026): antena di sisi samping"),
        (os.path.join(ROOT, "Sumber Dokumen", "sertifikasi-atex-gld-v2-2026", "GLD", "Casing_Belakang.jpg"),
         "Foto 8 — Iterasi sebelumnya (Agustus 2026), tampak belakang"),
    ]
    foto_grid(doc, items)

    doc.add_page_break()
    para(doc, "LAMPIRAN 2 — FOTO KUNJUNGAN PT PERTAMINA PATRA NIAGA KE ITB", 12, True,
         WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, "Lab IoT, Instrumentation and Computations, Gedung Laboratorium Fisika Terpadu, "
              "Institut Teknologi Bandung · 6 Agustus 2026", 9, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    fotos = sorted(p for p in glob.glob(os.path.join(FOTO_KUNJUNGAN, "*"))
                   if p.lower().endswith((".jpg", ".jpeg", ".png")))
    if fotos:
        def judul(p):  # "04_Paparan_data_..._a.jpg" -> "Paparan data ..."
            kata = os.path.splitext(os.path.basename(p))[0].split("_")[1:]
            if len(kata[-1]) == 1:
                kata = kata[:-1]
            return " ".join(kata).replace("Lab IoT", "Lab IoT ITB")
        items = [(p, f"Foto {i + 9} — {judul(p)}") for i, p in enumerate(fotos)]
    else:
        items = [(None, f"Foto {i} — Kunjungan Pertamina ke Lab IoT ITB, 6 Agustus 2026") for i in range(9, 13)]
    foto_grid(doc, items)


def build():
    os.makedirs(OUT_DIR, exist_ok=True)
    doc = docx.Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.1)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(10)
    footer(s)
    kop(doc)
    isi(doc)
    pengesahan(doc)
    lampiran(doc)
    out = os.path.join(OUT_DIR, NAMA + ".docx")
    doc.save(out)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT_DIR, out],
                   check=True, capture_output=True)
    print("->", out)
    print("->", os.path.join(OUT_DIR, NAMA + ".pdf"))
    print("Foto kunjungan ITB:", "ada" if os.path.isdir(FOTO_KUNJUNGAN) else "BELUM ADA (bingkai kosong)")


if __name__ == "__main__":
    build()
