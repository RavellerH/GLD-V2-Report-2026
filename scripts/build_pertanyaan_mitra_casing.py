# -*- coding: utf-8 -*-
"""Daftar pertanyaan untuk produsen/mitra casing GLD terkait persyaratan ATEX/IECEx.

Setiap pertanyaan diikat ke butir checklist ExCB yang ditutupnya, supaya jawaban
mitra langsung bisa dimasukkan ke dokumen submission
(`Dokumen_Teknis_Sertifikasi_GLD_IECEx_ATEX`).

Output: Deliverables/Daftar_Pertanyaan_Mitra_Casing_ATEX_GLD.docx
"""
import os

from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "Deliverables", "Daftar_Pertanyaan_Mitra_Casing_ATEX_GLD.docx")

NAVY = RGBColor(0x1B, 0x33, 0x5F)
BLUE = RGBColor(0x2B, 0x5F, 0xCB)
INK = RGBColor(0x26, 0x23, 0x21)
GRAY = RGBColor(0x55, 0x5B, 0x66)
RED = RGBColor(0xB4, 0x52, 0x0F)
HEAD_SHADE = "EAEFF9"
NOTE_SHADE = "F3F6FC"
WARN_SHADE = "FBEEE4"

doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.font.color.rgb = INK

sec = doc.sections[0]
sec.left_margin = Cm(1.9)
sec.right_margin = Cm(1.9)
sec.top_margin = Cm(1.7)
sec.bottom_margin = Cm(1.7)


def shade(cell, hex_color):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


def p(text="", size=10.5, bold=False, italic=False, color=None, space_after=8, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color or INK
    return para


def note_box(text, fill=NOTE_SHADE):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    shade(cell, fill)
    cell.paragraphs[0].paragraph_format.space_after = Pt(2)
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(9.5)
    run.font.color.rgb = GRAY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def qgroup(judul, pengantar, pertanyaan, widths=(0.42, 3.55, 1.35, 2.05)):
    """pertanyaan: list of (no, teks pertanyaan, 'untuk melengkapi')."""
    p(judul, size=12.5, bold=True, color=NAVY, space_after=4)
    if pengantar:
        p(pengantar, size=9.5, color=GRAY, space_after=6)
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0]
    trPr = hdr._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for i, h in enumerate(["NO", "PERTANYAAN", "UNTUK MELENGKAPI", "JAWABAN MITRA"]):
        shade(hdr.cells[i], HEAD_SHADE)
        hdr.cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
        r = hdr.cells[i].paragraphs[0].add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = NAVY
    for no, teks, untuk in pertanyaan:
        row = tbl.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        cells = row.cells
        for idx, val in enumerate([no, teks, untuk, ""]):
            para = cells[idx].paragraphs[0]
            para.paragraph_format.space_after = Pt(3)
            para.paragraph_format.space_before = Pt(3)
            run = para.add_run(val)
            run.font.size = Pt(9.5)
            if idx == 2:
                run.font.size = Pt(8.5)
                run.font.color.rgb = GRAY
    tbl.autofit = False
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl._tbl.tblPr.append(layout)
    grid = tbl._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    for row in tbl.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return tbl


# ---------------------------------------------------------------- kop
lt = doc.add_paragraph()
lt.paragraph_format.space_after = Pt(0)
r = lt.add_run("PT LAPI GANESHA UTAMA")
r.font.size = Pt(13)
r.font.bold = True
r.font.color.rgb = NAVY
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(14)
r = sub.add_run("Bekerja sama dengan Institut Teknologi Bandung — Lab IoT & Lab Fisika")
r.font.size = Pt(9)
r.font.color.rgb = GRAY

p("DAFTAR PERTANYAAN UNTUK MITRA CASING", size=8.5, bold=True, color=BLUE, space_after=2)
p("Persyaratan Enclosure ATEX / IECEx", size=18, bold=True, color=NAVY, space_after=4)
p("Gas Leak Detector (GLD) V2 — Node Sensor", size=11, color=GRAY, space_after=12)

meta = doc.add_table(rows=0, cols=2)
meta.style = "Table Grid"
for k, v in [
    ("Ditujukan kepada", "……………………………………  (mitra perancang & produsen casing/enclosure GLD)"),
    ("Dari", "PT LAPI Ganesha Utama — proyek sertifikasi hazardous area GLD V2"),
    ("Dasar", "IECEx/ATEX Certification Information Requirements (ExCB), Bagian 1 & Bagian 2.6"),
    ("Tujuan", "Mengumpulkan data enclosure yang belum tersedia agar berkas sertifikasi dapat dilengkapi"),
    ("Diminta kembali", "……………………………………"),
] :
    row = meta.add_row().cells
    shade(row[0], HEAD_SHADE)
    for idx, val in enumerate([k, v]):
        row[idx].paragraphs[0].paragraph_format.space_after = Pt(2)
        run = row[idx].paragraphs[0].add_run(val)
        run.font.size = Pt(9.5)
        run.font.bold = (idx == 0)
        if idx == 0:
            run.font.color.rgb = NAVY
    row[0].width = Inches(1.4)
    row[1].width = Inches(5.1)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

note_box(
    "Konteks singkat produk. GLD V2 adalah perangkat deteksi kebocoran gas titik-tetap untuk area kilang. "
    "Dimensi unit ±200 × 90 × 290 mm; catu daya 24 VDC dari jaringan listrik site; dipasang ke struktur "
    "existing memakai bracket L / U-bolt (tanpa pengeboran dan pengelasan pada struktur). Rencana material: "
    "paduan aluminium dan baja nirkarat, tanpa plastik/PVC. Perangkat memiliki delapan sensor gas MQ yang "
    "elemen sensingnya HARUS bersentuhan langsung dengan udara ambien (saat ini lewat penutup mesh), satu "
    "modul alarm suara/visual, antena SMA eksternal, serta port USB, sensor, daya, dan fan. Target yang "
    "diajukan: Zona 1, Kategori 2G, Grup II, kelas suhu T4 (≤135 °C), dan IP66/67 — semuanya masih berupa "
    "permintaan, belum disetujui lembaga sertifikasi."
)

note_box(
    "Catatan penting sebelum menjawab. Dua hal belum kami putuskan dan justru butuh masukan Anda: (1) metode "
    "proteksi ledakan yang akan dipakai (Ex d / Ex e / Ex i atau kombinasi), dan (2) cara menangani bukaan "
    "sensor yang tidak bisa disegel. Kedua hal ini menentukan hampir seluruh parameter enclosure, jadi mohon "
    "bagian B dijawab lebih dahulu bila waktu terbatas. Bila suatu data belum ada, tulis “belum ada” "
    "— jangan diisi perkiraan, karena jawaban ini menjadi bagian berkas resmi ke lembaga sertifikasi.",
    fill=WARN_SHADE,
)

# ---------------------------------------------------------------- A
qgroup(
    "A. Pengalaman, peran, dan legalitas mitra",
    "Menentukan apakah jalur tercepat adalah memakai enclosure yang sudah bersertifikat, atau mensertifikasi "
    "rancangan baru bersama-sama.",
    [
        ("A1", "Apakah perusahaan Anda pernah membuat enclosure untuk area berbahaya (Ex)? Sebutkan produk, "
               "lembaga sertifikasi, nomor sertifikat, dan tahunnya.", "Penilaian jalur sertifikasi"),
        ("A2", "Apakah Anda dapat menyediakan enclosure yang SUDAH memiliki sertifikat komponen Ex "
               "(Ex component certificate / certified empty enclosure)? Bila ya, mohon lampirkan sertifikatnya "
               "beserta batasan penggunaannya.", "Butir 2.6.i — sertifikat komponen Ex"),
        ("A3", "Bila belum bersertifikat: apakah Anda bersedia rancangan ini diajukan sertifikasi atas nama "
               "bersama, dan namanya dicantumkan sebagai lokasi manufaktur dalam berkas?", "Butir 1.4"),
        ("A4", "Apakah Anda bersedia menerima audit sistem mutu oleh lembaga sertifikasi (IECEx QAR atau "
               "asesmen produksi ATEX) di lokasi produksi Anda?", "Butir 1.5 — sistem mutu"),
        ("A5", "Alamat pabrik dan profil fasilitas: luas area, peralatan utama (mesin CNC, las, coating), "
               "jumlah personel, dan kapasitas produksi per bulan.", "Butir 1.4 — profil fasilitas"),
        ("A6", "Apakah perusahaan Anda memegang sertifikat ISO 9001? Bila ya: nomor, ruang lingkup, lembaga "
               "penerbit, dan masa berlaku.", "Butir 1.5"),
    ],
)

# ---------------------------------------------------------------- B
qgroup(
    "B. Metode proteksi ledakan — paling menentukan",
    "Jawaban bagian ini membuka empat pekerjaan yang saat ini tertahan: kalkulasi proteksi ledakan, penandaan "
    "flame path pada gambar, pemilihan material, dan sebagian BOM.",
    [
        ("B1", "Metode proteksi apa yang dapat Anda dukung untuk bentuk produk seperti ini — Ex d (flameproof), "
               "Ex e (increased safety), Ex t, atau kombinasi? Mana yang Anda rekomendasikan, dan mengapa?",
         "Butir 2.6.e — konsep proteksi"),
        ("B2", "Bila Ex d: apakah Anda mampu memenuhi DAN mendokumentasikan parameter flame path sesuai "
               "IEC 60079-1 — lebar celah maksimum, panjang jalur, kekasaran permukaan (Ra), dan volume "
               "internal? Berapa toleransi permesinan yang dapat Anda jamin secara konsisten?",
         "Butir 2.6.a & 2.6.e"),
        ("B3", "Bila Ex e / Ex t: bagaimana Anda memastikan derajat proteksi masuk (IP) dan jarak bebas/rambat "
               "pada bagian yang Anda buat?", "Butir 2.6.e"),
        ("B4", "Delapan sensor MQ harus terpapar udara ambien sehingga rumah sensor tidak dapat disegel penuh. "
               "Bagaimana Anda menangani bukaan ini secara Ex — apakah dengan flame arrestor sinter-metal "
               "bersertifikat, dan apakah Anda punya pemasok untuk komponen tersebut?",
         "Celah desain yang belum tertutup"),
        ("B5", "Bagaimana penetrasi lain ditangani: antena SMA eksternal, buzzer/modul alarm, port USB, port "
               "fan, dan masukan daya 24 VDC? Mana yang Anda sarankan dihilangkan atau dipindahkan agar "
               "sertifikasi lebih sederhana?", "Butir 2.6.a & 2.6.e"),
        ("B6", "Apakah desain memerlukan dua kompartemen terpisah (misalnya kompartemen elektronik dan "
               "kompartemen terminal)? Bagaimana pemisahannya direalisasikan?", "Butir 2.6.a"),
    ],
)

# ---------------------------------------------------------------- C
qgroup(
    "C. Material",
    "Data di bagian ini langsung mengisi butir 2.6.b (BOM komponen kritis keselamatan) dan 2.6.c (datasheet "
    "material non-logam) yang sekarang masih kosong.",
    [
        ("C1", "Paduan aluminium yang dipakai: grade spesifik dan komposisi kimianya, disertai sertifikat "
               "material (mill certificate). Mohon dipastikan kandungan magnesium dan logam ringan lain masih "
               "dalam batas yang diizinkan IEC 60079-0 untuk kelompok peralatan kami.",
         "Butir 2.6.b — BOM Ex-critical"),
        ("C2", "Baja nirkarat: grade (304 / 316 / 316L) dan dipakai untuk komponen yang mana.",
         "Butir 2.6.b"),
        ("C3", "Seluruh material NON-LOGAM (gasket, seal, o-ring, jendela/lensa, label, potting, kaki karet): "
               "jenis material, merek, rentang suhu kerja, ketahanan penuaan dan UV, sifat antistatik, "
               "ketahanan kimia, dan nilai CTI — sertakan datasheet pemasok atau deklarasi kesesuaian.",
         "Butir 2.6.c — datasheet material"),
        ("C4", "Pelapisan/pengecatan permukaan: jenis, ketebalan lapisan, dan bagaimana risiko muatan "
               "elektrostatik dikendalikan.", "Butir 2.6.c"),
        ("C5", "Baut dan pengencang: grade/property class, material, jumlah per tutup, serta torsi pengencangan "
               "yang disyaratkan.", "Butir 2.6.b & 2.6.g"),
        ("C6", "Cable gland dan penutup lubang (blanking plug): tipe, ukuran ulir, merek, dan apakah sudah "
               "bersertifikat Ex.", "Butir 2.6.b & 2.6.i"),
    ],
)

# ---------------------------------------------------------------- D
qgroup(
    "D. Konstruksi, pengujian, dan dokumentasi",
    "",
    [
        ("D1", "Apakah rancangan Anda pernah diuji IP66/IP67 menurut IEC 60529? Bila ya, lampirkan laporan uji "
               "dan sebutkan laboratoriumnya.", "Target IP66/67"),
        ("D2", "Apakah enclosure pernah lulus uji ketahanan impak dan uji jatuh sesuai IEC 60079-0 (umumnya "
               "7 J untuk risiko mekanis tinggi dan 4 J untuk risiko rendah — mohon dikonfirmasi terhadap "
               "edisi standar yang berlaku), serta uji kejut termal?", "Butir 2.6.e"),
        ("D3", "Bagaimana pembumian direalisasikan: terminal earth internal dan eksternal, ukuran konduktor, "
               "dan penandaannya?", "Butir 2.6.a & 2.6.g"),
        ("D4", "Sensor MQ memakai elemen pemanas. Bagaimana rancangan casing memengaruhi suhu permukaan luar, "
               "dan dapatkah Anda membantu pengukuran atau analisis titik terpanas untuk memverifikasi kelas "
               "suhu T4 (≤135 °C)?", "Butir 2.6.f — kalkulasi kelas suhu"),
        ("D5", "Dapatkah Anda menyediakan gambar teknik bertoleransi lengkap — gambar rakitan, gambar "
               "komponen, penampang yang menandai parameter flame path, serta gambar terminal dan pembumian — "
               "dalam format yang dapat dikirim ke lembaga sertifikasi?", "Butir 2.6.a — paket gambar"),
        ("D6", "Mohon uraikan proses manufaktur yang memengaruhi keselamatan proteksi ledakan: pengendalian "
               "akurasi permesinan, perlakuan permukaan jalur api, proses las, potting, die casting, dan "
               "bonding — termasuk cara pemeriksaannya.", "Butir 2.6.d — proses manufaktur"),
        ("D7", "Apakah nameplate dibuat oleh Anda? Sebutkan material, metode penandaan (laser/etsa/cetak), "
               "cara pemasangan, dan ketahanannya terhadap lingkungan kilang.", "Butir 2.6.h — nameplate"),
        ("D8", "Bagaimana rancangan mengakomodasi pemasangan bracket L / U-bolt ke struktur existing, termasuk "
               "orientasi tiang vertikal maupun handrail horizontal?", "Kompatibilitas mounting"),
    ],
)

# ---------------------------------------------------------------- E
qgroup(
    "E. Sampel uji, jadwal, dan komersial",
    "",
    [
        ("E1", "Berapa lead time untuk prototipe pertama dan untuk unit sampel uji? Perlu diketahui bahwa "
               "sebagian pengujian sertifikasi bersifat merusak, sehingga dibutuhkan lebih dari satu unit "
               "identik.", "Bagian 3 — informasi sampel"),
        ("E2", "Dapatkah Anda memproduksi beberapa unit yang benar-benar identik dan memberi nomor seri pada "
               "tiap unit?", "Butir 3.1 — model & nomor seri"),
        ("E3", "Perkiraan biaya dan jumlah pesanan minimum (MOQ) untuk tahap prototipe maupun produksi awal.",
         "Perencanaan biaya"),
        ("E4", "Setelah sertifikat terbit, perubahan pada enclosure akan memengaruhi keabsahan sertifikat. "
               "Bersediakah Anda terikat pada mekanisme pengendalian perubahan (tidak ada perubahan material, "
               "dimensi, atau proses tanpa persetujuan tertulis)?", "Butir 1.5 & pemeliharaan sertifikat"),
        ("E5", "Siapa PIC teknis di pihak Anda yang akan menjawab pertanyaan lembaga sertifikasi, dan melalui "
               "kanal apa?", "Butir 1.3 — kontak"),
    ],
)

# ---------------------------------------------------------------- penutup
doc.add_page_break()
p("Lampiran yang kami harapkan menyertai jawaban", size=12.5, bold=True, color=NAVY, space_after=6)
att = doc.add_table(rows=1, cols=3)
att.style = "Table Grid"
for i, h in enumerate(["NO", "BERKAS", "STATUS (ADA / BELUM ADA)"]):
    shade(att.rows[0].cells[i], HEAD_SHADE)
    att.rows[0].cells[i].paragraphs[0].paragraph_format.space_after = Pt(2)
    r = att.rows[0].cells[i].paragraphs[0].add_run(h)
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = NAVY
for i, nama in enumerate([
    "Sertifikat komponen Ex enclosure (bila ada)",
    "Sertifikat material (mill certificate) paduan aluminium dan baja nirkarat",
    "Datasheet seluruh material non-logam (gasket, seal, potting, label)",
    "Sertifikat Ex cable gland dan blanking plug",
    "Laporan uji IP66/IP67 (bila pernah dilakukan)",
    "Laporan uji impak / jatuh / kejut termal (bila pernah dilakukan)",
    "Gambar teknik bertoleransi: rakitan, komponen, penampang flame path",
    "Uraian proses manufaktur dan titik pemeriksaannya",
    "Profil perusahaan dan alamat pabrik",
    "Sertifikat ISO 9001 (bila ada)",
], 1):
    row = att.add_row()
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    cells = row.cells
    for idx, val in enumerate([str(i), nama, ""]):
        para = cells[idx].paragraphs[0]
        para.paragraph_format.space_before = Pt(4)
        para.paragraph_format.space_after = Pt(4)
        run = para.add_run(val)
        run.font.size = Pt(9.5)
    cells[0].width = Inches(0.4)
    cells[1].width = Inches(4.4)
    cells[2].width = Inches(2.2)
doc.add_paragraph().paragraph_format.space_after = Pt(14)

note_box(
    "Bila banyak butir belum dapat dijawab, kami usulkan satu sesi pembahasan teknis bersama untuk "
    "menyelesaikan bagian B lebih dahulu (metode proteksi dan penanganan bukaan sensor). Dua keputusan itu "
    "menentukan hampir seluruh jawaban di bagian C dan D, sehingga menuntaskannya lebih dulu akan lebih "
    "efisien daripada menjawab seluruh daftar sekaligus."
)

p("Dikirim oleh: ………………………………  ·  Jabatan: ………………………………  ·  Tanggal: ………………………",
  size=9.5, color=GRAY, space_after=4)
p("Dijawab oleh: ………………………………  ·  Jabatan: ………………………………  ·  Tanggal: ………………………",
  size=9.5, color=GRAY, space_after=14)

p("Jawaban atas daftar ini akan dimasukkan ke berkas teknis sertifikasi GLD V2 "
  "(Dokumen No. LGU/GLD/IECEX-TDF/2026-001). Mohon tidak mengisi dengan perkiraan.",
  size=8.5, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.save(OUT)
print("written", OUT)
