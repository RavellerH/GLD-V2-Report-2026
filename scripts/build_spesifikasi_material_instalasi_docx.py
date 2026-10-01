# -*- coding: utf-8 -*-
"""Build "Spesifikasi Teknis Material Instalasi GLD — RU IV Cilacap" (corporate .docx).

Dokumen untuk tim PT Pertamina Patra Niaga, bahan rapat 2 Oktober 2026: spesifikasi
tiang besi (ditanam & dicor), kabel catu daya 24 VDC (dua opsi skema PSU), jaringan
Gateway-server, fastener, grounding, dan BoQ. Konten di-hardcode di skrip ini
(bukan diparse dari HTML); helper format diambil dari
build_persiapan_instalasi_corporate_docx.py agar tampilan seragam.

Run:  python3 scripts/build_spesifikasi_material_instalasi_docx.py
PDF:  python3 scripts/docx_to_pdf_with_toc.py Deliverables/<nama>.docx
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from build_persiapan_instalasi_corporate_docx import (  # noqa: E402
    make_doc, set_cell_shading, add_field, NAVY, GRAY, WHITE, INK,
    GREEN, AMBER, RED, HEAD_SHADE, ZEBRA_SHADE,
)
from docx.shared import Pt, Inches, RGBColor  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "Deliverables",
                   "Spesifikasi_Material_Instalasi_GLD_RU-IV_Cilacap.docx")

DOC_NO = "LGU/GLD/INSTALASI-SPEK/2026-001"
REV = "1.5"
DATE = "2 Oktober 2026"
HEADER = "Spesifikasi Material Instalasi GLD — RU IV Cilacap"

# Asumsi perencanaan (diganti begitu data lapangan tersedia)
N_GLD = 3
N_CH = 6          # CH tersedia >5 unit; 6 dipakai sbg angka perencanaan
N_GW = 1

STATUS_STYLE = {
    "FINAL": (GREEN, "EAF4EC"),
    "USULAN": (AMBER, "FFF3DC"),
    "KONFIRMASI": (RED, "FBE6E4"),
}


# ---------------------------------------------------------------- helpers
def para(doc, text, size=10.2, bold=False, color=None, italic=False, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    _runs(p, text, size, bold, color, italic)
    return p


def _runs(p, text, size=10.2, bold=False, color=None, italic=False):
    """**tebal** didukung secara sederhana."""
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.font.size = Pt(size)
        r.font.bold = bold or (i % 2 == 1)
        r.font.italic = italic
        if color is not None:
            r.font.color.rgb = color


def bullets(doc, items, size=10):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        _runs(p, it, size)


def heading(doc, text, level=1):
    h = doc.add_heading(level=level)
    r = h.add_run(text)
    r.font.color.rgb = NAVY
    return h


def table(doc, head, rows, widths, size=9.2, status_col=None):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        set_cell_shading(c, "1B2A4A")
        c.paragraphs[0].paragraph_format.space_after = Pt(1)
        r = c.paragraphs[0].add_run(h)
        r.font.bold = True; r.font.size = Pt(size); r.font.color.rgb = WHITE
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            if status_col is not None and i == status_col and val in STATUS_STYLE:
                col, bg = STATUS_STYLE[val]
                set_cell_shading(c, bg)
                r = c.paragraphs[0].add_run(val)
                r.font.bold = True; r.font.size = Pt(size - 0.6); r.font.color.rgb = col
                continue
            if ri % 2 == 1:
                set_cell_shading(c, ZEBRA_SHADE)
            _runs(c.paragraphs[0], str(val), size)
    fix_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def fix_widths(t, widths):
    """Lebar kolom tetap (tblLayout=fixed + tblGrid) agar Word & LibreOffice patuh."""
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 1440)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)


def banner(doc, title, text, bg="FFF6E5"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    set_cell_shading(c, bg)
    c.paragraphs[0].paragraph_format.space_after = Pt(2)
    r = c.paragraphs[0].add_run(title)
    r.font.bold = True; r.font.size = Pt(9.8)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    _runs(p2, text, 9.4)
    fix_widths(t, [6.9])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


# ---------------------------------------------------------------- build
def build():
    doc, sec = make_doc()

    # letterhead
    lt = doc.add_table(rows=1, cols=1)
    lc = lt.rows[0].cells[0]
    fix_widths(lt, [6.9])
    set_cell_shading(lc, "1B2A4A")
    lc.paragraphs[0].paragraph_format.space_after = Pt(2)
    lr = lc.paragraphs[0].add_run("PT LAPI GANESHA UTAMA")
    lr.font.bold = True; lr.font.size = Pt(14); lr.font.color.rgb = WHITE
    lp2 = lc.add_paragraph()
    lp2.paragraph_format.space_after = Pt(3)
    lr2 = lp2.add_run("Bekerja sama dengan Lab IoT & Fisika Institut Teknologi Bandung")
    lr2.font.size = Pt(9.5); lr2.font.color.rgb = RGBColor(0xC7, 0xD2, 0xE0)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    title = doc.add_heading(level=0)
    tr = title.add_run("Spesifikasi Material & Kebutuhan Instalasi GLD — RU IV Cilacap")
    tr.font.size = Pt(18); tr.font.bold = True; tr.font.color.rgb = NAVY
    para(doc, "Tiang besi, kabel catu daya 24 VDC, jaringan Gateway–server, fastener, "
         "grounding, dan daftar kebutuhan (BoQ) untuk pemasangan sistem Gas Leak Detection "
         "di area aman perimeter Sulfur Recovery Unit (SRU). RU IV menyatakan bersedia mendukung "
         "penuh pelaksanaan instalasi; dokumen ini disusun agar tim RU IV dapat mulai menyiapkan "
         "material dan pekerjaan sipil/elektrikal.",
         size=10.6, color=GRAY, after=10)

    # document control
    rows = [
        ("Nomor dokumen", DOC_NO), ("Revisi", REV), ("Tanggal", DATE),
        ("Status", "Draf kerja — bahan rapat koordinasi 2 Oktober 2026, bukan dokumen pengadaan final"),
        ("Disiapkan oleh", "PT LAPI Ganesha Utama, bersama Lab IoT & Fisika Institut Teknologi Bandung"),
        ("Ditujukan kepada", "PT Pertamina Patra Niaga — RU IV Cilacap (pelaksana instalasi, termasuk kontraktor yang ditunjuk RU IV)"),
    ]
    p_ = doc.add_paragraph()
    r = p_.add_run("Document Control")
    r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = NAVY
    mt = doc.add_table(rows=0, cols=2)
    mt.style = "Table Grid"
    for k, v in rows:
        row = mt.add_row().cells
        set_cell_shading(row[0], HEAD_SHADE)
        row[0].paragraphs[0].paragraph_format.space_after = Pt(2)
        r0 = row[0].paragraphs[0].add_run(k)
        r0.font.bold = True; r0.font.size = Pt(9.3); r0.font.color.rgb = GRAY
        row[1].paragraphs[0].paragraph_format.space_after = Pt(2)
        r1 = row[1].paragraphs[0].add_run(v)
        r1.font.size = Pt(9.6)
    fix_widths(mt, [1.9, 5.0])
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    para(doc, "Cara membaca status", bold=True, color=NAVY, after=3)
    table(doc, ["Status", "Arti"], [
        ["FINAL", "Angka sudah tetap dari desain/keputusan proyek — boleh dipakai langsung untuk pengadaan."],
        ["USULAN", "Usulan teknis LGU berdasarkan praktik umum industri — boleh dipakai untuk estimasi, "
                   "disesuaikan bila standar RU IV berbeda."],
        ["KONFIRMASI", "Bergantung kondisi lapangan/keputusan RU IV — perlu dijawab di rapat sebelum pengadaan."],
    ], [1.2, 5.7], status_col=0)

    # header/footer
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run(HEADER + "  |  Draf Kerja")
    hr.font.size = Pt(8); hr.font.color.rgb = GRAY; hr.font.italic = True
    fp = sec.footer.paragraphs[0]
    fp.add_run(f"{DOC_NO}  ·  Rev. {REV}")
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.9), alignment=2)
    fp.add_run("\t")
    add_field(fp, "PAGE", "1")
    fp.add_run(" dari ")
    add_field(fp, "NUMPAGES", "1")
    for r in fp.runs:
        r.font.size = Pt(8); r.font.color.rgb = GRAY

    doc.add_page_break()

    # 1. ringkasan
    heading(doc, "1. Ringkasan — yang perlu disiapkan")
    table(doc, ["Kebutuhan", "Spesifikasi singkat", "Jumlah (perencanaan)", "Status"], [
        ["Tiang GLD", "Pipa baja galvanis 2\" Sch40 (OD 60,3 mm), panjang 1,6 m (1,0 m di atas tanah), ditanam & dicor", f"{N_GLD} batang", "USULAN"],
        ["Tiang Cluster Head", "Pipa baja galvanis 2\" Sch40 (OD 60,3 mm), panjang 4,0 m, ditanam & dicor", f"{N_CH} batang*", "USULAN"],
        ["Mast antena Gateway", "Pipa baja galvanis 2\" Sch40, panjang 6 m (± 4,8 m di atas tanah), ditanam & dicor di dekat ruang Gateway", f"{N_GW} batang", "USULAN"],
        ["Pondasi cor", "GLD 40×40×60 cm · CH 50×50×100 cm · mast GW 50×50×120 cm, beton fc' ≥ 20 MPa", f"{N_GLD + N_CH + N_GW} titik*", "USULAN"],
        ["Catu daya GLD", "24 VDC, kapasitas ≥ 1 A per unit (konsumsi aktual ≈ 0,33 A)", f"{N_GLD} unit beban", "FINAL"],
        ["Kabel daya 24 VDC", "2 inti tembaga; ukuran mengikuti jarak (Tabel 4.2): 1,5–6 mm²", "Σ jarak + cadangan", "KONFIRMASI"],
        ["Kabel antena Gateway", "Koaksial 50 Ω low-loss (LMR-400 atau setara), ≤ 15 m, + penangkal petir koaksial", "1 jalur", "USULAN"],
        ["Kabel jaringan", "Cat6 (router ↔ PC server, PC ↔ intranet kantor)", "2 jalur", "USULAN"],
        ["PC server", "PC fisik di ruang server lokal (bersebelahan dengan ruang Gateway), 2 port LAN", "1 unit", "FINAL"],
        ["Klem U-bolt & pelat", "U-bolt 2\"/DN50 thread M10, pelat 250×250 mm (desain LGU), untuk GLD", "2 U-bolt per titik GLD", "FINAL"],
        ["Bracket panel surya CH", "Bracket baja 2 panel (milik LGU), diklem ke tiang 2\", tanpa las/bor — dibawa LGU", "1 set per CH", "FINAL"],
        ["Grounding tiang", "Bonding ke sistem grounding RU IV", f"{N_GLD + N_CH + N_GW} titik*", "KONFIRMASI"],
    ], [1.35, 3.35, 1.25, 0.95], status_col=3)
    para(doc, f"* Unit Cluster Head (CH) yang tersedia lebih dari 5. Dokumen ini memakai **{N_CH} titik CH "
         "sebagai angka perencanaan**; jumlah final mengikuti penempatan di lapangan (prinsipnya CH "
         "sesedikit mungkin selama seluruh GLD terjangkau). CH memakai panel surya + baterai, "
         "**tidak memerlukan kabel daya**.", size=9.2, color=GRAY)

    # 2. konfigurasi
    heading(doc, "2. Konfigurasi sistem yang dipasang")
    table(doc, ["Perangkat", "Jumlah", "Sumber daya", "Penempatan"], [
        ["GLD (Node Sensor)", f"{N_GLD} unit", "24 VDC melalui kabel", "Titik deteksi di area aman perimeter SRU"],
        ["Cluster Head (CH)", f">5 tersedia (rencana {N_CH})", "Panel surya (2 panel) + baterai, tanpa kabel", "Di antara GLD dan Gateway, membentuk jaringan LoRa mesh"],
        ["Gateway (GW)", f"{N_GW} unit", "220 VAC via adaptor (rating dari tim LGU, lihat §5)", "Di dalam ruangan, dekat ruang server lokal; antena di luar pada mast tinggi"],
        ["PC server", "1 unit", "220 VAC", "Ruang server lokal; disediakan Pertamina, dikonfigurasi LGU"],
        ["Router lapangan", "1 unit", "220 VAC", "Satu ruangan dengan Gateway; disediakan LGU"],
    ], [1.4, 1.35, 1.95, 2.2])
    para(doc, "Alur data: GLD → (LoRa) → CH → (LoRa mesh) → antena Gateway di mast → (kabel koaksial) → Gateway "
         "di dalam ruangan → (Wi-Fi) → router LGU → (kabel Cat6) → PC server (MQTT broker + dashboard). PC server juga tersambung ke intranet kantor kilang lewat port LAN "
         "kedua agar dashboard dapat diakses dari kantor.", size=9.6)

    # 3. tiang
    heading(doc, "3. Spesifikasi tiang besi (ditanam & dicor)")
    para(doc, "Tiang harus berupa pipa ber-OD 60,3 mm (2\" NPS) karena klem U-bolt pada desain bracket "
         "GLD dirancang untuk diameter ini. Pipa ukuran lain tidak dapat dipakai tanpa mengubah U-bolt.",
         size=9.8)
    heading(doc, "3.1 Tiang GLD", level=2)
    table(doc, ["Parameter", "Spesifikasi", "Status"], [
        ["Material", "Pipa baja karbon galvanis celup panas (hot-dip galvanized), mis. ASTM A53 / API 5L Gr.B", "USULAN"],
        ["Ukuran", "2\" NPS Schedule 40 — OD 60,3 mm, tebal dinding 3,91 mm", "FINAL"],
        ["Panjang batang", "1,6 m = 1,0 m di atas tanah + 0,6 m tertanam di cor. "
                           "Satu batang pipa standar 6 m cukup untuk ketiga tiang GLD (3 × 1,6 m = 4,8 m)", "USULAN"],
        ["Tinggi pasang unit GLD", "Rekomendasi: titik tengah sensor ± 0,5 m dari tanah (unit 290 mm berada ± 0,35–0,65 m). "
                                   "Sisa tiang di atas unit memberi ruang geser bila HSE RU IV menetapkan tinggi lain", "USULAN"],
        ["Ujung atas", "Ditutup pipe cap (las di bengkel atau ulir) agar air hujan tidak masuk", "USULAN"],
        ["Ujung bawah (dalam cor)", "Diberi besi silang/angkur (mis. 2× besi Ø12 mm tembus pipa) agar tiang tidak berputar", "USULAN"],
        ["Pondasi", "Lubang 40 × 40 × 60 cm, beton fc' ≥ 20 MPa (± K-250), permukaan atas dibuat miring menjauhi tiang", "USULAN"],
        ["Masa tunggu", "Unit GLD dipasang setelah beton cukup kuat (umumnya ≥ 3 hari; ikuti praktik pelaksana RU IV)", "USULAN"],
        ["Tegak lurus", "Toleransi vertikal ≤ 1° (cek waterpass saat pengecoran)", "USULAN"],
    ], [1.6, 4.3, 1.0], status_col=2)
    para(doc, "Tinggi sensor menurut jenis gas", bold=True, color=NAVY, after=3)
    para(doc, "Tidak ada satu tinggi yang ideal untuk semua gas. Panduan umum industri (IEC 60079-29-2, "
         "ISA RP12.13 Part II) menentukan tinggi berdasarkan berat jenis gas dan letak sumber bocor:", size=9.6)
    table(doc, ["Jenis gas", "Sifat", "Tinggi sensor ideal"], [
        ["LPG (propana/butana), CO₂", "Jauh lebih berat dari udara, mengendap di tanah/parit", "0,3–0,5 m dari tanah"],
        ["H₂S", "Sedikit lebih berat dari udara, toksik", "0,3–0,6 m untuk deteksi kebocoran; zona napas 1,2–1,8 m untuk proteksi pekerja"],
        ["CO", "Hampir sama dengan udara", "Zona napas 1,5–1,8 m"],
        ["Metana", "Lebih ringan, naik", "± 0,5–1 m di atas sumber bocor"],
        ["H₂", "Sangat ringan, naik cepat", "Di atas sumber bocor / titik tertinggi tempat gas bisa terperangkap"],
    ], [1.7, 2.4, 2.8])
    banner(doc, "Rekomendasi untuk perimeter SRU: titik tengah sensor ± 0,5 m dari tanah",
           "Satu unit GLD membaca beberapa gas sekaligus, sehingga tingginya selalu kompromi. Di perimeter SRU, "
           "gas yang paling relevan (H₂S, LPG) cenderung turun, sehingga posisi rendah paling efektif — sekaligus "
           "tanpa pekerjaan ketinggian. Jangan di bawah ± 0,3 m karena sensor bermesh terbuka bisa terkena cipratan "
           "air, lumpur, dan genangan. Untuk titik yang terutama menyasar H₂/metana, tinggi ditentukan dari posisi "
           "sumber bocor (flange/valve), bukan dari tanah; tiang titik tersebut dibuat lebih panjang. "
           "Tinggi final per titik disepakati bersama HSE RU IV saat penentuan titik (letak sumber bocor & arah angin).")

    heading(doc, "3.2 Tiang Cluster Head", level=2)
    table(doc, ["Parameter", "Spesifikasi", "Status"], [
        ["Material & ukuran", "Sama dengan tiang GLD: pipa baja galvanis 2\" Sch40, OD 60,3 mm", "USULAN"],
        ["Panjang batang", "4,0 m = 3,0 m di atas tanah + 1,0 m tertanam", "USULAN"],
        ["Alasan tinggi", "Antena LoRa ditempatkan di atas halangan (orang, kendaraan, peralatan) untuk jangkauan antar-CH", "USULAN"],
        ["Pondasi", "Lubang 50 × 50 × 100 cm, beton fc' ≥ 20 MPa, angkur silang di ujung bawah", "USULAN"],
        ["Bracket panel surya", "Bracket baja 2 panel **disediakan dan dibawa LGU**, diklem ke tiang 2\" (OD 60,3 mm) "
                                "tanpa las/bor; dua panel dipasang berlawanan arah seperti Gambar 3.1", "FINAL"],
        ["Mounting CH & antena", "Unit CH dan antena dipasang di bagian atas tiang (perangkat & klem dari LGU); "
                                 "tiang CH tidak memakai pelat U-bolt 250×250 mm", "FINAL"],
        ["Ujung atas tiang", "Tertutup oleh dudukan CH/antena (bukan pipe cap); sambungan dibuat rapat agar air hujan tidak masuk", "FINAL"],
        ["Tinggi final per titik", "Ditentukan tim RF LGU saat penempatan (line of sight antar-CH)", "KONFIRMASI"],
    ], [1.6, 4.3, 1.0], status_col=2)
    ph = doc.add_paragraph()
    ph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ph.paragraph_format.keep_with_next = True
    for fn in ("ch_bracket_solar_detail.jpg", "ch_mast_antena_atap.jpg"):
        ph.add_run().add_picture(os.path.join(REPO, "scripts", "assets", "ch_mount_photos", fn), height=Inches(2.9))
        ph.add_run("    ")
    para(doc, "Gambar 3.1 Konfigurasi Cluster Head: bracket 2 panel surya diklem ke tiang (kiri), unit CH & antena "
              "di bagian atas tiang (kanan). Foto uji LGU & ITB; drum hanya dudukan sementara — di RU IV tiang "
              "ditanam & dicor.", size=8.8, color=GRAY, italic=True)

    heading(doc, "3.3 Mast antena Gateway", level=2)
    para(doc, "Gateway ditempatkan **di dalam ruangan** di dekat ruang server lokal, tetapi antenanya harus "
         "tinggi karena Gateway adalah titik pusat jaringan mesh: semua CH mengirim data ke antena ini. "
         "Antena dipasang di luar gedung pada mast, lalu disambung ke Gateway dengan kabel koaksial.", size=9.6)
    table(doc, ["Parameter", "Spesifikasi", "Status"], [
        ["Antena", "Antena fiber omni 8 dBi, 920–923 MHz (bawaan perangkat Gateway, disediakan LGU)", "FINAL"],
        ["Tinggi antena", "Puncak antena ≥ 2 m di atas atap gedung terdekat dan lebih tinggi dari antena CH (3 m). "
                          "Target perencanaan ± 5–6 m dari tanah", "USULAN"],
        ["Opsi 1 (disarankan): mast berdiri sendiri", "Pipa baja galvanis 2\" Sch40 panjang 6 m (1 batang standar): ± 4,8 m di atas tanah + "
                                                    "1,2 m tertanam; pondasi 50 × 50 × 120 cm; ditempatkan sedekat mungkin dengan dinding ruang Gateway", "USULAN"],
        ["Opsi 2: mast di dinding/atap gedung", "Pipa 2\" dengan wall bracket/klem ke dinding atau parapet; perlu izin RU IV "
                                               "karena melubangi/menjepit struktur gedung", "KONFIRMASI"],
        ["Penguat mast", "Bila tinggi > 5 m atau angin kencang: tambah guy wire 3 arah atau klem ke dinding gedung", "USULAN"],
        ["Klem antena", "Mengikuti klem bawaan antena (umumnya untuk pipa Ø 30–60 mm)", "USULAN"],
        ["Grounding & petir", "Mast di-bonding ke grounding RU IV; penangkal petir koaksial dipasang di titik kabel masuk gedung", "KONFIRMASI"],
    ], [1.75, 4.15, 1.0], status_col=2)

    heading(doc, "3.4 Grounding & penandaan", level=2)
    bullets(doc, [
        "Setiap tiang logam di-bonding ke sistem grounding RU IV (mis. kabel BC/NYA hijau-kuning 16 mm² + klem pipa "
        "+ sepatu kabel), atau ground rod tersendiri bila titik grounding jauh. **Metode dan nilai resistansi mengikuti "
        "standar RU IV.**",
        "Tiang CH dan mast Gateway lebih tinggi dari struktur sekitarnya: perlu tidaknya proteksi petir tambahan "
        "ditentukan HSE/elektrikal RU IV.",
        "Setiap tiang diberi label nomor titik (GLD-01…03, CH-01…, GW-01) yang sama dengan nomor di dashboard.",
    ])

    # 4. kabel daya
    heading(doc, "4. Catu daya & kabel 24 VDC untuk GLD")
    para(doc, "Hanya GLD yang membutuhkan kabel daya. Setiap GLD bekerja pada **24 VDC** dengan kapasitas "
         "sumber **≥ 1 A per unit**; konsumsi terukur maksimum ≈ 8 W (≈ 0,33 A). Unit GLD tidak pernah "
         "menerima 220 VAC — konversi AC/DC dilakukan PSU di luar unit.", size=9.8)

    heading(doc, "4.1 Dua opsi skema PSU", level=2)
    table(doc, ["", "Opsi A — PSU per titik", "Opsi B — PSU pusat"], [
        ["Konsep", "Tiap GLD punya PSU 220 VAC→24 VDC di box/panel terdekat yang ada stop kontak/220 VAC",
                   "Satu PSU 24 VDC di panel/ruang kontrol; kabel ditarik ke tiap GLD"],
        ["PSU", "24 VDC, ≥ 1 A (disarankan 2,5 A tipe DIN-rail industri), 1 per GLD",
                "24 VDC, ≥ 5 A (3 × 1 A + cadangan), tipe DIN-rail industri"],
        ["Proteksi", "MCB 2 A di sisi 220 VAC + terminal block", "MCB di sisi AC + sekring/MCB DC 2 A per jalur GLD"],
        ["Kabel 24 VDC", "Pendek (umumnya ≤ 50 m) → cukup 2 × 1,5 mm²", "Panjang → ukuran naik sesuai Tabel 4.2"],
        ["Kelebihan", "Kabel kecil & murah, gangguan satu titik tidak memengaruhi titik lain",
                      "Satu titik perawatan, bisa digabung dengan UPS"],
        ["Syarat", "Ada sumber 220 VAC di dekat tiap titik; box PSU di area aman", "Tiap GLD ditarik kabel sendiri (home-run), "
                                                                               "tidak disambung berantai"],
    ], [1.0, 2.95, 2.95])

    heading(doc, "4.2 Ukuran kabel vs panjang maksimum", level=2)
    para(doc, "Panjang maksimum satu arah (PSU → GLD) agar drop tegangan ≤ 5 % (1,2 V dari 24 V). "
         "Kabel tembaga 2 inti, resistansi pada 20 °C; dihitung untuk kapasitas desain 1 A dan untuk konsumsi aktual 0,33 A.",
         size=9.6)
    table(doc, ["Luas penampang", "Resistansi (Ω/km)", "Maks. @ 1 A (desain)", "Maks. @ 0,33 A (aktual)"], [
        ["2 × 1,5 mm²", "12,1", "50 m", "150 m"],
        ["2 × 2,5 mm²", "7,41", "80 m", "245 m"],
        ["2 × 4 mm²", "4,61", "130 m", "390 m"],
        ["2 × 6 mm²", "3,08", "195 m", "590 m"],
    ], [1.6, 1.6, 1.85, 1.85])
    para(doc, "**Rekomendasi:** pilih ukuran berdasarkan kolom 1 A (desain) agar ada cadangan. Bila jarak > 195 m, "
         "gunakan Opsi A (PSU lokal) untuk titik tersebut. Batas tegangan input minimum GLD belum dinyatakan "
         "tim desain; angka 5 % dipakai sebagai batas konservatif.", size=9.6)

    heading(doc, "4.3 Jenis kabel, rute & terminasi", level=2)
    table(doc, ["Parameter", "Spesifikasi", "Status"], [
        ["Konduktor", "2 inti tembaga, berlabel L+ / L− (positif / negatif 24 VDC)", "FINAL"],
        ["Jenis kabel", "Kabel daya/instrumen luar ruang berpelindung armor (mis. SWA), tahan UV & minyak", "USULAN"],
        ["Material selubung & rute", "Mengikuti standar kabel RU IV (ditanam, conduit, atau cable tray existing)", "KONFIRMASI"],
        ["Masuk ke unit GLD", "Cable gland M20 × 1,5 (IP66). Diameter luar kabel harus masuk rentang klem gland; "
                              "kabel armor memerlukan gland tipe armored", "FINAL"],
        ["Terminasi", "Ujung kabel diberi ferrule/skun; sambungan ke konduktor L+/L− di dalam unit dikerjakan teknisi LGU", "FINAL"],
        ["Pelaksana", "Penarikan kabel & PSU oleh RU IV; terminasi akhir, energize & commissioning oleh LGU", "FINAL"],
    ], [1.6, 4.3, 1.0], status_col=2)

    heading(doc, "4.4 Menghitung panjang kabel", level=2)
    para(doc, "Jarak titik GLD ke sumber listrik belum diukur. Panjang kabel per titik dihitung dengan:", size=9.8)
    para(doc, "Panjang kabel = (jarak rute aktual × 1,10) + 3 m cadangan di sisi PSU + 3 m cadangan di sisi GLD",
         bold=True, size=10, after=4)
    para(doc, "Faktor 1,10 menampung belokan dan naik-turun rute; ukur rute kabel, bukan jarak garis lurus. "
         "Contoh perencanaan (asumsi, bukan data lapangan):", size=9.6)
    table(doc, ["Asumsi jarak rute per GLD", "Panjang per titik", f"Total {N_GLD} GLD", "Ukuran kabel (desain 1 A)"], [
        ["20 m (Opsi A, PSU dekat)", "28 m", "84 m → beli 100 m", "2 × 1,5 mm²"],
        ["50 m", "61 m", "183 m → beli 200 m", "2 × 1,5 mm²"],
        ["100 m (Opsi B, PSU pusat)", "116 m", "348 m → beli 400 m", "2 × 4 mm²"],
    ], [1.95, 1.4, 1.75, 1.8])

    # 5. jaringan & server
    heading(doc, "5. Gateway, PC server & jaringan")
    para(doc, "Tata letak: ruang Gateway (Gateway + router LGU) bersebelahan/dekat dengan ruang server lokal "
         "(PC server). Gateway tersambung ke router lewat Wi-Fi, sehingga keduanya sebaiknya satu ruangan; "
         "router ke PC server memakai kabel Cat6.", size=9.6)
    heading(doc, "5.1 Ruang Gateway & ruang server", level=2)
    table(doc, ["Item", "Spesifikasi", "Penyedia", "Status"], [
        ["PC server", "PC fisik, 2 port LAN (NIC kedua bisa USB-LAN), di ruang server lokal. "
                      "Usulan minimum: 4 core / 8 GB RAM / SSD 256 GB", "Pertamina", "USULAN"],
        ["Router lapangan", "Router Wi-Fi milik LGU, satu ruangan dengan Gateway; dikonfigurasi LGU sebelum mobilisasi (lihat §5.2)", "LGU", "FINAL"],
        ["Gateway ↔ router", "Wi-Fi 2,4 GHz, satu ruangan, jarak ± 1–5 m, tanpa sekat logam", "—", "FINAL"],
        ["Kabel router ↔ PC server", "Cat6 sesuai jarak antar-ruang (maks. 90 m per segmen), lewat jalur kabel/conduit gedung", "RU IV", "KONFIRMASI"],
        ["Kabel PC ↔ intranet kantor", "Cat6 ke titik jaringan kantor terdekat; panjang mengikuti lokasi (> 90 m → switch/fiber). "
                                       "Akses hanya untuk API/dashboard", "Pertamina (IT RU IV)", "KONFIRMASI"],
        ["Daya Gateway", "Adaptor 220 VAC dari stop kontak ruang Gateway; tegangan & rating adaptor dikonfirmasi tim LGU", "LGU", "KONFIRMASI"],
        ["Stop kontak 220 VAC", "Ruang Gateway: min. 2 (Gateway, router). Ruang server: min. 3 (PC, monitor, cadangan)", "RU IV", "USULAN"],
        ["UPS", "Disarankan ≥ 1 kVA untuk PC + router + Gateway agar alarm tetap tercatat saat listrik padam", "Pertamina", "USULAN"],
    ], [1.45, 3.2, 1.25, 1.0], status_col=3)

    heading(doc, "5.2 Skema jaringan lokal", level=2)
    table(doc, ["Segmen", "Media & ketentuan", "Status"], [
        ["Gateway → router", "Wi-Fi 2,4 GHz (Gateway berbasis ESP32-S3 hanya mendukung 2,4 GHz), WPA2, SSID khusus jaringan GLD", "FINAL"],
        ["Router → PC server (LAN 1)", "Kabel Cat6 dari ruang Gateway ke ruang server lokal; IP tetap untuk PC server (MQTT broker)", "FINAL"],
        ["PC server → intranet kantor (LAN 2)", "Port LAN kedua ke titik jaringan kantor; hanya untuk akses dashboard/API dari kantor. "
                                                "PC tidak meneruskan (routing) lalu lintas antara jaringan GLD dan intranet", "FINAL"],
        ["Internet", "Tidak ada. Router tanpa WAN/SIM; seluruh sistem berjalan lokal", "FINAL"],
        ["Spesifikasi minimum router", "Wi-Fi 2,4 GHz 802.11n, ≥ 1 port LAN gigabit, catu 220 VAC; disediakan & dikonfigurasi LGU", "FINAL"],
        ["Panjang Cat6 router ↔ PC", "Mengikuti jarak rute antar-ruang (≤ 90 m per segmen); diukur saat survey ruangan", "KONFIRMASI"],
    ], [1.85, 4.05, 1.0], status_col=2)

    heading(doc, "5.3 Kabel koaksial antena Gateway", level=2)
    table(doc, ["Parameter", "Spesifikasi", "Status"], [
        ["Jenis kabel", "Koaksial 50 Ω low-loss, LMR-400 atau setara (tahan UV untuk bagian luar)", "USULAN"],
        ["Panjang", "Sependek mungkin; target ≤ 15 m dari antena ke Gateway. Rumus: rute aktual × 1,10 + 1 m", "USULAN"],
        ["Konektor", "Tipe N di sisi antena & penangkal petir; adaptor/pigtail ke konektor antena Gateway "
                     "(tipe konektor Gateway dikonfirmasi tim LGU)", "KONFIRMASI"],
        ["Penangkal petir", "Coaxial lightning arrester 50 Ω untuk 900 MHz (tipe gas discharge), konektor N, "
                            "dipasang di titik kabel masuk gedung dan di-grounding", "USULAN"],
        ["Kabel masuk gedung", "Lewat lubang/sleeve dinding yang disegel, dengan drip loop (lengkungan turun) sebelum masuk", "USULAN"],
        ["Kedap air", "Semua sambungan konektor di luar ruangan dibalut self-amalgamating tape", "USULAN"],
        ["Radius tekuk", "Tidak ditekuk tajam (LMR-400: radius minimum ± 2,5 cm sekali tekuk, ± 10 cm untuk tekukan berulang)", "USULAN"],
    ], [1.6, 4.3, 1.0], status_col=2)
    para(doc, "Redaman kabel pada 921 MHz (nilai katalog umum, belum termasuk ± 0,5 dB rugi konektor). "
         "Setiap 3 dB redaman memangkas setengah daya sinyal dari dan ke seluruh CH, jadi kabel harus pendek dan berkualitas:",
         size=9.6)
    table(doc, ["Panjang kabel", "LMR-240 (≈ 0,25 dB/m)", "LMR-400 (≈ 0,13 dB/m)", "LMR-600 (≈ 0,08 dB/m)"], [
        ["5 m", "1,2 dB", "0,6 dB", "0,4 dB"],
        ["10 m", "2,5 dB", "1,3 dB", "0,8 dB"],
        ["15 m", "3,7 dB", "1,9 dB", "1,2 dB"],
        ["25 m", "6,2 dB", "3,2 dB", "2,1 dB"],
    ], [1.6, 1.75, 1.75, 1.8])
    para(doc, "**Rekomendasi:** LMR-400 sampai 15 m. Bila rute lebih dari 20 m, pakai LMR-600 atau pindahkan "
         "Gateway ke ruangan yang lebih dekat dengan mast. Hindari kabel RG-58 (redaman ± 0,5 dB/m).", size=9.6)

    # 6. fastener
    heading(doc, "6. Bracket, fastener & aksesori per titik")
    table(doc, ["Item", "Spesifikasi", "Per titik", "Status"], [
        ["U-bolt", "2\"/DN50, thread M10, lebar dalam 62–65 mm, tinggi dalam 95–100 mm, panjang ulir 40 mm", "2 buah", "FINAL"],
        ["Mur & ring", "M10 — mur + ring datar + ring per, material anti-karat (disarankan SS316 karena dekat laut)", "4 set", "USULAN"],
        ["Pelat mounting", "250 × 250 mm, dibuat pelaksana RU IV sesuai gambar desain LGU", "1 buah", "FINAL"],
        ["Spacer", "45–60 mm (bila perlu)", "sesuai kebutuhan", "FINAL"],
        ["Kabel ties", "Tahan UV, stainless atau nylon UV", "± 10 buah", "USULAN"],
        ["Larangan", "Tanpa PVC pada housing/bracket; tanpa bor/las pada struktur existing kilang", "—", "FINAL"],
    ], [1.3, 3.65, 1.0, 0.95], status_col=3)
    banner(doc, "Perlu dicek sebelum fabrikasi pelat",
           "Desain pelat U-bolt saat ini secara visual paling sesuai untuk pipa horizontal (handrail). Karena "
           "seluruh titik GLD memakai tiang baru yang berdiri vertikal, kecocokan orientasi pelat pada tiang vertikal "
           "perlu dicek tim desain LGU bersama pelaksana RU IV sebelum pelat difabrikasi.", bg="FBE9E9")

    # 7. BoQ
    heading(doc, "7. Daftar kebutuhan (BoQ) perencanaan")
    para(doc, f"Basis: {N_GLD} GLD + {N_CH} CH + {N_GW} Gateway. Angka CH dan panjang kabel adalah asumsi perencanaan; "
         "sesuaikan setelah titik final ditetapkan.", size=9.6)
    n_pole = N_GLD                 # tiang yang memakai pelat U-bolt (CH memakai bracket surya LGU)
    n_all = N_GLD + N_CH + N_GW    # semua tiang/mast yang dicor & di-grounding
    beton = N_GLD * 0.096 + N_CH * 0.25 + N_GW * 0.30
    beton_s = f"{beton:.1f}".replace(".", ",")
    table(doc, ["No", "Item", "Spesifikasi", "Qty", "Penyedia"], [
        ["1", "Pipa tiang GLD", "Galvanis 2\" Sch40, 1,6 m (dipotong dari 1 batang 6 m)", f"{N_GLD} batang", "RU IV"],
        ["2", "Pipa tiang CH", "Galvanis 2\" Sch40, 4,0 m", f"{N_CH} batang", "RU IV"],
        ["3", "Pipa mast antena Gateway", "Galvanis 2\" Sch40, 6,0 m (+ guy wire bila perlu)", f"{N_GW} batang", "RU IV"],
        ["4", "Pipe cap 2\"", "Penutup ujung atas tiang GLD & mast GW (tiang CH tidak)", f"{N_GLD + N_GW} buah", "RU IV"],
        ["5", "Beton cor", f"GLD {N_GLD}×0,10 + CH {N_CH}×0,25 + GW {N_GW}×0,30 m³ ≈ {beton_s} m³ "
                          "(belum termasuk susut/sisa)", "± 2,5 m³", "RU IV"],
        ["6", "U-bolt 2\" M10 + mur/ring", "Lihat §6 (GLD saja)", f"{2*n_pole} U-bolt, {4*n_pole} set mur", "RU IV"],
        ["7", "Pelat mounting 250×250 mm", "Sesuai gambar LGU", f"{n_pole} buah", "RU IV"],
        ["8", "Bracket panel surya CH", "Bracket baja 2 panel, klem ke tiang 2\" (Gambar 3.1)", f"{N_CH} set", "LGU"],
        ["9", "PSU 24 VDC", "Opsi A: 3 × (≥1 A) · Opsi B: 1 × (≥5 A)", "3 atau 1 unit", "RU IV"],
        ["10", "Box/panel PSU + MCB + terminal", "Untuk area aman", "3 atau 1 set", "RU IV"],
        ["11", "Kabel daya 2 inti", "Ukuran & panjang per §4", "Σ per §4.4", "RU IV"],
        ["12", "Kabel koaksial LMR-400 + konektor N", "Antena GW → Gateway, ≤ 15 m (§5.3)", "1 jalur", "RU IV"],
        ["13", "Penangkal petir koaksial", "50 Ω, 900 MHz, konektor N (§5.3)", "1 buah", "RU IV"],
        ["14", "Kabel grounding + klem", "Mengikuti standar RU IV (tiang, mast, penangkal petir)", f"{n_all} titik + 1", "RU IV"],
        ["15", "PC server + UPS", "Lihat §5.1", "1 set", "Pertamina"],
        ["16", "Kabel Cat6", "Router ↔ PC server, PC ↔ intranet kantor", "2 jalur", "RU IV / IT RU IV"],
        ["17", "Router lapangan", "Lihat §5.2", "1 unit", "LGU"],
        ["18", "Unit GLD, CH, Gateway + antena, panel surya CH", "Perangkat sistem", f"{N_GLD} / {N_CH} / {N_GW}", "LGU"],
    ], [0.4, 1.85, 2.5, 1.15, 1.0])

    # 8. pertanyaan rapat
    heading(doc, "8. Yang perlu diputuskan di rapat 2 Oktober")
    table(doc, ["No", "Pertanyaan untuk RU IV", "Dampak bila belum dijawab"], [
        ["1", "Titik/koordinat final 3 GLD di perimeter SRU, dan jarak rute ke sumber 220 VAC terdekat", "Panjang & ukuran kabel tidak bisa dipesan"],
        ["2", "Pilih Opsi A (PSU per titik) atau Opsi B (PSU pusat)", "Jumlah PSU & ukuran kabel"],
        ["3", "Standar kabel & rute RU IV (ditanam / conduit / cable tray), termasuk material selubung", "Jenis kabel & cable gland"],
        ["4", "Standar & titik grounding, serta perlu tidaknya proteksi petir untuk tiang CH dan mast Gateway", "Item grounding di BoQ"],
        ["5", "Izin galian & pengecoran di area perimeter SRU, serta dokumen klasifikasi area tertulis", "Jadwal mulai pekerjaan sipil"],
        ["6", "Ruang Gateway & ruang server lokal: lokasi, jarak antar-ruang, 220 VAC, UPS, titik intranet kantor", "Panjang Cat6 & penempatan router"],
        ["7", "Lokasi mast antena Gateway (berdiri sendiri vs di dinding/atap) & jarak rute koaksial ke ruang Gateway", "Panjang & tipe kabel koaksial"],
        ["8", "Tim/kontraktor pelaksana RU IV & jadwal fabrikasi/pemasangan", "Jadwal instalasi"],
    ], [0.4, 4.0, 2.5])

    heading(doc, "9. Batasan dokumen")
    bullets(doc, [
        "Angka berstatus USULAN berasal dari praktik umum industri, bukan dari kajian struktur/elektrikal formal. "
        "RU IV dapat menyesuaikan dengan standar yang berlaku di kilang.",
        "Pembagian kerja: RU IV mendukung penuh instalasi, yaitu penyediaan material, pekerjaan sipil, "
        "kelistrikan, dan pemasangan fisik (termasuk pekerjaan ketinggian & K3) oleh tim/kontraktor RU IV. "
        "LGU menyediakan perangkat sistem, basis desain, supervisi teknis/QA, serta terminasi elektrikal, "
        "energize, dan commissioning.",
        "Instalasi permanen hanya di lokasi yang tidak memerlukan ATEX (area aman), sesuai kesepakatan sebelumnya.",
    ], size=9.6)

    doc.save(OUT)
    print("written", OUT)


if __name__ == "__main__":
    build()
