"""Build the two sample-handover letters (GTS Shanghai, EN; Bea Cukai, ID).

Placeholders that still need real data are written in [square brackets] and
highlighted yellow, so they are easy to find and must not be sent unfilled.

Usage: python scripts/build_surat_sampel_gts_bc.py
Outputs: Deliverables/Surat_Pengantar_Sampel_GLD_ke_GTS_Shanghai.docx
         Deliverables/Surat_Keterangan_Sampel_GLD_ke_Bea_Cukai.docx
"""
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "..", "Deliverables")

LGU_NAME = "PT LAPI GANESHA UTAMA"
LGU_ADDR = "Jl. Dederuk No. 30, Sadang Serang, Bandung, West Java"
GALAKSI = "PT Galaksi Megatama Indonesia"
GTS = "Shanghai Global Testing Services Co., Ltd."

PH = re.compile(r"(\[[^\]]+\])")


def new_doc():
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.5)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.12
    return d


def para(d, text, bold=False, align=None, size=None, after=None, indent=None):
    p = d.add_paragraph()
    for part in PH.split(text):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if PH.fullmatch(part):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "j":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def table(d, rows, widths=(5.2, 10.8)):
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for k, v in rows:
        c = t.add_row().cells
        c[0].width, c[1].width = Cm(widths[0]), Cm(widths[1])
        c[0].paragraphs[0].add_run(k).bold = True
        p = c[1].paragraphs[0]
        for part in PH.split(v):
            if not part:
                continue
            r = p.add_run(part)
            if PH.fullmatch(part):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(10)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def numbered(d, items, indent=0.6):
    for i, t in enumerate(items, 1):
        p = para(d, f"{i}.\t{t}", align="j", indent=indent + 0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(indent + 0.6))


def letterhead(d):
    para(d, LGU_NAME, bold=True, size=14, after=0)
    para(d, LGU_ADDR + ", Indonesia", size=9.5, after=0)
    para(d, "Telephone: [●]  ·  E-mail: [●]", size=9.5, after=8)


# --------------------------------------------------------------------------
# 1. Letter to GTS Shanghai (English)
# --------------------------------------------------------------------------
def build_gts():
    d = new_doc()
    letterhead(d)
    para(d, "Reference No.: [LGU/…/…/2026]", after=0)
    para(d, "Bandung, [date of issue] 2026", after=10)

    para(d, "To:", after=0)
    para(d, GTS, bold=True, after=0)
    para(d, "[Registered address of GTS in Shanghai]", after=0)
    para(d, "People’s Republic of China", after=10)

    para(d, "Subject: Letter of Introduction and Equipment Sample Handover "
            "for ATEX Certification Testing and Evaluation", bold=True, after=8)

    para(d, "Dear Sir/Madam,")
    para(d, "PT LAPI Ganesha Utama (“LGU”), Indonesia, hereby confirms that "
            "its authorised representative named below is carrying the "
            "equipment sample described in this letter to the People’s Republic of China "
            f"for handover to {GTS} (“GTS”).", align="j")
    para(d, "The sample relates to the ATEX/IECEx certification application for the "
            f"Gas Leak Detector GLD V2 submitted to GTS by {GALAKSI} as applicant and "
            "manufacturer. The equipment is intended exclusively for technical evaluation, "
            "testing, and coordination of the applicable assessment and certification "
            "process relating to ATEX/IECEx requirements.", align="j")

    para(d, "A. Equipment Details", bold=True, after=2)
    table(d, [
        ("Equipment name", "Gas Leak Detector (fixed point-type gas detector), GLD V2"),
        ("Manufacturer", f"{GALAKSI}, Indonesia"),
        ("Model / Type", "GLD V2"),
        ("Serial number", "GLD2-0x1001"),
        ("Quantity", "1 (one) unit"),
        ("Country of origin", "Indonesia"),
        ("Description", "Electronic gas-sensing instrument in a metal (aluminium "
                        "alloy ADC12) enclosure, approx. 200 × 90 × 290 mm, approx. "
                        "2.4 kg. Supply 24 VDC, max. 8 W. Integrated LoRa radio "
                        "(920–923 MHz). Contains no gas, gas cylinders or chemical "
                        "substances."),
        ("Battery", "Lithium-ion, 7 × 18650 cells in parallel (1S7P, 3.7 V "
                    "nominal), installed in the equipment (UN 3481, lithium ion "
                    "batteries contained in equipment). Capacity: [… mAh per cell "
                    "→ … Wh total]. The equipment is switched off and protected "
                    "against short circuit and accidental activation during "
                    "transport. UN 38.3 test summary: [attached / to be obtained]."),
        ("HS code (proposed)", "[9027.10 — to be confirmed with customs/forwarder]"),
        ("Declared customs value", "[USD … — value that can be justified, e.g. "
                                   "production cost of the prototype]"),
        ("Receiving company", f"{GTS}"),
        ("Contact person at GTS", "[Name, telephone, e-mail]"),
    ])

    para(d, "B. Authorised Representative Carrying the Sample", bold=True, after=2)
    table(d, [
        ("Name", "[Full name as in passport]"),
        ("Position", "[Position at LGU]"),
        ("Passport number", "[●]"),
        ("Flight / travel date", "[Flight no. · departure date · arrival airport]"),
    ])

    para(d, "C. Purpose and Conditions of Handover", bold=True, after=2)
    numbered(d, [
        "The equipment is provided solely for technical evaluation, testing, and "
        "activities associated with the ATEX/IECEx certification process.",
        "The equipment is not being transferred pursuant to a commercial sale. Its "
        "declared customs value is stated for customs purposes only and does not "
        "represent a sale price.",
        "The equipment is intended to remain in China for the duration of the agreed "
        "testing and certification activities, subject to applicable import, customs, "
        "and regulatory requirements.",
        f"Ownership of the equipment remains with PT LAPI Ganesha Utama unless otherwise "
        f"agreed in writing by the parties. The certification application itself is held "
        f"by {GALAKSI}.",
        "The receiving company is authorised to receive the equipment and to coordinate "
        "the relevant technical testing, evaluation, and certification activities with "
        "the applicable laboratory or certification body.",
        "Any destructive testing, material alteration, disposal, transfer to another "
        "party, or other disposition of the equipment shall be subject to prior written "
        "agreement between LGU and the receiving company. "
        "[Optional: state whether the sample is to be returned to LGU after testing.]",
        "The receiving company shall coordinate the applicable import and customs "
        "procedures in China, including any requirement relating to the integrated "
        "radio module, and ensure that the equipment is handled in accordance with "
        "applicable laws and regulations.",
    ])

    para(d, "This letter is issued to confirm the purpose and intended handling of the "
            "equipment sample and to support the relevant technical and customs "
            "documentation. Should further information be required, please contact the "
            "undersigned.", align="j")

    para(d, "Enclosures: (1) photographs of the sample; (2) copy of the ATEX application "
            "form; (3) [copy of the representative’s passport]; (4) [UN 38.3 test summary of the battery cells].", size=9.5)

    para(d, "Yours faithfully,", after=0)
    para(d, LGU_NAME, bold=True, after=40)
    para(d, "[Name of Authorised Signatory]", bold=True, after=0)
    para(d, "[Position]", after=0)
    para(d, "[Signature and company stamp]", size=9.5)

    path = os.path.join(OUT, "Surat_Pengantar_Sampel_GLD_ke_GTS_Shanghai.docx")
    d.save(path)
    return path


# --------------------------------------------------------------------------
# 2. Letter to Bea Cukai (Indonesian)
# --------------------------------------------------------------------------
def build_bc():
    d = new_doc()
    para(d, LGU_NAME, bold=True, size=14, after=0)
    para(d, "Jl. Dederuk No. 30, Sadang Serang, Bandung, Jawa Barat, Indonesia",
         size=9.5, after=0)
    para(d, "Telepon: [●]  ·  E-mail: [●]", size=9.5, after=8)

    para(d, "Nomor\t: [Nomor Surat]", after=0)
    para(d, "Lampiran\t: 4 (empat) berkas", after=0)
    para(d, "Perihal\t: Surat Keterangan dan Pengantar Sampel Alat untuk Pengujian "
            "dan Sertifikasi ATEX", after=10)

    para(d, "Bandung, [Tanggal Surat] 2026", after=8)
    para(d, "Kepada Yth.", after=0)
    para(d, "Petugas Direktorat Jenderal Bea dan Cukai", bold=True, after=0)
    para(d, "Kantor Pengawasan dan Pelayanan Bea dan Cukai "
            "[Bandara Internasional Soekarno-Hatta / nama bandara keberangkatan]", after=0)
    para(d, "di tempat", after=10)

    para(d, "Dengan hormat,")
    para(d, "Dengan ini kami menerangkan bahwa perwakilan PT LAPI Ganesha Utama (LGU) "
            "yang identitasnya tercantum di bawah membawa satu unit sampel alat berupa "
            "Gas Leak Detector ke Republik Rakyat Tiongkok untuk keperluan evaluasi "
            "teknis, pengujian, dan proses sertifikasi ATEX/IECEx melalui perusahaan "
            "penerima di Tiongkok.", align="j")

    para(d, "A. Rincian sampel alat", bold=True, after=2)
    table(d, [
        ("Nama barang", "Gas Leak Detector (detektor kebocoran gas tipe titik), GLD V2"),
        ("Merek / Produsen", f"{GALAKSI}, Indonesia"),
        ("Model / Tipe", "GLD V2"),
        ("Nomor seri", "GLD2-0x1001"),
        ("Jumlah", "1 (satu) unit"),
        ("Uraian singkat", "Alat elektronik pendeteksi gas dalam casing logam "
                           "(paduan aluminium ADC12), ±200 × 90 × 290 mm, ±2,4 kg, "
                           "catu daya 24 VDC maks. 8 W, dilengkapi radio LoRa "
                           "(920–923 MHz). Tidak berisi gas, tabung gas, maupun bahan "
                           "kimia."),
        ("Baterai", "Litium-ion, 7 sel 18650 paralel (1S7P, nominal 3,7 V), "
                    "terpasang di dalam alat (UN 3481, baterai litium-ion dalam "
                    "peralatan). Kapasitas: [… mAh per sel → … Wh total]. Alat "
                    "dalam keadaan mati dan terlindung dari hubung singkat serta "
                    "aktivasi tidak sengaja selama perjalanan. Ringkasan uji "
                    "UN 38.3: [dilampirkan / menyusul]."),
        ("Pos tarif HS (usulan)", "[9027.10 — dikonfirmasi ke petugas Bea Cukai/PPJK]"),
        ("Pemilik barang", "PT LAPI Ganesha Utama"),
        ("Nilai pabean", "[USD … — nilai yang dapat dipertanggungjawabkan, mis. biaya "
                         "produksi prototipe]"),
        ("Negara tujuan", "Republik Rakyat Tiongkok"),
        ("Perusahaan penerima", f"{GTS}, [alamat di Shanghai]"),
    ])

    para(d, "B. Identitas pembawa barang", bold=True, after=2)
    table(d, [
        ("Nama", "[Nama lengkap sesuai paspor]"),
        ("Jabatan", "[Jabatan di LGU]"),
        ("Nomor paspor", "[●]"),
        ("Penerbangan / tanggal", "[Nomor penerbangan · tanggal keberangkatan]"),
    ])

    para(d, "C. Keterangan", bold=True, after=2)
    numbered(d, [
        "Sampel alat dibawa bukan untuk tujuan penjualan atau distribusi komersial, "
        "melainkan khusus untuk mendukung evaluasi teknis dan sertifikasi ATEX/IECEx. "
        f"Permohonan sertifikasi diajukan oleh {GALAKSI} selaku pemohon dan produsen.",
        "Sampel alat diserahkan kepada perusahaan penerima di Tiongkok dan tetap berada "
        "di Tiongkok selama kegiatan pengujian dan proses sertifikasi, dengan mengikuti "
        "ketentuan impor, ekspor, dan kepabeanan yang berlaku. "
        "[Pilih salah satu: sampel tidak direncanakan kembali ke Indonesia / sampel "
        "direncanakan kembali ke Indonesia setelah pengujian.]",
        "Kepemilikan alat tetap berada pada PT LAPI Ganesha Utama, kecuali kemudian "
        "disepakati lain secara tertulis oleh para pihak.",
        "Surat ini merupakan keterangan mengenai identitas, kepemilikan, dan tujuan "
        "pembawaan sampel alat untuk mendukung pemeriksaan serta konsultasi prosedur "
        "kepabeanan. Surat ini bukan pengganti pemberitahuan ekspor, izin, atau dokumen "
        "kepabeanan lain yang diwajibkan peraturan yang berlaku; apabila diperlukan, "
        "kami siap melengkapi dokumen tersebut.",
    ])

    para(d, "Lampiran: (1) foto sampel alat; (2) salinan formulir aplikasi sertifikasi "
            "ATEX; (3) salinan paspor pembawa barang; (4) [ringkasan uji UN 38.3 sel baterai].", size=9.5)

    para(d, "Demikian surat keterangan ini kami buat dengan sebenar-benarnya untuk "
            "digunakan sebagaimana mestinya.", align="j")

    para(d, "Hormat kami,", after=0)
    para(d, LGU_NAME, bold=True, after=40)
    para(d, "[Nama Pejabat Berwenang]", bold=True, after=0)
    para(d, "[Jabatan]", after=0)
    para(d, "[Tanda tangan dan stempel perusahaan]", size=9.5)

    path = os.path.join(OUT, "Surat_Keterangan_Sampel_GLD_ke_Bea_Cukai.docx")
    d.save(path)
    return path


if __name__ == "__main__":
    for p in (build_gts(), build_bc()):
        print(os.path.normpath(p))
