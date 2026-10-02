# -*- coding: utf-8 -*-
"""Tanggapan LGU atas desain & layout Pertamina RU IV (bahan rapat teknis 2 Okt 2026).

Output: Deliverables/Tanggapan_Desain_Pemasangan_GLD_RU-IV_Cilacap_2Oktober2026.docx
"""
import os

from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from build_persiapan_instalasi_corporate_docx import (
    make_doc, set_cell_shading, set_cell_border, add_field, set_cell_margins, set_table_borders,
    NAVY, GRAY, WHITE, HEAD_SHADE, ZEBRA_SHADE, OK_BG, OK_BD, WARN_BG, WARN_BD,
)

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(REPO, "Deliverables", "Tanggapan_Desain_Pemasangan_GLD_RU-IV_Cilacap_2Oktober2026.docx")
ASSET = os.path.join(REPO, "scripts", "assets", "ppn_design")
DOC_NO = "LGU/GLD/INSTALASI-TANGGAPAN/2026-001"
REVISION = "1.0"
DOC_DATE = "2 Oktober 2026"

doc, sec = make_doc()

STATUS = {"Sesuai": "EAF4EC", "Sesuaikan": "FFF3DC", "Konfirmasi": "FBE6E4"}


def p(text="", size=10.2, bold=False, italic=False, color=None, space_after=7, align=None):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = para.add_run(part)
        r.font.size = Pt(size)
        r.font.bold = bold or (i % 2 == 1)
        r.font.italic = italic
        if color is not None:
            r.font.color.rgb = color
    return para


def box(text, kind="ok", label=None):
    bg, bd = {"ok": (OK_BG, OK_BD), "warn": (WARN_BG, WARN_BD)}[kind]
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_shading(cell, bg)
    set_cell_border(cell, bd, sz=6)
    set_cell_margins(cell, top=110, bottom=110, start=130, end=130)
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    if label:
        r = para.add_run(label + " ")
        r.font.bold = True
        r.font.size = Pt(9.6)
    r = para.add_run(text)
    r.font.size = Pt(9.6)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def table(headers, rows, widths, font=8.8, status_col=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, h in enumerate(headers):
        c = tbl.rows[0].cells[i]
        set_cell_shading(c, "1B2A4A")
        c.paragraphs[0].paragraph_format.space_after = Pt(1)
        r = c.paragraphs[0].add_run(h)
        r.font.bold = True; r.font.size = Pt(8.6); r.font.color.rgb = WHITE
    for ri, row in enumerate(rows):
        cells = tbl.add_row().cells
        for ci, val in enumerate(row):
            c = cells[ci]
            if status_col is not None and ci == status_col:
                set_cell_shading(c, STATUS[val])
            elif ri % 2 == 1:
                set_cell_shading(c, ZEBRA_SHADE)
            for li, line in enumerate(str(val).split("\n")):
                para = c.paragraphs[0] if li == 0 else c.add_paragraph()
                para.paragraph_format.space_after = Pt(1)
                r = para.add_run(line)
                r.font.size = Pt(font)
                if ci == 0 or (status_col is not None and ci == status_col):
                    r.font.bold = True
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    for row in tbl.rows:
        rp = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); rp.append(cs)
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    set_table_borders(tbl)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def figure(fn, width, caption):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(os.path.join(ASSET, fn), width=Inches(width))
    p(caption, size=8.8, italic=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER)


# ---------------------------------------------------------------- cover
lt = doc.add_table(rows=1, cols=1)
lc = lt.rows[0].cells[0]
set_cell_shading(lc, "1B2A4A")
lc.paragraphs[0].paragraph_format.space_after = Pt(2)
r = lc.paragraphs[0].add_run("PT LAPI GANESHA UTAMA")
r.font.bold = True; r.font.size = Pt(14); r.font.color.rgb = WHITE
lp = lc.add_paragraph()
lp.paragraph_format.space_after = Pt(3)
r = lp.add_run("Bekerja sama dengan Lab IoT & Fisika Institut Teknologi Bandung")
r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xC7, 0xD2, 0xE0)
doc.add_paragraph().paragraph_format.space_after = Pt(8)

title = doc.add_heading(level=0)
title.paragraph_format.space_after = Pt(4)
tr = title.add_run("Tanggapan atas Desain & Layout Pemasangan GLD")
tr.font.size = Pt(19); tr.font.bold = True; tr.font.color.rgb = NAVY
p("RU IV Cilacap — bahan rapat teknis penetapan titik & kebutuhan instalasi", size=12, color=GRAY, space_after=10)

mt = doc.add_table(rows=0, cols=2)
mt.style = "Table Grid"
for k, v in [
    ("Nomor dokumen", DOC_NO), ("Revisi", REVISION), ("Tanggal", DOC_DATE),
    ("Ditujukan kepada", "Tim PT Pertamina Patra Niaga — RU IV Cilacap"),
    ("Menanggapi", "Block diagram lingkup kerja, layout titik pemasangan, dan MoM Koordinasi Persiapan "
                   "Pemasangan GLD 29 September 2026 dari Pertamina RU IV"),
    ("Dokumen LGU terkait", "Daftar Perangkat & Kebutuhan Pemasangan rev 1.2; Spesifikasi Material Instalasi rev 1.5"),
]:
    row = mt.add_row().cells
    set_cell_shading(row[0], HEAD_SHADE)
    row[0].paragraphs[0].paragraph_format.space_after = Pt(2)
    r0 = row[0].paragraphs[0].add_run(k)
    r0.font.bold = True; r0.font.size = Pt(9.3); r0.font.color.rgb = GRAY
    row[1].paragraphs[0].paragraph_format.space_after = Pt(2)
    r1 = row[1].paragraphs[0].add_run(v)
    r1.font.size = Pt(9.6)
    row[0].width = Inches(1.9); row[1].width = Inches(4.6)
doc.add_paragraph().paragraph_format.space_after = Pt(6)

hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
hr = hp.add_run("Tanggapan Desain & Layout Pemasangan GLD — RU IV Cilacap")
hr.font.size = Pt(8); hr.font.color.rgb = GRAY; hr.font.italic = True
fp = sec.footer.paragraphs[0]
fp.add_run(f"{DOC_NO}  ·  Rev. {REVISION}  ·  {DOC_DATE}  ·  Halaman ")
add_field(fp, "PAGE", "1")
fp.add_run(" dari ")
add_field(fp, "NUMPAGES", "1")
for run in fp.runs:
    run.font.size = Pt(8); run.font.color.rgb = GRAY

# ---------------------------------------------------------------- 1
doc.add_heading("1. Ringkasan", level=1)
p("LGU & ITB telah menerima block diagram lingkup kerja, layout titik pemasangan, dan MoM rapat 29 September "
  "2026 dari Pertamina RU IV. Secara umum desain tersebut **sejalan** dengan arsitektur sistem dan pembagian "
  "lingkup yang diusulkan LGU: perangkat GLD, Cluster Head, panel surya, dan Gateway oleh LGU; tiang/penyangga, "
  "junction box, PSU, kabel daya, server, UPS, dan jaringan IT oleh Pertamina.")
p("Dokumen ini memetakan desain Pertamina terhadap dokumen LGU, menandai butir yang **sesuai**, butir yang "
  "**akan disesuaikan** LGU, dan butir yang **perlu dikonfirmasi** dalam rapat teknis.")

# ---------------------------------------------------------------- 2
doc.add_heading("2. Desain dari Pertamina RU IV", level=1)
figure("block_diagram_scope.jpg", 6.4, "Gambar 1. Block diagram lingkup kerja LGU dan Pertamina (sumber: Pertamina RU IV).")
figure("layout_titik.jpg", 6.4, "Gambar 2. Layout titik pemasangan: GLD-1, GLD-2, JB-1, JB-2, CH-1–3, GW-1 "
       "(sumber: Pertamina RU IV).")

# ---------------------------------------------------------------- 3
doc.add_page_break()
doc.add_heading("3. Pemetaan terhadap Dokumen LGU", level=1)
table(["Item", "Desain Pertamina", "Dokumen LGU", "Status", "Catatan / tindak lanjut"], [
    ["Jumlah titik", "2 GLD, 3 CH, 1 Gateway, 2 junction box",
     "3 GLD + 1 cadangan; CH 6 (angka perencanaan); 1 Gateway", "Sesuaikan",
     "LGU menyesuaikan dokumen ke layout ini. Perlu konfirmasi: unit GLD ke-3 dipasang di titik tambahan atau "
     "menjadi cadangan"],
    ["Tiang / penyangga", "Stanchion 2\", hot-dip galvanized (HDG)", "Pipa galvanis 2\" Sch40, OD 60,3 mm",
     "Sesuai", "Cocok dengan U-bolt 2\"/DN50 pada pelat mounting GLD. Konfirmasi: stanchion dengan base plate "
     "dibaut ke lantai beton (seperti foto GLD-2) atau ditanam & dicor"],
    ["Mounting GLD & Gateway", "Mounting support oleh Pertamina", "Pelat 250 × 250 mm + U-bolt sesuai gambar LGU",
     "Sesuai", "Gambar CAD pelat dari LGU menjadi acuan fabrikasi"],
    ["Catu daya GLD", "Local receptacle 220 VAC → PSU di junction box → 24 VDC ke GLD",
     "Opsi A: PSU per titik di box dekat sumber 220 VAC", "Sesuai",
     "PSU 24 VDC ≥ 1 A per GLD (disarankan 2,5 A tipe DIN-rail) + MCB 2 A sisi AC"],
    ["Junction box", "SS316", "Box PSU di area aman", "Sesuai", "SS316 cocok untuk lingkungan dekat laut"],
    ["Kabel 24 VDC", "2C, 16 AWG (≈ 1,31 mm²)", "2 × 1,5 mm² untuk jarak ≤ 50 m", "Sesuai",
     "16 AWG memenuhi bila JB → GLD ≤ ± 45 m (drop 5 % pada 1 A desain) atau ≤ ± 135 m pada konsumsi aktual "
     "0,33 A. Pada layout, JB berada dekat GLD"],
    ["Cluster Head", "Lingkup LGU, catu daya panel surya; mounting tidak ditandai",
     "Bracket 2 panel surya + unit CH & antena di atas tiang dibawa LGU; tiang 2\", ± 3 m di atas tanah",
     "Konfirmasi", "Siapa yang menyediakan tiang CH (3 titik di perimeter)"],
    ["Gateway", "Gateway di field dengan link RF; mounting support oleh Pertamina; GW-1 di gedung",
     "Unit Gateway di dalam ruangan dekat server; antena omni 8 dBi di mast + koaksial ≤ 15 m", "Konfirmasi",
     "Posisi unit Gateway (dalam/luar gedung), sumber daya Gateway (adaptor 5 V dari 220 VAC — belum ada "
     "di diagram), dan jalur Gateway → server"],
    ["Router jaringan field", "Belum digambar", "Router milik LGU antara Gateway dan PC server, tanpa internet",
     "Sesuaikan", "Ditambahkan pada block diagram (lingkup LGU, di dalam gedung)"],
    ["Server, UPS, jaringan IT", "Lingkup Pertamina, di dalam gedung", "PC server + UPS disediakan Pertamina",
     "Sesuai", "Spesifikasi di MoM (i5, RAM 16 GB, SSD 128–256 GB) memenuhi minimum LGU (4 core, 8 GB, 100 GB)"],
    ["Klasifikasi area", "Area field ditandai \"CL.I DV.I\"",
     "Instalasi permanen pada lokasi non-ATEX di perimeter SRU", "Konfirmasi",
     "Perlu konfirmasi HSE RU IV atas klasifikasi area di tiap titik GLD dan CH (lihat Bagian 4)"],
], [1.0, 1.35, 1.45, 0.75, 2.0], status_col=3)

# ---------------------------------------------------------------- 4
doc.add_heading("4. Catatan Klasifikasi Area", level=1)
box("Block diagram menandai area field sebagai Class I Division 1. Unit GLD saat ini sedang dalam proses "
    "sertifikasi ATEX/IECEx dan belum memiliki sertifikat Ex, sehingga sesuai kesepakatan sebelumnya pemasangan "
    "permanen diarahkan pada lokasi non-ATEX. Penetapan titik GLD-1, GLD-2, dan CH-1–3 menunggu konfirmasi "
    "klasifikasi area dari HSE RU IV. Bila sebagian titik berada di area terklasifikasi, LGU dapat membantu "
    "mencari titik alternatif terdekat yang berada di area aman.", kind="warn", label="Perlu konfirmasi HSE.")

# ---------------------------------------------------------------- 5
doc.add_heading("5. Pertanyaan untuk Rapat Teknis", level=1)
table(["No", "Pertanyaan", "Pihak"], [
    ["1", "Klasifikasi area di titik GLD-1, GLD-2, CH-1, CH-2, CH-3 (apakah termasuk area aman)", "HSE RU IV"],
    ["2", "Unit GLD ke-3: titik pemasangan tambahan atau disimpan sebagai cadangan", "RU IV & LGU"],
    ["3", "Tiang CH (3 titik): penyedia, tinggi (usulan ± 3 m di atas tanah), dan metode dudukan", "RU IV"],
    ["4", "Posisi unit Gateway, sumber 220 VAC untuk adaptornya, serta jalur ke ruang server", "RU IV & LGU"],
    ["5", "Stanchion GLD: base plate dibaut ke beton atau ditanam & dicor; tinggi sensor ± 0,5–1 m dari lantai", "RU IV"],
    ["6", "Jarak kabel JB → GLD per titik (untuk verifikasi 16 AWG)", "RU IV"],
    ["7", "Ruang & titik jaringan untuk PC server, UPS, dan router LGU", "RU IV (IT)"],
], [0.45, 4.6, 1.45])

# ---------------------------------------------------------------- 6
doc.add_heading("6. Tindak Lanjut LGU", level=1)
for t in [
    "Daftar perangkat, spesifikasi, dimensi, dan kebutuhan utilitas (tindak lanjut no. 1 MoM 29 Sep) telah "
    "disampaikan dalam Daftar Perangkat & Kebutuhan Pemasangan rev 1.2 dan Spesifikasi Material Instalasi rev 1.5.",
    "Setelah rapat teknis, LGU memperbarui Spesifikasi Material (jumlah titik, tiang CH, Gateway, router) dan BoQ "
    "sesuai layout final sebagai dasar rekomendasi kebutuhan lapangan dan kontrak jasa pemasangan RU IV.",
    "LGU mendampingi plot posisi perangkat dan penentuan aksesori serta wiring pada sesi teknis bersama perwakilan lokasi.",
]:
    bp = doc.add_paragraph(style="List Bullet")
    bp.paragraph_format.space_after = Pt(3)
    r = bp.add_run(t)
    r.font.size = Pt(10)

doc.save(OUT)
print("written", OUT)
