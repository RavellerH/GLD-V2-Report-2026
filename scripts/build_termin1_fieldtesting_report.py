from pathlib import Path

from docx import Document
from PIL import Image
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "Paket Pertamina" / "01_Pilot_Field_Testing"
OUT_DOCX = OUT_DIR / "Laporan_Pemenuhan_Deliverable_Termin_1_FieldTesting_GLD.docx"

CONTRACT_IMAGE = Path(
    r"C:\Users\HP\AppData\Local\Temp\codex-clipboard-783d6c24-1253-4e73-99bd-c702ded1db56.png"
)
PROTOTYPE_IMAGE = (
    ROOT
    / "Sumber Dokumen"
    / "sertifikasi-atex-gld-v2-2026"
    / "GLD"
    / "3. Motherboard_ModulSensor_PenutupMesh_Casing_Antena_ModulAlarm.jpg"
)
BRACKET_IMAGE = (
    ROOT
    / "Sumber Dokumen"
    / "GLD U Bolt Bracket"
    / "GLD U-Bolt Bracket V2.png"
)

CHARCOAL = "2F4050"
TEAL = "1ABB9C"
LIGHT_TEAL = "E8F7F4"
PALE_GRAY = "F3F3F4"
MID_GRAY = "D9D9D9"
TEXT = "20262B"
MUTED = "5B6670"
WHITE = "FFFFFF"
GREEN = "187A63"
AMBER = "9A6500"


def prepare_image(source, target, max_width=1800, quality=88):
    target.parent.mkdir(parents=True, exist_ok=True)
    with Image.open(source) as image:
        image = image.convert("RGB")
        if image.width > max_width:
            ratio = max_width / image.width
            image = image.resize((max_width, round(image.height * ratio)), Image.Resampling.LANCZOS)
        image.save(target, "JPEG", quality=quality, optimize=True)
    return target


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=10.5, bold=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def format_cell(cell, bold=False, color=TEXT, size=9.1, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows, widths=None, font_size=9.1, status_col=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        cell.text = header
        set_cell_shading(cell, CHARCOAL)
        if widths:
            cell.width = Inches(widths[idx])
        format_cell(cell, bold=True, color=WHITE, size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.text = str(value)
            if widths:
                cell.width = Inches(widths[idx])
            if row_idx % 2:
                set_cell_shading(cell, PALE_GRAY)
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or idx == status_col else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(cell, size=font_size, align=align)
            if idx == status_col:
                for run in cell.paragraphs[0].runs:
                    status = str(value).upper()
                    color = GREEN if "MEMENUHI" in status or "TERVERIFIKASI" in status else AMBER
                    set_run_font(run, size=font_size, bold=True, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_paragraph(doc, text="", bold_lead=None, style=None, align=None, size=10.5):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=size, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=size)
    else:
        r = p.add_run(text)
        set_run_font(r, size=size)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(text)
    set_run_font(r, size=10.2)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    set_run_font(r, size=15 if level == 1 else 12, bold=True, color="000000")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run_font(r, size=8.5, color=MUTED)
    r.italic = True
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)

    for level, size in ((1, 15), (2, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)


def add_document_control(doc):
    rows = [
        ("Nomor dokumen", "LGU-GLD-T1-FIT-2026-001"),
        ("Revisi", "0.2"),
        ("Tanggal", "11 September 2026"),
        ("Status", "Untuk review dan pengesahan Termin 1"),
        ("Disiapkan oleh", "LAPI Ganesha Utama bersama Lab IoT/Instrumentation and Computation ITB"),
        ("Ditujukan kepada", "PT Pertamina Patra Niaga"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for i, (label, value) in enumerate(rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        table.rows[i].cells[0].width = Inches(1.65)
        table.rows[i].cells[1].width = Inches(4.85)
        set_cell_shading(table.rows[i].cells[0], PALE_GRAY)
        format_cell(table.rows[i].cells[0], bold=True, size=9.2)
        format_cell(table.rows[i].cells[1], size=9.2)


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prototype_for_doc = PROTOTYPE_IMAGE
    if PROTOTYPE_IMAGE.exists():
        prototype_for_doc = prepare_image(
            PROTOTYPE_IMAGE,
            ROOT / "tmp" / "termin1_assets" / "prototype_gld.jpg",
            max_width=1600,
        )
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    # A4 is the standard business-paper size for formal project reports in Indonesia.
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    header = section.header.paragraphs[0]
    header.text = "GLD TAHAP 2   |   TERMIN 1 FIELD TESTING"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_font(run, size=8.2, bold=True, color=CHARCOAL)
    add_page_number(section.footer.paragraphs[0])

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LAPORAN PEMENUHAN DELIVERABLE TERMIN 1")
    set_run_font(r, size=24, bold=True, color="000000")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Pengembangan dan Field Testing Sistem Gas Leak Detection Tahap 2")
    set_run_font(r, size=15, bold=False, color=CHARCOAL)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(34)
    r = p.add_run("PENGAJUAN PEMBAYARAN TERMIN 1 SEBESAR 20 PERSEN")
    set_run_font(r, size=11.5, bold=True, color=TEAL)

    add_document_control(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("RUANG LINGKUP PENGAJUAN")
    set_run_font(r, size=9.5, bold=True, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "Detail engineering dan desain, kesiapan komponen, konfigurasi firmware, "
        "factory integration test, serta dokumentasi instalasi dan as-built tahap pabrik/laboratorium"
    )
    set_run_font(r, size=10.5, color=TEXT)

    doc.add_page_break()

    # Executive summary
    add_heading(doc, "1 Ringkasan Eksekutif", 1)
    add_paragraph(
        doc,
        "Kesimpulan. Berdasarkan bukti teknis dan dokumentasi yang tersedia sampai 11 September 2026, "
        "pekerjaan Termin 1 telah mencapai penyelesaian substansial pada tahap rekayasa dan integrasi "
        "laboratorium. Paket ini layak diajukan kepada PT Pertamina Patra Niaga untuk evaluasi, penerimaan, "
        "dan pembayaran Termin 1 sebesar 20 persen.",
        bold_lead="Kesimpulan. ",
    )
    add_paragraph(
        doc,
        "Pemenuhan tersebut ditopang oleh dokumen detail engineering dan arsitektur, kesiapan perangkat, "
        "konfigurasi firmware, pengujian end-to-end GLD-Cluster Head-Gateway-Server, pengujian mesh dan "
        "alarm push, serta peninjauan perkembangan sistem oleh PT Pertamina Patra Niaga di Lab IoT, "
        "Instrumentation and Computations, Gedung Laboratorium Fisika Terpadu, Institut Teknologi Bandung, sebelum "
        "kunjungan RU IV Cilacap. Berdasarkan konfirmasi PIC proyek pada 11 September 2026, kunjungan "
        "Pertamina tersebut merupakan witness terhadap perkembangan dan integrasi sistem untuk tahap FAT/FIT."
    )
    add_paragraph(
        doc,
        "Batas tahap. Pengajuan ini tidak menyatakan bahwa instalasi, commissioning, site acceptance test, "
        "atau sertifikasi hazardous area telah selesai. Status RU IV Cilacap masih survey lokasi. Kegiatan "
        "lapangan tersebut merupakan tahapan sesudah factory integration dan akan dibuktikan pada termin "
        "berikutnya.",
        bold_lead="Batas tahap. ",
    )

    rows = [
        ("1", "Detail engineering dan desain", "MEMENUHI", "Arsitektur, datasheet sistem, TDS R4, desain mounting, serta dokumen HSE dan instalasi tersedia."),
        ("2", "Persiapan komponen", "MEMENUHI", "Empat unit GLD dan enam belas unit Cluster Head telah dilaporkan tersedia untuk pengembangan dan integrasi."),
        ("3", "Konfigurasi firmware", "MEMENUHI", "Akuisisi delapan sensor, inferensi lokal, komunikasi LoRa STAR/MESH, alarm, downlink, dan integrasi gateway-server terdokumentasi."),
        ("4", "Factory integration test", "MEMENUHI DENGAN PENGESAHAN", "Rangkaian integrasi lab dan alarm push telah diuji; PT Pertamina Patra Niaga meninjau sistem di Lab IoT/Instrumentation and Computation ITB sebelum kunjungan Cilacap. Tanda tangan pengesahan disediakan pada laporan ini."),
        ("5", "Dokumen instalasi dan as-built", "MEMENUHI SESUAI TAHAP", "Dokumen instalasi dan konfigurasi as-built prototipe/laboratorium tersedia. As-built site final diterbitkan setelah instalasi lapangan."),
    ]
    add_table(doc, ["No", "Deliverable", "Status", "Dasar penilaian"], rows, [0.35, 1.65, 1.35, 3.15], font_size=8.8, status_col=2)

    doc.add_page_break()
    add_heading(doc, "2 Dasar Kontraktual", 1)
    add_paragraph(
        doc,
        "Termin 1 pada ketentuan pembayaran pekerjaan field testing bernilai 20 persen dari harga SPK. "
        "Milestone ini didasarkan pada pelaksanaan detail engineering, persiapan komponen, konfigurasi "
        "firmware, dan factory integration test, yang dibuktikan melalui dokumen detail engineering dan "
        "desain, dokumen instalasi dan as-built, serta laporan factory acceptance test."
    )
    if CONTRACT_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(CONTRACT_IMAGE), width=Inches(6.45))
        add_caption(doc, "Gambar 1 Ketentuan pembayaran dan deliverable Termin 1 pada dokumen SPK")

    add_paragraph(
        doc,
        "Interpretasi tahap. Struktur termin menempatkan field testing kelompok Unit Kilang pada Termin 2 "
        "dan Termin 3. Karena itu, as-built pada Termin 1 diperlakukan sebagai baseline konfigurasi perangkat "
        "dan integrasi yang telah diuji di pabrik/laboratorium. As-built lokasi dipenuhi setelah instalasi aktual "
        "dan tidak diklaim sebagai bagian yang sudah selesai dalam laporan ini.",
        bold_lead="Interpretasi tahap. ",
    )

    doc.add_page_break()

    # Engineering evidence
    add_heading(doc, "3 Detail Engineering dan Desain", 1)
    add_paragraph(
        doc,
        "Paket engineering telah berkembang dari konsep sistem menjadi konfigurasi teknis yang dapat diuji. "
        "Dokumentasi mencakup aliran data end-to-end, fungsi tiap perangkat, protokol radio, keamanan payload, "
        "mode operasi, antarmuka server, kebutuhan daya, dan konsep pemasangan."
    )
    rows = [
        ("DE-01", "Arsitektur sistem end-to-end", "Datasheet Sistem GLD Arsitektur Server dan Jaringan", "GLD ke CH melalui STAR, antar-CH melalui MESH, GW ke broker dan server."),
        ("DE-02", "Spesifikasi Gas Leak Detector", "Technical Datasheet GasleakDetector Revision 4.0", "Akuisisi 8 sensor MQ, ADS1256, ESP32-S3, LoRa, mode alarm dan downlink."),
        ("DE-03", "Spesifikasi CH, Gateway, Server", "Paket Technical Datasheet Revision 4.0", "Kontrak komunikasi, cache, queue, parent selection, broker, dan pemrosesan server."),
        ("DE-04", "Desain mekanik pemasangan", "Desain Bracket L U-Bolt GLD Mounting", "Basis mounting non-invasif pada struktur existing tanpa las atau bor."),
        ("DE-05", "Dokumen instalasi dan HSE", "JSA HSE RU IV dan Checklist Kesiapan Instalasi", "Prasyarat HSE, pembagian tanggung jawab, mobilisasi, commissioning, dan acceptance test."),
    ]
    add_table(doc, ["ID", "Cakupan", "Dokumen", "Isi yang dibuktikan"], rows, [0.6, 1.45, 2.15, 2.55], font_size=8.6)

    if BRACKET_IMAGE.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(BRACKET_IMAGE), width=Inches(6.3))
        add_caption(doc, "Gambar 2 Basis desain bracket L dan U-bolt untuk pemasangan GLD")

    doc.add_page_break()
    add_heading(doc, "4 Kesiapan Komponen dan Perangkat", 1)
    add_paragraph(
        doc,
        "Notulen rapat 6 Agustus 2026 mencatat kesiapan 4 unit GLD, 9 Cluster Head besar, dan 7 Cluster Head "
        "kecil. Jumlah tersebut menunjukkan bahwa perangkat inti untuk pengujian integrasi telah tersedia. "
        "Untuk pengiriman lapangan, identitas nomor seri dan daftar packing per lokasi akan ditetapkan pada "
        "tahap mobilisasi."
    )
    rows = [
        ("GLD", "4 unit", "Tersedia untuk integrasi dan kebutuhan pilot saat ini"),
        ("Cluster Head besar", "9 unit", "Tersedia untuk pengujian jaringan dan cakupan"),
        ("Cluster Head kecil", "7 unit", "Tersedia untuk konfigurasi jaringan"),
        ("Gateway dan server", "Konfigurasi lab", "Rantai komunikasi end-to-end telah diuji di lingkungan lab"),
        ("Gas test chamber", "Selesai", "Mendukung pengumpulan data dan pengujian sensor di lingkungan terkendali"),
    ]
    add_table(doc, ["Komponen", "Status kuantitas", "Keterangan"], rows, [1.7, 1.35, 3.75], font_size=9.0)

    if prototype_for_doc.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(prototype_for_doc), width=Inches(2.75))
        add_caption(doc, "Gambar 3 Unit GLD terakit dengan enclosure, sensor cartridge, antena, dan modul alarm")

    doc.add_page_break()

    # Firmware and FAT
    add_heading(doc, "5 Konfigurasi Firmware", 1)
    add_paragraph(
        doc,
        "Firmware telah dikonfigurasi untuk menjalankan fungsi utama pada rantai sensor hingga server. "
        "Konfigurasi yang telah dibuktikan di lingkungan laboratorium dirangkum sebagai berikut."
    )
    for item in [
        "Akuisisi delapan kanal sensor gas melalui ADC dan pemrosesan pada ESP32-S3.",
        "Inferensi AI dilakukan secara lokal pada perangkat GLD.",
        "Komunikasi GLD ke serving Cluster Head menggunakan pola STAR dan komunikasi antar-Cluster Head menggunakan MESH.",
        "Payload lapangan dilindungi AES-128-GCM; sesi Gateway ke broker dan Server ke broker menggunakan TLS terpisah.",
        "Telemetri normal disimpan di cache Cluster Head dan diambil melalui mekanisme pull; alarm dan clear dikirim secara push.",
        "Downlink dan receive window endpoint telah didefinisikan dalam Technical Datasheet Revision 4.0.",
    ]:
        add_bullet(doc, item)
    add_paragraph(
        doc,
        "Catatan konfigurasi. Profil dokumentasi model AI pada TDS R4 masih perlu diselaraskan dengan profil "
        "CNN Dual-Branch empat kelas yang telah dikonfirmasi tim. Perbedaan dokumentasi ini tidak membatalkan "
        "fungsi inferensi lokal yang telah berjalan, tetapi harus ditutup sebelum penerbitan dokumen konfigurasi final.",
        bold_lead="Catatan konfigurasi. ",
    )

    add_heading(doc, "6 Pelaksanaan Factory Integration Test", 1)
    add_paragraph(
        doc,
        "Factory Integration Test dilaksanakan sebagai rangkaian pengujian di lingkungan ITB sebelum "
        "kunjungan lapangan RU IV Cilacap. Pada 6 Agustus 2026, perwakilan PT Pertamina Patra Niaga hadir di "
        "Lab IoT, Instrumentation and Computations, Gedung Laboratorium Fisika Terpadu, Institut Teknologi Bandung, "
        "bersama LAPI Ganesha Utama dan Lab IoT/Instrumentation and Computation untuk meninjau perkembangan perangkat, desain mekanik, jaringan komunikasi, "
        "strategi instalasi, serta kesiapan implementasi. Pengujian lanjutan pada rentang 6 sampai 8 Agustus "
        "membuktikan alarm push melalui konfigurasi mesh kampus."
    )
    add_paragraph(
        doc,
        "Kehadiran dan witness. Pihak yang terlibat adalah LAPI Ganesha Utama, Lab IoT/Instrumentation and "
        "Computation ITB, dan PT Pertamina Patra Niaga. Konfirmasi PIC proyek pada 11 September 2026 menegaskan bahwa kunjungan tersebut "
        "merupakan witness perkembangan dan integrasi sistem sebelum visit Cilacap. Lembar pengesahan pada "
        "bagian akhir laporan disediakan untuk mengubah bukti witness tersebut menjadi acceptance formal Termin 1.",
        bold_lead="Kehadiran dan witness. ",
    )

    doc.add_page_break()
    # Keep the continuation table clear of the running header in LibreOffice/PDF.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = Pt(4)
    rows = [
        ("FAT-01", "Akuisisi sensor dan telemetri GLD", "TERVERIFIKASI DI LAB", "Data delapan sensor dapat dibaca dan direkam; dataset konsisten tersedia."),
        ("FAT-02", "Inferensi AI lokal", "TERVERIFIKASI DI LAB", "Inferensi on-device pada ESP32-S3 telah dikonfirmasi dan didukung secara generik oleh TDS R4."),
        ("FAT-03", "Komunikasi GLD ke CH", "TERVERIFIKASI DI LAB", "Link STAR dan payload lapangan berfungsi pada konfigurasi uji."),
        ("FAT-04", "Mesh multi-hop dan failover CH", "TERVERIFIKASI DI LAB", "Topologi multi-hop 8 CH telah diuji untuk menyiasati keterbatasan jangkauan per-hop."),
        ("FAT-05", "CH ke Gateway dan Server", "TERVERIFIKASI DI LAB", "Rantai GLD-CH-GW-Server telah berjalan end-to-end di laboratorium."),
        ("FAT-06", "Alarm push", "TERVERIFIKASI DI LAB", "GLD disemprot LPG dan alarm diterima server otomatis tanpa pull request pada uji 6-8 Agustus."),
        ("FAT-07", "Monitoring dan dokumentasi data", "TERVERIFIKASI DI LAB", "Data dan alarm dapat dipantau pada lingkungan server lab; commissioning produksi belum termasuk tahap ini."),
    ]
    add_table(doc, ["Test ID", "Fungsi", "Hasil", "Bukti hasil"], rows, [0.65, 1.65, 1.55, 2.95], font_size=8.5, status_col=2)

    add_paragraph(
        doc,
        "Hasil FAT. Fungsi inti yang diperlukan untuk melanjutkan ke persiapan lapangan telah ditunjukkan pada "
        "lingkungan laboratorium. Tidak ditemukan kegagalan integrasi yang menghentikan persiapan pilot. "
        "Temuan yang masih terbuka bersifat penyempurnaan dokumentasi, penambahan kelas gas, commissioning "
        "produksi, serta kepatuhan area berbahaya yang ditangani pada tahap berikutnya.",
        bold_lead="Hasil FAT. ",
    )

    doc.add_page_break()

    # Installation / as-built and evidence
    add_heading(doc, "7 Dokumen Instalasi dan As Built", 1)
    add_paragraph(
        doc,
        "Dokumen instalasi yang tersedia telah mencakup konsep mounting, kebutuhan daya 24 VDC, arsitektur "
        "jaringan, kebutuhan server, pembagian tanggung jawab, draft JSA/HSE, dan checklist pra-instalasi. "
        "Konfigurasi yang diuji di laboratorium menjadi baseline as-built factory untuk Termin 1."
    )
    rows = [
        ("As-built perangkat", "Tersedia sesuai tahap", "Foto unit terakit, spesifikasi hardware, TDS, konfigurasi firmware, dan arsitektur integrasi lab."),
        ("Desain pemasangan", "Tersedia", "Bracket L dan U-bolt, konsep struktur existing, tanpa las atau bor."),
        ("Metode dan HSE", "Tersedia sebagai draft", "JSA/HSE dan checklist pra-instalasi; pengesahan spesifik lokasi dilakukan sebelum pekerjaan fisik."),
        ("As-built lokasi RU IV", "Tahap berikutnya", "Diterbitkan setelah instalasi, pengukuran aktual, commissioning, dan SAT. Tidak diklaim selesai pada Termin 1."),
    ]
    add_table(doc, ["Dokumen", "Status", "Batas dan isi"], rows, [1.65, 1.45, 3.65], font_size=9.0, status_col=1)

    add_heading(doc, "8 Register Bukti", 1)
    rows = [
        ("E-01", "Notulensi Kick Off Meeting 12 Juni 2026", "Scope, desain end-to-end, casing, mode operasi, tanggung jawab, dan target implementasi."),
        ("E-02", "Notulen Meeting GLD 6 Agustus 2026", "Witness PT Pertamina Patra Niaga di Lab IoT/Instrumentation and Computation ITB, desain mekanik, arsitektur jaringan, status perangkat, dan tindak lanjut."),
        ("E-03", "Technical Datasheet Revision 4.0", "Kontrak teknis GLD, CH, Gateway, Server, radio, keamanan, cache, alarm, dan downlink."),
        ("E-04", "Datasheet Sistem GLD Arsitektur Server dan Jaringan", "Dokumen detail engineering konsolidasi untuk perangkat, jaringan, server, daya, dan instalasi."),
        ("E-05", "Laporan Progres By Date GLD", "Kronologi pengujian, penyelesaian integrasi, dan perkembangan perangkat."),
        ("E-06", "Test Sinyal LoRa dan log serial", "Data uji komunikasi serta bukti operasi perangkat."),
        ("E-07", "Dataset sensor gas", "Data akuisisi delapan kanal untuk LPG, CO2, udara bersih, dan sesi baseline."),
        ("E-08", "Desain Bracket L U-Bolt GLD Mounting", "Basis desain instalasi mekanik pada struktur existing."),
        ("E-09", "JSA HSE dan Checklist Kesiapan Instalasi RU IV", "Prasyarat HSE, mobilisasi, commissioning, dan acceptance test."),
        ("E-10", "Konfirmasi PIC proyek 11 September 2026", "PT Pertamina Patra Niaga telah hadir di Lab IoT/Instrumentation and Computation ITB untuk menyaksikan perkembangan/integrasi sebelum visit Cilacap."),
    ]
    add_table(doc, ["ID", "Bukti", "Relevansi"], rows, [0.55, 2.55, 3.65], font_size=8.7)

    add_heading(doc, "9 Outstanding dan Tahap Berikutnya", 1)
    add_paragraph(
        doc,
        "Outstanding berikut tidak mengubah kesimpulan substantial completion Termin 1, tetapi wajib "
        "ditindaklanjuti pada tahapan lapangan dan sertifikasi."
    )
    rows = [
        ("1", "Tanda tangan witness dan acceptance FAT", "Administrasi Termin 1", "Tutup melalui lembar pengesahan laporan dan BAST."),
        ("2", "Nomor seri dan packing list per RU", "Mobilisasi", "Tetapkan sebelum pengiriman perangkat."),
        ("3", "Instalasi, commissioning, dan SAT RU IV", "Termin lapangan", "Belum dimulai; dilaksanakan setelah izin dan kesiapan lokasi."),
        ("4", "As-built site final", "Pasca-instalasi", "Susun berdasarkan ukuran, jalur, dan konfigurasi aktual."),
        ("5", "H2S, Benzena, dan kelas gas tambahan", "Pengembangan AI", "Menunggu ketersediaan sampel dan validasi model."),
        ("6", "Sertifikasi hazardous area", "Jalur sertifikasi", "Masih persiapan dokumen; tidak menjadi klaim laporan Termin 1 field testing."),
    ]
    add_table(doc, ["No", "Outstanding", "Tahap", "Tindakan"], rows, [0.35, 2.15, 1.35, 2.9], font_size=8.7)

    doc.add_page_break()

    # Conclusion and sign-off
    add_heading(doc, "10 Kesimpulan dan Rekomendasi Penerimaan", 1)
    add_paragraph(
        doc,
        "Pekerjaan detail engineering, persiapan komponen, konfigurasi firmware, dan factory integration test "
        "telah dilaksanakan secara substansial. Sistem inti telah diuji end-to-end di lingkungan laboratorium, "
        "dan PT Pertamina Patra Niaga telah melakukan peninjauan langsung di Lab IoT, Instrumentation and Computations, "
        "Gedung Laboratorium Fisika Terpadu, Institut Teknologi Bandung, sebelum kunjungan RU IV Cilacap. "
        "Dokumen instalasi serta baseline as-built prototipe/laboratorium juga telah tersedia."
    )
    add_paragraph(
        doc,
        "Rekomendasi. Deliverable Termin 1 direkomendasikan untuk diterima dan diproses sebagai dasar "
        "pembayaran 20 persen dari harga SPK. Penandatanganan lembar pengesahan dan BAST akan menjadi "
        "penutupan administratif atas pekerjaan yang telah dilaksanakan. Instalasi, site acceptance test, "
        "as-built lapangan, serta sertifikasi tetap dilaporkan pada tahap berikutnya.",
        bold_lead="Rekomendasi. ",
    )

    add_heading(doc, "11 Lembar Pengesahan", 1)
    add_paragraph(
        doc,
        "Dengan menandatangani lembar ini, para pihak menyatakan bahwa laporan telah diperiksa dan bahwa "
        "deliverable Termin 1 dapat diterima sebagai dasar proses pembayaran sesuai ketentuan SPK, dengan "
        "tindak lanjut yang tercantum pada Bagian 9."
    )
    rows = [
        ("Disiapkan oleh", "LAPI Ganesha Utama", "Nama dan jabatan\n\n\nTanda tangan dan tanggal"),
        ("Diverifikasi oleh", "Lab IoT/Instrumentation and Computation ITB", "Nama dan jabatan\n\n\nTanda tangan dan tanggal"),
        ("Diterima oleh", "PT Pertamina Patra Niaga", "Nama dan jabatan\n\n\nTanda tangan dan tanggal"),
    ]
    add_table(doc, ["Peran", "Pihak", "Pengesahan"], rows, [1.25, 2.25, 3.25], font_size=9.3)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Keputusan penerimaan")
    set_run_font(r, size=10.5, bold=True)
    for choice in [
        "[  ] Diterima untuk proses pembayaran Termin 1 sebesar 20 persen",
        "[  ] Diterima dengan catatan sebagaimana Bagian 9",
        "[  ] Memerlukan perbaikan sebelum diproses",
    ]:
        add_paragraph(doc, choice, size=10.2)

    doc.add_page_break()
    add_heading(doc, "12 Referensi Dokumen", 1)
    refs = [
        "Ketentuan Cara Pembayaran pada SPK pekerjaan field testing GLD Tahap 2.",
        "Notulensi Kick Off Meeting Pertamina, 12 Juni 2026.",
        "Notulen Meeting GLD, 6 Agustus 2026, Lab IoT, Instrumentation and Computations, Gedung Laboratorium Fisika Terpadu, Institut Teknologi Bandung.",
        "Technical Datasheets Lab IoT ITB Revision 4.0, 4 September 2026.",
        "Datasheet Sistem GLD Arsitektur Server dan Jaringan.",
        "Laporan Progres By Date GLD sampai 31 Agustus 2026.",
        "Desain Bracket L U-Bolt GLD Mounting.",
        "JSA HSE RU IV Cilacap GLD dan Checklist Kesiapan Instalasi RU IV Cilacap.",
        "Konfirmasi PIC proyek, 11 September 2026, mengenai witness PT Pertamina Patra Niaga di Lab IoT/Instrumentation and Computation ITB sebelum visit Cilacap.",
    ]
    for i, ref in enumerate(refs, 1):
        add_paragraph(doc, f"{i}. {ref}", size=9.4)

    doc.core_properties.title = "Laporan Pemenuhan Deliverable Termin 1 Field Testing GLD"
    doc.core_properties.subject = "Pengajuan pembayaran Termin 1 sebesar 20 persen"
    doc.core_properties.author = "LAPI Ganesha Utama bersama Lab IoT/Instrumentation and Computation ITB"
    doc.core_properties.keywords = "GLD, Termin 1, FAT, Factory Integration Test, Pertamina"
    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    build_document()
