"""Versi Word (.docx) Laporan Termin 1 Field Testing rev 0.3 LENGKAP dengan Lampiran A (Laporan FAT),
dibangun di Linux tanpa Word COM (python-docx + docxcompose).

Isi: halaman depan LGU (cover, kontrol dokumen, pengesahan, kata pengantar, daftar isi) + isi laporan
Termin 1 (dapat diedit) + Lampiran A: separator, Laporan FAT (dapat diedit), lampiran bukti FAT 1-6
(gambar halaman dari PDF bukti asli) dan Laporan Uji Lab 01-03 (dapat diedit).

Pakai: python3 scripts/build_word_termin1_ft_lengkap.py
"""
import os
import sys

import docx
import pymupdf
from docx.shared import Mm, Pt
from docxcompose.composer import Composer

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import build_cover_laporan_lgu as cv  # noqa: E402
import build_word_penagihan_lgu as wp  # noqa: E402

ROOT = cv.ROOT
TMP = os.path.join(ROOT, "Paket LGU", "02_Penagihan_Termin_1", "src")
D = cv.DOCS[0]
FINAL_PDF = os.path.join(ROOT, D["out"])
OUT = FINAL_PDF[:-4] + ".docx"
FAT_PDF = os.path.join(ROOT, D["lampiran_fat"])
BODY = os.path.join(ROOT, "Paket Pertamina", "04_Laporan_Termin_1",
                    "Laporan_Pemenuhan_Deliverable_Termin_1_FieldTesting_GLD_Rev03.docx")
FAT_DOCX = os.path.join(ROOT, "Deliverables", "Laporan_FAT_GLD_Tahap2.docx")
LAB_DOCX = {
    "7": os.path.join(ROOT, "Deliverables", "Laporan_Uji_Lab_01_Model_AI_GLD.docx"),
    "8": os.path.join(ROOT, "Deliverables", "Laporan_Uji_Lab_02_Komunikasi_LoRa_GLD.docx"),
    "9": os.path.join(ROOT, "Deliverables", "Laporan_Uji_Lab_03_Mesh_Integrasi_GLD.docx"),
}


def image_pages(pdf, pages, out, tag):
    """Dokumen berisi halaman PDF sebagai gambar, satu halaman A4 per halaman sumber."""
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(s, m, Mm(8))
    s.header_distance = s.footer_distance = Mm(0)
    first = True
    for i in pages:
        pg = pdf[i]
        png = os.path.join(TMP, f"img_{tag}_{i}.jpg")
        pix = pg.get_pixmap(dpi=110)
        pix.save(png, jpg_quality=72)
        w, h = pg.rect.width, pg.rect.height
        maxw, maxh = 194, 280
        scale = min(maxw / (w * 25.4 / 72), maxh / (h * 25.4 / 72))
        p = d.add_paragraph()
        if not first:
            p.paragraph_format.page_break_before = True
        first = False
        p.paragraph_format.space_after = Pt(0)
        p.alignment = 1
        p.add_run().add_picture(png, width=Mm(w * 25.4 / 72 * scale), height=Mm(h * 25.4 / 72 * scale))
    d.save(out)
    return out


def main():
    os.makedirs(TMP, exist_ok=True)
    final = pymupdf.open(FINAL_PDF)
    fat = pymupdf.open(FAT_PDF)
    D["_jml"] = final.page_count
    cover_png = os.path.join(TMP, "cover_01.png")
    final[0].get_pixmap(dpi=200).save(cover_png)

    # daftar isi = daftar isi PDF resmi (bab + Lampiran A + lampiran FAT)
    toc = final.get_toc()
    entries = [(1, t, p) for lvl, t, p in toc if lvl == 1 and p > 5]
    entries += [(2, cv.fat_label(t), p) for lvl, t, p in toc if lvl == 2]
    entries.sort(key=lambda e: (e[2], e[0]))
    front_docx = os.path.join(TMP, "front_01_rev03.docx")
    tempat = cv.TTD["tempat"]
    cv.TTD["tempat"] = D.get("tempat", tempat)
    try:
        wp.front(D, entries, cover_png, front_docx)
    finally:
        cv.TTD["tempat"] = tempat
    f = docx.Document(front_docx)
    for p in f.paragraphs:
        if p.text.startswith("Nomor halaman mengikuti"):
            for r in p.runs:
                r.text = ""
            p.runs[0].text = "Nomor halaman mengacu pada berkas PDF resmi laporan ini."
    f.add_page_break()  # isi laporan mulai di halaman baru
    f.save(front_docx)

    body_docx = os.path.join(TMP, "body_01_rev03.docx")
    wp.clean_body(BODY, body_docx)

    # posisi halaman di PDF gabungan FAT (0-based), 5 halaman depan FAT dilewati
    fat_toc = [(t, p - 1) for lvl, t, p in fat.get_toc() if lvl == 1 and t.startswith("Lampiran ")]
    sep_a = final.page_count - fat.page_count + 5 - 1  # halaman separator Lampiran A di PDF final
    parts = [image_pages(final, [sep_a], os.path.join(TMP, "lampA_sep.docx"), "A"), FAT_DOCX]
    for k, (t, start) in enumerate(fat_toc):
        no = t.split()[1]
        end = fat_toc[k + 1][1] if k + 1 < len(fat_toc) else fat.page_count
        if no in LAB_DOCX:
            parts.append(image_pages(fat, [start], os.path.join(TMP, f"lamp{no}_sep.docx"), f"s{no}"))
            parts.append(LAB_DOCX[no])
        else:
            parts.append(image_pages(fat, list(range(start, end)), os.path.join(TMP, f"lamp{no}.docx"), f"l{no}"))

    master = docx.Document(front_docx)
    comp = Composer(master)
    for path in [body_docx] + parts:
        comp.append(docx.Document(path))
    comp.save(OUT)
    print(OUT, round(os.path.getsize(OUT) / 1e6, 1), "MB")


if __name__ == "__main__":
    main()
