from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
NOTULEN = ROOT / "Deliverables" / "Notulen_Meeting_GLD_6Agustus2026.docx"
BAST = (
    ROOT
    / "Paket Pertamina"
    / "01_Pilot_Field_Testing"
    / "Draft_BAST_Termin_1_FieldTesting_20Persen.docx"
)

CHARCOAL = "2F4050"
TEAL = "1ABB9C"
PALE_TEAL = "E8F7F4"
PALE_GRAY = "F3F3F4"
MID_GRAY = "D9D9D9"
TEXT = "20262B"
MUTED = "5B6670"
WHITE = "FFFFFF"


def set_run_font(run, size=10.5, bold=None, italic=None, color=TEXT, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=MID_GRAY, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(round(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def format_cell(cell, size=9.2, bold=False, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        for run in paragraph.runs:
            set_run_font(run, size=size, bold=bold, color=color)


def fill_cell(cell, text, *, bold=False, size=9.2, color=TEXT, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    format_cell(cell, size=size, bold=bold, color=color, align=align)


def add_table(doc, headers, rows, widths, *, font_size=9.0, center_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for idx, header in enumerate(headers):
        set_width(hdr.cells[idx], widths[idx])
        set_cell_shading(hdr.cells[idx], CHARCOAL)
        fill_cell(
            hdr.cells[idx],
            header,
            bold=True,
            size=8.6,
            color=WHITE,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            set_width(cells[idx], widths[idx])
            if row_idx % 2:
                set_cell_shading(cells[idx], PALE_GRAY)
            fill_cell(
                cells[idx],
                value,
                size=font_size,
                align=WD_ALIGN_PARAGRAPH.CENTER if idx in center_cols else WD_ALIGN_PARAGRAPH.LEFT,
            )
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Halaman ")
    set_run_font(run, size=8.2, color=MUTED)
    spacer = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    spacer._r.addnext(field)


def remove_paragraph_borders(element):
    p_pr = element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    title._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("000000")
    title.paragraph_format.space_after = Pt(4)
    remove_paragraph_borders(title._element)

    for name, size in (("Heading 1", 13), ("Heading 2", 11)):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(11 if name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "GLD TAHAP 2  |  BERITA ACARA SERAH TERIMA PEKERJAAN"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.runs[0], size=8.1, bold=True, color=MUTED)
    section.footer.paragraphs[0].text = ""


def add_paragraph(doc, text="", *, bold_lead=None, align=None, size=10.5, italic=False):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.12
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=size, bold=True)
        tail = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(tail, size=size, italic=italic)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=size, italic=italic)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    for run in paragraph.runs:
        set_run_font(run, size=10.0)
    if not paragraph.runs:
        set_run_font(paragraph.add_run(text), size=10.0)
    else:
        paragraph.runs[0].text = text
    return paragraph


def replace_exact_paragraph_text(doc, old, new):
    changed = 0
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == old:
            if paragraph.runs:
                paragraph.runs[0].text = new
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(new)
            changed += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip() == old:
                        if paragraph.runs:
                            paragraph.runs[0].text = new
                            for run in paragraph.runs[1:]:
                                run.text = ""
                        else:
                            paragraph.add_run(new)
                        changed += 1
    return changed


def count_exact_paragraph_text(doc, target):
    count = sum(1 for paragraph in doc.paragraphs if paragraph.text.strip() == target)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                count += sum(1 for paragraph in cell.paragraphs if paragraph.text.strip() == target)
    return count


def revise_notulen():
    doc = Document(NOTULEN)
    replacements = {
        "Labtek XV, Institut Teknologi Bandung": (
            "Lab IoT, Instrumentation and Computations, Gedung Laboratorium Fisika Terpadu, "
            "Institut Teknologi Bandung"
        ),
        "Tim ITB, PT Pertamina, dan PT LAPI Ganesha Utama.": (
            "LAPI Ganesha Utama, Lab IoT/Instrumentation and Computation ITB, dan "
            "PT Pertamina Patra Niaga."
        ),
    }
    results = {old: replace_exact_paragraph_text(doc, old, new) for old, new in replacements.items()}
    for old, new in replacements.items():
        if results[old] == 0 and count_exact_paragraph_text(doc, new) < 1:
            raise RuntimeError(f"Target revisi notulen tidak ditemukan: {old}")
        if results[old] > 1:
            raise RuntimeError(f"Target revisi notulen ditemukan lebih dari satu kali: {old}")
    doc.save(NOTULEN)
    return NOTULEN


def build_bast():
    BAST.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    setup_styles(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_borders(title._element)
    set_run_font(title.add_run("BERITA ACARA SERAH TERIMA PEKERJAAN TERMIN 1"), size=20, bold=True, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    set_run_font(
        subtitle.add_run("Pengembangan dan Field Testing Sistem Gas Leak Detection Tahap 2"),
        size=11.5,
        bold=True,
        color=MUTED,
    )

    meta_rows = [
        ("Nomor BAST", "........................................................"),
        ("Nomor dan tanggal SPK", "........................................................"),
        ("Termin", "Termin 1 sebesar 20 persen dari nilai SPK"),
        ("Acuan laporan", "LGU-GLD-T1-FIT-2026-001 Revisi 0.2 tanggal 11 September 2026"),
        ("Tanggal serah terima", "........................................................"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for idx, (label, value) in enumerate(meta_rows):
        cells = table.add_row().cells
        set_width(cells[0], 5.0)
        set_width(cells[1], 11.0)
        set_cell_shading(cells[0], CHARCOAL)
        fill_cell(cells[0], label, bold=True, size=9.0, color=WHITE)
        if idx % 2:
            set_cell_shading(cells[1], PALE_GRAY)
        fill_cell(cells[1], value, size=9.1)

    doc.add_heading("Pernyataan Para Pihak", level=1)
    add_paragraph(
        doc,
        "Pada hari ............ tanggal ............ bulan ............ tahun ............, para pihak di bawah ini menyatakan telah melaksanakan pemeriksaan dan serah terima pekerjaan Termin 1 untuk pekerjaan Pengembangan dan Field Testing Sistem Gas Leak Detection Tahap 2.",
    )

    parties = [
        ("Pihak Pertama", "PT Pertamina Patra Niaga", "Nama, jabatan, dan unit kerja diisi oleh pihak berwenang"),
        ("Pihak Kedua", "LAPI Ganesha Utama", "Nama dan jabatan penanggung jawab kontrak diisi oleh pihak berwenang"),
        (
            "Saksi teknis",
            "Lab IoT/Instrumentation and Computation ITB",
            "Perwakilan teknis yang menyaksikan atau mendukung pelaksanaan pekerjaan",
        ),
    ]
    add_table(doc, ["Kedudukan", "Organisasi", "Keterangan"], parties, [3.2, 5.3, 7.5], font_size=8.8)
    add_paragraph(
        doc,
        "Dalam berita acara ini, PT Pertamina Patra Niaga selanjutnya disebut Pihak Pertama dan LAPI Ganesha Utama selanjutnya disebut Pihak Kedua. Lab IoT/Instrumentation and Computation ITB dicantumkan sebagai saksi atau pelaksana teknis dan bukan sebagai pihak kontraktual, kecuali ditentukan lain dalam dokumen SPK.",
        size=9.6,
    )

    doc.add_heading("Ruang Lingkup Serah Terima", level=1)
    add_paragraph(
        doc,
        "Pihak Kedua menyerahkan hasil pekerjaan Termin 1 pada tahap rekayasa dan integrasi laboratorium. Pekerjaan tersebut mencakup:",
    )
    for item in (
        "detail engineering dan desain sistem GLD;",
        "persiapan komponen dan perangkat untuk integrasi;",
        "konfigurasi firmware perangkat GLD, Cluster Head, Gateway, dan antarmuka server;",
        "factory integration test atau FAT/FIT pada rantai GLD-Cluster Head-Gateway-Server; dan",
        "dokumentasi instalasi serta baseline as-built prototipe atau laboratorium.",
    ):
        add_bullet(doc, item)
    add_paragraph(
        doc,
        "Pihak Pertama telah melakukan peninjauan perkembangan dan integrasi sistem di Lab IoT, Instrumentation and Computations, Gedung Laboratorium Fisika Terpadu, Institut Teknologi Bandung, sebelum kunjungan ke RU IV Cilacap. Peninjauan tersebut menjadi bagian dari bukti witness tahap FAT/FIT.",
    )

    doc.add_heading("Hasil Pemeriksaan Deliverable", level=1)
    rows = [
        ("1", "Detail engineering dan desain", "Memenuhi", "Arsitektur, datasheet, TDS, desain mounting, dan dokumen instalasi tersedia"),
        ("2", "Persiapan komponen", "Memenuhi", "Perangkat inti tersedia untuk integrasi dan pengujian laboratorium"),
        ("3", "Konfigurasi firmware", "Memenuhi", "Akuisisi sensor, inferensi lokal, komunikasi LoRa, alarm, dan integrasi terdokumentasi"),
        ("4", "Factory integration test", "Memenuhi dengan pengesahan", "Integrasi end-to-end, mesh, failover, dan alarm push telah diuji pada tahap laboratorium"),
        ("5", "Dokumen instalasi dan as-built", "Memenuhi sesuai tahap", "Baseline prototipe atau laboratorium tersedia; as-built lokasi diterbitkan setelah instalasi lapangan"),
    ]
    add_table(doc, ["No", "Deliverable", "Hasil", "Keterangan"], rows, [1.0, 4.0, 3.4, 7.6], font_size=8.3, center_cols=(0, 2))

    doc.add_heading("Batas Penerimaan", level=1)
    add_paragraph(
        doc,
        "Serah terima Termin 1 ini terbatas pada hasil detail engineering, persiapan komponen, konfigurasi firmware, factory integration test, serta baseline dokumentasi prototipe atau laboratorium. Berita acara ini tidak menyatakan bahwa instalasi, commissioning, site acceptance test, as-built lokasi RU IV Cilacap, atau sertifikasi hazardous area telah selesai.",
    )

    doc.add_heading("Nilai Pekerjaan", level=1)
    values = [
        ("Nilai SPK", "Rp ........................................................"),
        ("Nilai Termin 1", "20 persen x nilai SPK = Rp ........................................................"),
        ("Pajak dan potongan", "Sesuai ketentuan SPK dan administrasi yang berlaku"),
        ("Nilai pembayaran bersih", "Rp ........................................................"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for idx, (label, value) in enumerate(values):
        cells = table.add_row().cells
        set_width(cells[0], 5.0)
        set_width(cells[1], 11.0)
        set_cell_shading(cells[0], PALE_TEAL)
        fill_cell(cells[0], label, bold=True, size=9.0)
        if idx % 2:
            set_cell_shading(cells[1], PALE_GRAY)
        fill_cell(cells[1], value, size=9.0)

    doc.add_heading("Lampiran", level=1)
    attachments = [
        ("1", "Laporan Pemenuhan Deliverable Termin 1 Field Testing GLD Revisi 0.2", "Dokumen utama"),
        ("2", "Klausul pembayaran dan proposal pekerjaan", "Dasar kontraktual"),
        ("3", "Notulen rapat 6 Agustus 2026", "Witness dan pembahasan teknis"),
        ("4", "Datasheet sistem, TDS perangkat, arsitektur, dan desain mounting", "Detail engineering"),
        ("5", "Bukti uji LoRa, log serial, presentasi AI, dan foto unit", "Bukti teknis terkurasi"),
        ("6", "Laporan progres, JSA, checklist instalasi, dan notulen survey RU IV", "Tahap berikutnya"),
    ]
    add_table(doc, ["No", "Lampiran", "Fungsi"], attachments, [1.0, 9.5, 5.5], font_size=8.7, center_cols=(0,))

    doc.add_heading("Kesepakatan", level=1)
    clauses = [
        "Pihak Pertama telah menerima dan memeriksa dokumen serta bukti pekerjaan Termin 1 sebagaimana tercantum dalam berita acara dan lampirannya.",
        "Apabila terdapat catatan hasil pemeriksaan, catatan tersebut dituangkan pada bagian Catatan Pemeriksaan dan diselesaikan sesuai kesepakatan para pihak tanpa mengubah batas tahap yang dinyatakan dalam berita acara ini.",
        "Setelah berita acara ini ditandatangani oleh pihak yang berwenang, dokumen dapat digunakan sebagai dasar proses administrasi pembayaran Termin 1 sebesar 20 persen sesuai ketentuan SPK.",
        "Berita acara dibuat dalam rangkap yang diperlukan dan mempunyai kekuatan administrasi yang sama setelah ditandatangani.",
    ]
    for clause in clauses:
        add_bullet(doc, clause)

    doc.add_heading("Catatan Pemeriksaan", level=1)
    notes = doc.add_table(rows=3, cols=1)
    notes.alignment = WD_TABLE_ALIGNMENT.CENTER
    notes.autofit = False
    set_table_borders(notes)
    set_width(notes.cell(0, 0), 16.0)
    set_cell_shading(notes.cell(0, 0), PALE_GRAY)
    fill_cell(notes.cell(0, 0), "Diisi bila terdapat catatan atau tindak lanjut sebelum pengesahan", bold=True, size=9.0)
    for idx in (1, 2):
        fill_cell(notes.cell(idx, 0), "................................................................................................................................................", size=9.0)

    doc.add_heading("Pengesahan", level=1)
    add_paragraph(
        doc,
        "Para pihak menyatakan bahwa informasi dalam berita acara ini telah diperiksa dan disetujui sesuai kewenangan masing-masing.",
    )

    signatures = doc.add_table(rows=4, cols=2)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    signatures.autofit = False
    set_table_borders(signatures, color=WHITE, size="0")
    for cell in signatures.rows[0].cells:
        set_cell_shading(cell, CHARCOAL)
    fill_cell(signatures.cell(0, 0), "PIHAK PERTAMA", bold=True, size=9.5, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(signatures.cell(0, 1), "PIHAK KEDUA", bold=True, size=9.5, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(signatures.cell(1, 0), "PT Pertamina Patra Niaga", bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(signatures.cell(1, 1), "LAPI Ganesha Utama", bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(signatures.cell(2, 0), "\n\n\n\n", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(signatures.cell(2, 1), "\n\n\n\n", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(signatures.cell(3, 0), "Nama: ........................................\nJabatan: .....................................\nTanggal: ......................................", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(signatures.cell(3, 1), "Nama: ........................................\nJabatan: .....................................\nTanggal: ......................................", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in signatures.rows:
        for idx, cell in enumerate(row.cells):
            set_width(cell, 8.0)

    doc.add_paragraph()
    witness = doc.add_table(rows=4, cols=1)
    witness.alignment = WD_TABLE_ALIGNMENT.CENTER
    witness.autofit = False
    set_table_borders(witness, color=WHITE, size="0")
    set_width(witness.cell(0, 0), 16.0)
    set_cell_shading(witness.cell(0, 0), CHARCOAL)
    fill_cell(witness.cell(0, 0), "SAKSI TEKNIS", bold=True, size=9.5, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(witness.cell(1, 0), "Lab IoT/Instrumentation and Computation ITB", bold=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(witness.cell(2, 0), "\n\n\n", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    fill_cell(witness.cell(3, 0), "Nama: ........................................    Jabatan: ........................................    Tanggal: ........................................", size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER)

    doc.add_heading("Daftar Isian Wajib Sebelum Penandatanganan", level=1)
    checks = [
        "Nomor dan tanggal SPK telah diisi.",
        "Nilai SPK, nilai Termin 1, pajak, potongan, dan nilai bersih telah diverifikasi.",
        "Nomor BAST dan tanggal serah terima telah diisi.",
        "Nama serta jabatan seluruh penandatangan telah sesuai kewenangan.",
        "Catatan pemeriksaan telah ditutup atau disepakati para pihak.",
        "Lampiran yang disebutkan telah disertakan pada paket pengajuan.",
    ]
    for item in checks:
        add_bullet(doc, "[  ] " + item)

    doc.save(BAST)
    return BAST


def main():
    print(revise_notulen())
    print(build_bast())


if __name__ == "__main__":
    main()
