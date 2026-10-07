# -*- coding: utf-8 -*-
"""Shared content: photo set of the current production configuration (latest enclosure and PCB), 7 Oct 2026.

Source: Sumber Dokumen/Foto_Assembly_GLD_07Okt2026/ (23 photographs, assembly sequence of one unit).
Rendered into section 2.4 of both certification documents.
"""
import os

from PIL import Image
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = os.path.join(REPO, "scripts", "assets", "assembly_photos_0710")
NAVY = RGBColor(0x1B, 0x33, 0x5F)
GRAY = RGBColor(0x55, 0x5B, 0x66)

PHOTOS = [
    ("IMG_14.02.34.jpg", "Base enclosure before assembly",
     "Machined bore of the base enclosure with the PCB mounting bosses; the external antenna is fitted at the top "
     "and its internal coaxial lead (U.FL) is routed inside."),
    ("IMG_14.03.09.jpg", "Main board placed in the base enclosure",
     "Main PCB with the ESP32-S3 module and LoRa module seated on the mounting bosses."),
    ("IMG_14.03.33.jpg", "Antenna lead connected to the main board",
     "U.FL coaxial connector of the antenna lead plugged onto the main board."),
    ("IMG_14.03.58.jpg", "PCB cover / terminal board fitted",
     "Terminal board with screw terminals labelled RS485 A/B, FAN 5V/0V, ALARM, BAT −/+ and 24V +/−; eight 2×4 "
     "headers for the sensor modules and a micro-USB service port."),
    ("IMG_14.04.17.jpg", "Sensor module",
     "MQ-series sensor with its integrated stainless-steel protective mesh, mounted on a round carrier board."),
    ("IMG_14.04.34.jpg", "Sensor module being inserted", "Sensor module plugged into one of the eight headers."),
    ("IMG_14.04.49.jpg", "Sensor modules fitted", "Sensor modules installed on the terminal board positions."),
    ("IMG_14.05.01.jpg", "DC sampling fan",
     "CIXIKEJI DC brushless fan, model CX5010B5H, DC 5 V 0.23 A (label marks CE, FC, RoHS), with ferrule-terminated "
     "leads."),
    ("IMG_14.05.32.jpg", "Fan leads terminated", "Fan leads connected to the FAN 5V/0V terminal."),
    ("IMG_14.06.08.jpg", "Alarm beacon fitted",
     "Visual/audible alarm beacon (stainless-steel body, red dome) threaded into the right-hand entry; its leads "
     "are routed to the terminal board."),
    ("IMG_14.07.08.jpg", "Alarm leads terminated", "Beacon leads connected to the ALARM terminal."),
    ("IMG_14.08.05.jpg", "Cable gland fitted to the left-hand entry",
     "Cable gland of the current prototype (black polymer body) installed in the power-cable entry."),
    ("IMG_14.09.23.jpg", "24 VDC cable inserted", "24 VDC supply cable passed through the cable gland."),
    ("IMG_14.10.55.jpg", "24 VDC conductors terminated", "Supply conductors connected to the 24V +/− terminal."),
    ("IMG_14.12.07.jpg", "Cover with sampling fan",
     "Enclosure cover with the DC fan mounted directly behind the stainless-steel mesh; cover being fitted."),
    ("IMG_14.17.20.jpg", "Assembled unit — grounding",
     "Cover closed; the external grounding screw beside the cable gland is being tightened."),
    ("IMG_14.34.54.jpg", "Threaded cover being closed", "The threaded cover is screwed onto the base by hand."),
    ("IMG_14.35.32.jpg", "Assembled unit — front view",
     "Stainless-steel mesh sensing face; cable gland (left), alarm beacon (right), antenna (top)."),
    ("IMG_14.35.54.jpg", "Assembled unit — side view with blanking plug",
     "Hexagon blanking plug closing the lower entry; alarm beacon on the right."),
    ("IMG_14.36.39.jpg", "Assembled unit — side view with cable entry",
     "Cable gland, external grounding screw and antenna."),
    ("IMG_14.37.03.jpg", "Assembled unit — side view with beacon", "Alarm beacon and mounting lugs."),
    ("IMG_14.37.30.jpg", "Assembled unit — side view of the base",
     "Beacon entry (left), cable gland (right) and antenna."),
    ("IMG_14.38.01.jpg", "Assembled unit — rear view",
     "Rear face of the base with mounting lugs; antenna (top), beacon (left), cable gland (right)."),
    ("Nameplate_position_beside_antenna.jpg", "Nameplate position",
     "Flat top face of the base enclosure beside the antenna, between the antenna and the mounting lug — the "
     "intended location of the nameplate (marking still draft; not yet fitted)."),
    ("Antenna_interface_outside.jpg", "Antenna interface — outside",
     "External antenna screwed onto its connector on the top face of the base enclosure; the threaded cover joint "
     "with its black O-ring and the hexagon blanking plug with O-ring are also visible."),
    ("Antenna_interface_inside.jpg", "Antenna interface — inside the enclosure",
     "Brass SMA bulkhead seated in a spot-faced hole through the enclosure wall, with the U.FL pigtail running to "
     "the main board; PCB mounting boss below."),
    ("Cover_thread_and_entry_thread.jpg", "Cover thread and side-entry thread",
     "External thread on the base neck that carries the threaded cover (about 10 visible threads), black O-ring at "
     "the root of the thread, and an internally threaded side entry (estimated M20 × 1.5). Caliper measurements of "
     "the same neck are given in Figures 4-128 to 4-131; see Section 2.6.e.3."),
    ("Base_caliper_wall_at_Oring.jpg", "Base neck wall at the O-ring — caliper 10.50 mm",
     "Vernier caliper (0.05 mm) across the neck wall at its root, just above the O-ring: reading 10.50 mm."),
    ("Base_caliper_wall_at_top.jpg", "Base neck wall at the top — caliper 5.80 mm",
     "Caliper across the threaded neck wall at its top edge (inner bore to thread crest): reading 5.80 mm."),
    ("Base_caliper_thread_band_height.jpg", "Threaded neck height — caliper ≈14.8 mm",
     "Caliper held vertically from the base shoulder (O-ring level) to the top of the threaded neck: reading "
     "≈14.8 mm; about 10 thread crests are visible over this height."),
    ("Base_caliper_inner_bore.jpg", "Base inner bore — caliper ≈84.7 mm",
     "Inside jaws across the bore of the base at the top of the neck: reading ≈84.7 mm."),
    ("Gasket_caliper_cord_section.jpg", "Cover O-ring — cord cross-section 2.70 mm",
     "Black rubber O-ring of the cover-to-base joint measured across its cord: reading 2.70 mm."),
    ("Base_ruler_internal_depth.jpg", "Base internal depth — steel rule ≈45 mm",
     "Steel rule standing on the PCB-boss ledge inside the base: rim top at about 45 mm (rule reading, ±1 mm); "
     "the floor below the ledge is not included."),
    ("Oring_caliper_outer_diameter.jpg", "Cover O-ring — outer diameter ≈91.0 mm",
     "O-ring removed and measured across its outside with the caliper jaws: reading ≈91.0 mm (unstretched; the "
     "inner diameter is ≈85.6 mm by subtracting 2 × 2.70 mm cord)."),
    ("Cover_caliper_internal_thread_length.jpg", "Cover internal thread length — caliper ≈18.9 mm",
     "Caliper held vertically inside the cover from the cover rim to the end of the internal thread: reading "
     "≈18.9 mm."),
    ("Base_caliper_neck_height_from_body.jpg", "Base neck height from the body shoulder — caliper ≈20.2 mm",
     "Caliper from the painted body shoulder below the O-ring to the top of the threaded neck: reading ≈20.2 mm "
     "(compare ≈14.8 mm measured from the O-ring level, Figure 4-130)."),
    ("Cover_ruler_internal_depth.jpg", "Cover internal depth — steel rule ≈47.5 mm",
     "Steel rule standing on the stainless-steel mesh inside the cover: cover rim at about 47.5 mm (rule reading, "
     "±1 mm)."),
]


# (group title, first index, end index) — figure numbers are first_fig + index
PHOTO_GROUPS = [
    ("Assembly sequence — base, main board, sensors, fan, alarm, cable entry, cover (Figures 4-101 to 4-117)", 0, 17),
    ("Assembled unit — external views (Figures 4-118 to 4-123)", 17, 23),
    ("Nameplate position, antenna interface and threads (Figures 4-124 to 4-127)", 23, 27),
    ("Dimensional check — caliper measurements of the base neck, bore and O-ring cord (Figures 4-128 to 4-132)", 27, 32),
    ("Dimensional check — depths, cover thread length and O-ring diameter (Figures 4-133 to 4-137)", 32, 37),
]


def _fig_style(d):
    """Paragraph style for drawing/figure titles: outline level 6 (PDF bookmark) and source of the list of drawings."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement as _OE
    from docx.oxml.ns import qn as _qn
    name = "GLD Figure Title"
    try:
        return d.styles[name]
    except KeyError:
        st = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = d.styles["Normal"]
        st.paragraph_format.keep_with_next = True
        ppr = st.element.get_or_add_pPr()
        ol = _OE("w:outlineLvl")
        ol.set(_qn("w:val"), "5")
        ppr.append(ol)
        return st


def _group_heading(doc, text):
    para = doc.add_heading("", level=5)
    para.paragraph_format.space_before = Pt(8)
    para.paragraph_format.space_after = Pt(4)
    para.paragraph_format.keep_with_next = True
    r = para.add_run(text)
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.italic = False
    r.font.name = "Calibri"
    r.font.color.rgb = NAVY


def _fit(path, max_w, max_h):
    with Image.open(path) as im:
        w, h = im.size
    width = max_w
    if width * h / w > max_h:
        width = max_h * w / h
    return width


def render(doc, h=None, first_fig=101):
    title = "2.4.2 Current production configuration — assembly photo set (7 October 2026)"
    if h:
        h(title)
    else:
        para = doc.add_paragraph()
        r = para.add_run(title)
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = NAVY
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(6)
    ri = intro.add_run("Photographs of one unit of the current production configuration (latest enclosure and main "
                       "board), taken in assembly order from the empty base enclosure to the closed unit. No Ex "
                       "marking is shown; the MQ sensor mesh is not described as flame-arresting.")
    ri.font.size = Pt(9.5)
    for g_title, i0, i1 in PHOTO_GROUPS:
        if h:
            _group_heading(doc, g_title)
        tbl = doc.add_table(rows=0, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i in range(i0, i1, 2):
            row = tbl.add_row()
            for j in range(2):
                if i + j >= i1:
                    continue
                fn, title_, desc = PHOTOS[i + j]
                cell = row.cells[j]
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                path = os.path.join(A, fn)
                p.add_run().add_picture(path, width=Inches(_fit(path, 3.0, 3.1)))
                tp = cell.add_paragraph(style=_fig_style(doc))
                tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                tp.paragraph_format.space_after = Pt(0)
                r1 = tp.add_run(f"Figure 4-{first_fig + i + j}. GLD V2 – {title_}.")
                r1.font.bold = True
                r1.font.size = Pt(8.5)
                r1.font.color.rgb = NAVY
                cp = cell.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_after = Pt(8)
                r2 = cp.add_run(desc)
                r2.font.size = Pt(8)
                r2.font.color.rgb = GRAY
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
